"""真实材料 fixture —— 按实际会遇到的糟糕情况造，不造理想化的干净文件。"""

from __future__ import annotations

import json
from pathlib import Path

import pytest


@pytest.fixture
def messy_xlsx(tmp_path) -> Path:
    """一份真实感的梳理表：标题占前两行、表头在第 3 行、有合并单元格分组、
    关键口径写在批注里。"""
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    wb = Workbook()
    ws = wb.active
    ws.title = "业务对象实体梳理"

    ws["A1"] = "采购计划管理实体及业务规则梳理 v2"
    ws["A2"] = "填写人：供应链部　更新日期：2026-07-31"
    for c, h in enumerate(["业务域", "业务对象", "实体名称", "字段", "类型", "口径说明"],
                          start=1):
        ws.cell(row=3, column=c, value=h)

    rows = [
        ("采购计划", "采购需求计划", "pbpHeader", "plan_id", "VARCHAR", "主键"),
        (None, None, None, "plan_amount", "DECIMAL", "计划金额"),
        (None, "采购包", "purchasePackage", "package_id", "VARCHAR", "主键"),
        ("合同", "采购合同", "clmContract", "contract_id", "VARCHAR", "主键"),
        (None, None, None, "plan_amount", "DECIMAL", "计划金额"),
    ]
    for i, row in enumerate(rows, start=4):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(row=i, column=c, value=v)

    ws.merge_cells("A4:A6")  # 业务域跨三行合并
    ws.merge_cells("B4:B5")  # 业务对象跨两行合并
    # 口径藏在批注里 —— 真实梳理表里最关键的信息经常在这
    ws.cell(row=5, column=6).comment = Comment("含税，年度累计，CNY", "王明")
    ws.cell(row=8, column=6).comment = Comment("不含税，单次，CNY", "李强")

    p = tmp_path / "实体梳理.xlsx"
    wb.save(p)
    return p


@pytest.fixture
def real_ddl(tmp_path) -> Path:
    """带行内注释与外键的 DDL。口径全在 `--` 注释里。"""
    p = tmp_path / "schema.ddl"
    p.write_text(
        """-- 采购中台物理模型
CREATE TABLE pbp_header (
  plan_id      VARCHAR(32) NOT NULL PRIMARY KEY,
  plan_name    VARCHAR(200),
  plan_amount  DECIMAL(18,2),   -- 含税·年度累计·CNY
  created_at   TIMESTAMP
);

CREATE TABLE clm_contract (
  contract_id  VARCHAR(32) NOT NULL PRIMARY KEY,
  plan_id      VARCHAR(32) NOT NULL,
  plan_amount  DECIMAL(18,2),   -- 不含税·单次·CNY
  CONSTRAINT fk_plan FOREIGN KEY (plan_id) REFERENCES pbp_header(plan_id)
);
""",
        encoding="utf-8")
    return p


@pytest.fixture
def real_openapi(tmp_path) -> Path:
    p = tmp_path / "openapi.json"
    p.write_text(json.dumps({
        "openapi": "3.0.0",
        "info": {"title": "采购中台 API"},
        "paths": {
            "/purchase-plans/{id}/submit": {
                "post": {"operationId": "submitPurchasePlan", "summary": "提交采购计划",
                         "requestBody": {"content": {"application/json": {
                             "schema": {"$ref": "#/components/schemas/SubmitReq"}}}}}},
            "/purchase-packages": {
                "post": {"operationId": "createPurchasePackage", "summary": "创建采购包"}},
            "/suppliers": {"get": {"operationId": "listSuppliers"}},
        },
        "components": {"schemas": {
            "Plan": {"type": "object", "required": ["planId"], "properties": {
                "planId": {"type": "string"},
                "planAmount": {"type": "number", "format": "double",
                               "description": "计划金额（含税，年度累计）"},
                "effectiveDate": {"type": "string", "format": "date"},
                "status": {"type": "string", "enum": ["DRAFT", "SUBMITTED"]},
            }},
            "SubmitReq": {"type": "object", "properties": {"remark": {"type": "string"}}},
        }},
    }, ensure_ascii=False), encoding="utf-8")
    return p


@pytest.fixture
def gbk_csv(tmp_path) -> Path:
    """中文 CSV 常见 GBK 编码，直接按 utf-8 读会炸。"""
    p = tmp_path / "material.csv"
    rows = ["物料编码,物料名称,计量单位,单价"]
    rows += [f"M-1002{i:02d},低压电缆 YJV-4×{i},米,{12.5 + i}" for i in range(30)]
    p.write_bytes("\n".join(rows).encode("gb18030"))
    return p


@pytest.fixture
def flow_docx(tmp_path) -> Path:
    """流程说明。价值集中在规则句里 —— 基数就是从这种句子里读出来的。"""
    import docx

    d = docx.Document()
    d.core_properties.author = "wubin"
    d.add_heading("采购业务流程说明", level=1)
    d.add_heading("3.2 采购包组建", level=2)
    d.add_paragraph(
        "执行计划下达后，采购员按物料类别、供应商能力、交付窗口三个因子组包。"
        "一个执行计划可拆入多个采购包，一个采购包也可包含多个执行计划的行项。"
        "组包完成后进入询比价环节。")
    d.add_heading("4.1 金额口径", level=2)
    d.add_paragraph(
        "预算控制以年度累计含税金额为准；合同签订时按不含税单次金额记录，"
        "两者之间的换算由财务共享中心维护税率表。")
    t = d.add_table(rows=3, cols=3)
    for c, h in enumerate(["字段", "所属对象", "口径"]):
        t.cell(0, c).text = h
    t.cell(1, 0).text = "plan_amount"
    t.cell(1, 1).text = "pbpHeader"
    t.cell(1, 2).text = "含税年度累计"
    t.cell(2, 0).text = "plan_amount"
    t.cell(2, 1).text = "clmContract"
    t.cell(2, 2).text = "不含税单次"

    p = tmp_path / "流程说明.docx"
    d.save(p)
    return p


@pytest.fixture
def corpus(messy_xlsx, real_ddl, real_openapi, gbk_csv, flow_docx) -> list[Path]:
    return [messy_xlsx, real_ddl, real_openapi, gbk_csv, flow_docx]
