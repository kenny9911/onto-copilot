"""解析器 —— 针对真实材料里实际会遇到的糟糕情况。"""

from __future__ import annotations

from ontocopilot.onto.parse import (
    build_index,
    collect_endpoints,
    collect_profiles,
    corpus_summary,
    default_registry,
)
from ontocopilot.onto.parse.tabular import (
    _fill_hierarchy,
    detect_header_row,
    infer_type,
)


# ══════════════════════════════════════════════════════════════════
#  xlsx
# ══════════════════════════════════════════════════════════════════
def test_header_is_found_below_title_rows(messy_xlsx):
    """真实梳理表前面有标题和填写人信息，表头不在第 1 行。"""
    doc = default_registry().parse(messy_xlsx)
    sheet = doc.structured["sheets"][0]
    assert sheet["header_row"] == 3
    assert sheet["columns"] == ["业务域", "业务对象", "实体名称", "字段", "类型", "口径说明"]
    assert any(f.kind == "header_offset" for f in doc.findings)


def test_merged_cells_are_filled_down(messy_xlsx):
    """合并区只有左上角有值，不还原就会得到一堆空格。"""
    doc = default_registry().parse(messy_xlsx)
    rows = [c.raw for c in doc.chunks if "row" in c.tags]
    # 「采购计划」跨三行合并，每行都该带上它
    assert sum(1 for r in rows if r.get("业务域") == "采购计划") == 3
    assert sum(1 for r in rows if r.get("业务对象") == "采购需求计划") == 2


def test_blank_continuation_entity_is_carried_into_render(messy_xlsx):
    """实体名称是"写一次、下面留空继承"（不是合并单元格），plan_amount 字段行
    的 实体名称/业务对象/业务域 全靠视觉缩进继承。不还原进 render，切片进了检索
    索引也不知道自己属于哪个实体 —— 跨实体口径冲突（含税 vs 不含税）无从归因。"""
    doc = default_registry().parse(messy_xlsx)
    amt = [c for c in doc.chunks
           if "row" in c.tags and c.raw.get("字段") == "plan_amount"]
    assert len(amt) == 2
    # 两个同名字段靠 render 里的实体名区分
    by_entity = {}
    for c in amt:
        if "pbpHeader" in c.render:
            by_entity["pbpHeader"] = c
        if "clmContract" in c.render:
            by_entity["clmContract"] = c
    assert set(by_entity) == {"pbpHeader", "clmContract"}, \
        f"plan_amount 行没带上实体名：{[c.render for c in amt]}"
    # 含税·年度累计 归 pbpHeader，不含税·单次 归 clmContract
    assert "年度累计" in by_entity["pbpHeader"].render
    assert "单次" in by_entity["clmContract"].render
    # R8 三个分组列全靠继承 —— 业务对象/业务域 也要跟上
    assert "采购合同" in by_entity["clmContract"].render
    assert "合同" in by_entity["clmContract"].render


def test_cell_comments_are_extracted(messy_xlsx):
    """口径经常只写在批注里。丢了批注，口径冲突就永远发现不了。"""
    doc = default_registry().parse(messy_xlsx)
    text = "\n".join(c.render for c in doc.chunks)
    assert "含税，年度累计，CNY" in text
    assert "不含税，单次，CNY" in text
    assert "〔批注〕" in text


def test_render_pairs_binds_comment_to_its_column():
    """批注要绑在被批注的那一列后面，不是甩到行尾 —— 否则批注是哪一列的口径就丢了。"""
    from ontocopilot.onto.parse.tabular import _render_pairs
    header = ["字段", "类型", "口径"]
    row = ["plan_amount", "DECIMAL", "计划金额"]
    joined = " | ".join(_render_pairs(header, row, comments={0: "含税"}))
    assert "字段=plan_amount　〔批注〕含税" in joined


def test_middle_column_comment_binds_to_that_column(tmp_path):
    """一份批注在中间列（不是最后一列）时，也要就地绑定，而不是落在行尾。"""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    wb = Workbook()
    ws = wb.active
    ws.title = "T"
    for c, h in enumerate(["实体", "字段", "类型"], start=1):
        ws.cell(row=1, column=c, value=h)
    ws.cell(row=2, column=1, value="order")
    ws.cell(row=2, column=2, value="amount")
    ws.cell(row=2, column=3, value="DECIMAL")
    ws.cell(row=2, column=2).comment = Comment("单位：分", "a")  # 批注在中间列
    p = tmp_path / "t.xlsx"
    wb.save(p)
    doc = default_registry().parse(p)
    row = next(c for c in doc.chunks if "row" in c.tags)
    assert "字段=amount　〔批注〕单位：分" in row.render
    assert row.render.index("单位：分") < row.render.index("类型=DECIMAL")


def test_office_metadata_leak_is_reported_and_indexed(messy_xlsx):
    """作者名和绝对保存路径不在表格内容里，但常常泄漏客户名与项目代号。"""
    doc = default_registry().parse(messy_xlsx)
    leaks = [f for f in doc.findings if f.kind == "metadata_leak"]
    assert leaks and leaks[0].severity == "warn"
    # 元数据本身也要可检索 —— FDE 问"你怎么知道的"时要答得出来
    meta = [c for c in doc.chunks if "meta" in c.tags]
    assert meta and meta[0].locator["kind"] == "meta"


def test_row_chunks_carry_precise_locators(messy_xlsx):
    doc = default_registry().parse(messy_xlsx)
    row = next(c for c in doc.chunks if "row" in c.tags)
    assert row.locator["sheet"] == "业务对象实体梳理"
    assert row.cite().startswith("实体梳理.xlsx!业务对象实体梳理!R")


def test_fill_hierarchy_carries_grouping_not_optional_attribute():
    """分组列（键列左侧、更粗）向下继承；可选属性列（键列右侧）的空是"没有值"，不碰。"""
    header = ["模块", "字段", "默认值"]
    body = [["订单", "id", "0"],
            ["", "amount", ""],
            ["用户", "uid", "1"],
            ["", "name", ""]]
    out = _fill_hierarchy(header, body)
    assert [r[0] for r in out] == ["订单", "订单", "用户", "用户"]  # 分组列继承
    assert [r[2] for r in out] == ["0", "", "1", ""]              # 属性列原样
    # 输入不被就地改写
    assert body[1][0] == ""


def test_fill_hierarchy_without_record_key_is_noop():
    """没有"逐行都填、逐行变化"的记录键列 —— 不是一行一记录的表，不猜、原样返回。"""
    header = ["A", "B"]
    body = [["x", ""], ["", "y"], ["", ""]]
    assert _fill_hierarchy(header, body) == body


def test_header_detection_prefers_the_row_above_data():
    rows = [["采购计划梳理", "", ""], ["", "", ""],
            ["编码", "名称", "金额"], ["M-001", "电缆", "12.5"], ["M-002", "断路器", "88.0"]]
    assert detect_header_row(rows) == 2


def test_type_inference_from_actual_values():
    assert infer_type(["1", "2", "30"]) == "INTEGER"
    assert infer_type(["1.5", "2", "30.25"]) == "DECIMAL"
    assert infer_type(["2026-01-01", "2026/2/3"]) == "DATE"
    assert infer_type(["是", "否", "是"]) == "BOOLEAN"
    assert infer_type(["A", "B", "A", "B", "A", "B", "A", "B", "A", "B"]) == "ENUM"
    assert infer_type([]) == "STRING"


# ══════════════════════════════════════════════════════════════════
#  DDL
# ══════════════════════════════════════════════════════════════════
def test_ddl_columns_types_and_primary_keys(real_ddl):
    doc = default_registry().parse(real_ddl)
    tables = {t["name"]: t for t in doc.structured["tables"]}
    assert set(tables) == {"pbp_header", "clm_contract"}
    assert tables["pbp_header"]["primary_key"] == ["plan_id"]
    amount = next(c for c in tables["pbp_header"]["columns"] if c["name"] == "plan_amount")
    assert amount["type"] == "DECIMAL(18, 2)"
    assert next(c for c in tables["pbp_header"]["columns"]
                if c["name"] == "plan_id")["nullable"] is False


def test_ddl_inline_comments_carry_the_definition(real_ddl):
    """两张表的 plan_amount 长得一模一样，区别全在注释里。"""
    doc = default_registry().parse(real_ddl)
    tables = {t["name"]: t for t in doc.structured["tables"]}
    a = next(c for c in tables["pbp_header"]["columns"] if c["name"] == "plan_amount")
    b = next(c for c in tables["clm_contract"]["columns"] if c["name"] == "plan_amount")
    assert a["comment"] == "含税·年度累计·CNY"
    assert b["comment"] == "不含税·单次·CNY"


def test_ddl_inline_comment_syntax_is_captured(tmp_path):
    """口径不止写在 `-- 注释` 里，MySQL/Oracle 常用行内 `COMMENT '...'` —— 也要抽出来。"""
    p = tmp_path / "s.ddl"
    p.write_text("CREATE TABLE t (amount DECIMAL(18,2) COMMENT '含税·年度');",
                 encoding="utf-8")
    doc = default_registry().parse(p)
    col = next(c for c in doc.structured["tables"][0]["columns"] if c["name"] == "amount")
    assert col["comment"] == "含税·年度"


def test_composite_foreign_key_maps_columns_positionally(tmp_path):
    """复合外键 (x,y)→(x,y) 要按位配对，不能全指向第一个引用列。"""
    p = tmp_path / "s.ddl"
    p.write_text(
        "CREATE TABLE p (x VARCHAR(8), y VARCHAR(8), PRIMARY KEY (x,y));\n"
        "CREATE TABLE c (x VARCHAR(8), y VARCHAR(8),\n"
        "  CONSTRAINT fk FOREIGN KEY (x, y) REFERENCES p(x, y));",
        encoding="utf-8")
    doc = default_registry().parse(p)
    fks = doc.structured["tables"][1]["foreign_keys"]
    assert {(f["column"], f["ref_column"]) for f in fks} == {("x", "x"), ("y", "y")}


def test_ddl_emits_a_table_level_chunk(real_ddl):
    """除了 per-column/per-fk，还要有一个"这张表是干嘛的"的表级切片。"""
    doc = default_registry().parse(real_ddl)
    tbl = [c for c in doc.chunks if "table" in c.tags]
    body = " ".join(c.render for c in tbl)
    assert "pbp_header" in body and "clm_contract" in body
    assert any("plan_id" in c.render for c in tbl)  # 概览带上列名/主键


def test_constraint_lines_are_not_mistaken_for_columns(real_ddl):
    """`CONSTRAINT fk_plan FOREIGN KEY ...` 不是列定义。"""
    doc = default_registry().parse(real_ddl)
    cols = {c["name"] for t in doc.structured["tables"] for c in t["columns"]}
    assert "CONSTRAINT" not in cols and "fk_plan" not in cols


def test_foreign_keys_become_cardinality_evidence(real_ddl):
    doc = default_registry().parse(real_ddl)
    fk = doc.structured["tables"][1]["foreign_keys"]
    assert fk and fk[0]["column"] == "plan_id"
    assert fk[0]["ref_table"] == "pbp_header" and fk[0]["ref_column"] == "plan_id"
    chunk = next(c for c in doc.chunks if "fk" in c.tags)
    assert "一对多" in chunk.render


def test_bad_sql_is_reported_not_swallowed(tmp_path):
    p = tmp_path / "broken.ddl"
    p.write_text("CREATE TABLE (((", encoding="utf-8")
    doc = default_registry().parse(p)
    assert any(f.kind in ("parse_failed", "no_tables") for f in doc.findings)


# ══════════════════════════════════════════════════════════════════
#  OpenAPI
# ══════════════════════════════════════════════════════════════════
def test_write_endpoints_are_separated_from_reads(real_openapi):
    """GET 不是 Action —— 读操作不改变世界状态。"""
    doc = default_registry().parse(real_openapi)
    eps = {e["operationId"]: e for e in doc.structured["endpoints"]}
    assert eps["submitPurchasePlan"]["write"] and eps["createPurchasePackage"]["write"]
    assert not eps["listSuppliers"]["write"]
    assert collect_endpoints([doc]) and all(e["write"] for e in collect_endpoints([doc]))


def test_schemas_give_object_and_property_candidates(real_openapi):
    doc = default_registry().parse(real_openapi)
    plan = doc.structured["schemas"]["Plan"]
    assert plan["properties"]["planId"]["required"]
    assert plan["properties"]["planAmount"]["base_type"] == "DECIMAL"
    assert plan["properties"]["effectiveDate"]["base_type"] == "DATE"
    assert plan["properties"]["status"]["enum"] == ["DRAFT", "SUBMITTED"]
    # description 常常就是口径说明，必须带出来
    assert "含税" in plan["properties"]["planAmount"]["description"]


def test_openapi_schema_chunk_render_carries_description_and_enum(real_openapi):
    """property 的 description 常常就是口径，enum 是取值域 —— 都要进 render 才可检索。"""
    doc = default_registry().parse(real_openapi)
    sc = next(c for c in doc.chunks if "schema" in c.tags and "Plan" in c.render)
    assert "含税" in sc.render      # description 进 render
    assert "DRAFT" in sc.render     # enum 进 render


def test_write_endpoint_links_to_its_request_schema(real_openapi):
    """写端点→请求体 schema 是一条关系，像 DDL 外键那样单独成一等切片。"""
    doc = default_registry().parse(real_openapi)
    links = [c for c in doc.chunks if "link" in c.tags]
    assert any("submitPurchasePlan" in c.render and "SubmitReq" in c.render for c in links)


def test_openapi_null_description_does_not_crash(tmp_path):
    """description 显式为 null 时（合法 JSON）不能崩 —— 真实 spec 里常见。"""
    import json
    p = tmp_path / "o.json"
    p.write_text(json.dumps({"openapi": "3.0.0", "paths": {
        "/x": {"post": {"operationId": "doX", "description": None}}}}), encoding="utf-8")
    doc = default_registry().parse(p)
    assert any("doX" in c.render for c in doc.chunks)


def test_endpoint_chunks_carry_json_pointers(real_openapi):
    doc = default_registry().parse(real_openapi)
    ep = next(c for c in doc.chunks if "endpoint" in c.tags)
    assert ep.locator["kind"] == "json" and ep.locator["pointer"].startswith("$.paths.")


# ══════════════════════════════════════════════════════════════════
#  CSV / docx
# ══════════════════════════════════════════════════════════════════
def test_gbk_csv_is_decoded_and_flagged(gbk_csv):
    doc = default_registry().parse(gbk_csv)
    assert doc.structured["columns"] == ["物料编码", "物料名称", "计量单位", "单价"]
    assert any(f.kind == "encoding_guess" for f in doc.findings)


def test_large_csv_is_sampled_not_fully_indexed(gbk_csv):
    """几万行逐行入索引会撑爆上下文，且毫无检索价值。"""
    doc = default_registry().parse(gbk_csv)
    assert doc.structured["rows"] == 30
    assert len([c for c in doc.chunks if "sample" in c.tags]) == 20
    assert any(f.kind == "sampled" for f in doc.findings)
    prof = doc.structured["profile"]["物料编码"]
    assert prof["unique"] and prof["inferred_type"] == "STRING"
    assert doc.structured["profile"]["单价"]["inferred_type"] == "DECIMAL"


def test_docx_rule_sentences_are_tagged(flow_docx):
    """基数就是从「一个执行计划可拆入多个采购包」这种句子里读出来的。"""
    doc = default_registry().parse(flow_docx)
    rules = [c for c in doc.chunks if "rule" in c.tags]
    assert rules
    assert any("多个采购包" in c.render for c in rules)
    assert any("年度累计含税" in c.render for c in doc.chunks)


def test_docx_chunks_carry_heading_breadcrumb(flow_docx):
    """"4.1 金额口径"下的规则句要带上祖先标题"采购业务流程说明"——只带最近一级
    标题会丢掉层级语境,规则句脱离它所属的流程就容易误读。"""
    doc = default_registry().parse(flow_docx)
    amt = next(c for c in doc.chunks if "para" in c.tags and "年度累计含税" in c.render)
    assert "采购业务流程说明" in amt.render   # 顶层标题在面包屑里
    assert "金额口径" in amt.render            # 直属标题也在
    assert " > " in amt.locator.get("section", "")  # locator 记的是路径


def test_docx_tables_are_extracted_separately(flow_docx):
    doc = default_registry().parse(flow_docx)
    tbl = [c for c in doc.chunks if "table" in c.tags]
    assert len(tbl) == 2
    assert "含税年度累计" in " ".join(c.render for c in tbl)


def test_english_rule_sentences_are_tagged(tmp_path):
    """Copilot 面向很多业务域和语言 —— 英文的基数/约束句也要能被识别为规则。"""
    import docx
    d = docx.Document()
    d.add_heading("Rules", level=1)
    d.add_paragraph("Each order must reference exactly one customer; "
                    "an order can contain multiple line items.")
    p = tmp_path / "rules.docx"
    d.save(p)
    doc = default_registry().parse(p)
    rules = [c for c in doc.chunks if "rule" in c.tags]
    assert rules and any("customer" in c.render.lower() for c in rules)


def test_rule_hints_are_not_procurement_specific():
    """规则线索里不该再夹带 含税/不含税 这种某个业务域独有的词。"""
    from ontocopilot.onto.parse.text import _RULE_HINTS
    assert not _RULE_HINTS.search("金额含税")
    assert not _RULE_HINTS.search("按不含税记录")
    # 但通用建模线索照常命中
    assert _RULE_HINTS.search("每个订单必须对应一个客户")
    assert _RULE_HINTS.search("references the parent table")


def test_docx_author_leak_is_reported(flow_docx):
    doc = default_registry().parse(flow_docx)
    assert any(f.kind == "metadata_leak" for f in doc.findings)


# ══════════════════════════════════════════════════════════════════
#  语料级
# ══════════════════════════════════════════════════════════════════
def test_whole_corpus_parses_and_is_searchable(corpus):
    docs = default_registry().parse_all(corpus)
    assert len(docs) == 5
    ix = build_index(docs)
    assert len(ix) > 20

    hits = ix.search("plan_amount 计划金额 口径", top_k=8, expand=0)
    cites = [c.cite() for c in hits]
    # 同一个字段的两种口径分别来自 xlsx 批注和 DDL 注释，都要能检出来
    assert any("schema.ddl" in c for c in cites)
    body = "\n".join(h.render for h in hits)
    assert "含税" in body


def test_summary_surfaces_every_finding(corpus):
    docs = default_registry().parse_all(corpus)
    s = corpus_summary(docs)
    kinds = {f["kind"] for f in s["findings"]}
    assert {"metadata_leak", "header_offset", "encoding_guess", "sampled"} <= kinds
    assert s["chunks"] == sum(len(d.chunks) for d in docs)


def test_unknown_extension_falls_back_not_skipped(tmp_path):
    """跳过一份材料而不告诉任何人，是这类系统最阴的失败模式。"""
    p = tmp_path / "notes.xyz"
    p.write_text("一个执行计划可拆入多个采购包。", encoding="utf-8")
    doc = default_registry().parse(p)
    assert doc.chunks and "采购包" in doc.chunks[0].render


def test_profiles_are_collected_for_type_mismatch_detection(corpus):
    docs = default_registry().parse_all(corpus)
    profiles = collect_profiles(docs)
    assert any(k.endswith(".单价") for k in profiles)
    assert profiles[next(k for k in profiles if k.endswith(".单价"))]["inferred_type"] == "DECIMAL"


def test_retrieval_gives_every_matching_file_a_seat(corpus):
    """跨文件口径冲突的前提是两边都被检出来。

    纯按分数取 top-k 时中文密集的切片会把 DDL 挤掉 —— 那样「计划金额」的双口径
    就永远发现不了。按文件轮转保证来源多样性。
    """
    ix = build_index(default_registry().parse_all(corpus))
    hits = ix.search("plan_amount 计划金额 口径", top_k=8, expand=0)
    files = {h.file_name for h in hits}
    assert {"schema.ddl", "实体梳理.xlsx"} <= files, f"只检出了 {files}"

    flat = ix.search("plan_amount 计划金额 口径", top_k=8, expand=0, diversify_by_file=False)
    assert len({h.file_name for h in flat}) <= len(files), "关掉多样性后来源应更集中"
