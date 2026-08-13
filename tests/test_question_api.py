from __future__ import annotations

import io
import json
import zipfile

import httpx
import pytest

from ontocopilot import appconfig, authgate, server
from ontocopilot.kernel.events import EventKind
from ontocopilot.kernel.journal import FileJournal
from ontocopilot.onto.conflict import Conflict, ConflictKind, Option
from ontocopilot.onto.oir import (
    OIR,
    BaseType,
    ObjectType,
    OpenQuestion,
    PropertyType,
    inferred,
)
from ontocopilot.onto.questions import Decision as QuestionDecision
from ontocopilot.onto.questions import (
    Question,
    QuestionBacklog,
    QuestionPriority,
    QuestionStatus,
)
from ontocopilot.onto.template import compile_template, write_xlsx
from ontocopilot.server import SESSIONS, Session, app
from ontocopilot.store.deps import set_repo_for_tests
from ontocopilot.store.repo import MemoryRepo, QuestionRow, SessionRow


@pytest.fixture(autouse=True)
def _reset(monkeypatch, tmp_path):
    monkeypatch.delenv("ONTOCOPILOT_AUTH", raising=False)
    monkeypatch.setattr(server, "ROOT", tmp_path)
    authgate._ATTEMPTS.clear()
    appconfig._CACHE = {}
    yield
    set_repo_for_tests(None)
    SESSIONS.clear()


def _client(repo: MemoryRepo) -> httpx.AsyncClient:
    set_repo_for_tests(repo)
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t")


async def _inject(repo: MemoryRepo, sid: str = "qapi") -> Session:
    set_repo_for_tests(repo)
    await repo.create_session(SessionRow(id=sid))
    s = Session(id=sid, title="采购项目")
    s.dir.mkdir(parents=True)
    oir = OIR()
    oir.add_question(OpenQuestion(
        rid="oq_threshold", text=inferred("金额正好等于50万元是否总监审批？"),
        options=["是", "否"], owner="business.chen"))
    s.state["_oir"] = oir
    s.state["oir"] = oir.to_dict()
    SESSIONS[sid] = s
    await server._sync_question_backlog(s, oir=oir)
    return s


async def test_get_questions_includes_summary_batch_and_active_decision():
    repo = MemoryRepo()
    s = await _inject(repo)
    async with _client(repo) as c:
        first = await c.get(f"/api/sessions/{s.id}/questions")
        assert first.status_code == 200
        data = first.json()
        assert data["summary"]["total"] == 1
        assert data["nextBatch"] == ["oq_threshold"]
        assert data["questions"][0]["activeDecision"] is None

        answered = await c.post(
            f"/api/sessions/{s.id}/questions/oq_threshold/answer",
            json={"answer": "是", "answerText": "是，含边界值", "actor": "fde",
                  "idempotencyKey": "qapi-answer-1"})
        assert answered.status_code == 200, answered.text
        got = (await c.get(f"/api/sessions/{s.id}/questions")).json()["questions"][0]
        assert got["answer"] == "是"
        assert got["activeDecision"]["actor"] == "fde"


async def test_patch_assignment_status_and_cas():
    repo = MemoryRepo()
    s = await _inject(repo)
    async with _client(repo) as c:
        r = await c.patch(f"/api/sessions/{s.id}/questions/oq_threshold", json={
            "ownerUserId": "u.chen", "audienceRole": "business_owner",
            "priority": "blocking", "expected_revision": 0})
        assert r.status_code == 200, r.text
        q = r.json()["question"]
        assert (q["ownerUserId"], q["audienceRole"], q["priority"], q["version"]) == (
            "u.chen", "business_owner", "blocking", 1)
        stale = await c.patch(f"/api/sessions/{s.id}/questions/oq_threshold",
                              json={"status": "deferred", "expected_revision": 0})
        assert stale.status_code == 409
        ok = await c.patch(f"/api/sessions/{s.id}/questions/oq_threshold",
                           json={"status": "deferred", "expected_revision": 1})
        assert ok.status_code == 200 and ok.json()["question"]["status"] == "deferred"


async def test_deferred_question_answer_fails_before_any_side_effect():
    """直接 API 也必须先 reopen；失败不能留下 Decision、OIR 答案或 Revision。"""
    repo = MemoryRepo()
    s = await _inject(repo, "deferred-answer")
    async with _client(repo) as c:
        deferred = await c.patch(
            f"/api/sessions/{s.id}/questions/oq_threshold",
            json={"status": "deferred", "expected_revision": 0},
        )
        assert deferred.status_code == 200
        before_oir = s.state["_oir"].to_dict()

        response = await c.post(
            f"/api/sessions/{s.id}/questions/oq_threshold/answer",
            json={"answer": "是", "answerText": "确认值", "actor": "fde",
                  "idempotencyKey": "must-not-commit"},
        )

    assert response.status_code == 409
    assert "重新打开" in response.text
    assert await repo.list_decisions_v1(s.id) == []
    assert await repo.list_revisions(s.id) == []
    assert s.state["_oir"].to_dict() == before_oir
    assert (await repo.get_question(s.id, "oq_threshold")).status == "deferred"


async def test_blocking_question_prevents_bundle_but_not_draft_artifact_download():
    """重开的发布 blocker 不能让旧 Bundle 继续冒充当前版本。"""
    repo = MemoryRepo()
    s = await _inject(repo, "blocked-bundle")
    (s.dir / "oir.json").write_text("{}", encoding="utf-8")
    s.state["release_state"] = "RELEASED"  # 模拟刚重开问题、投影尚未重算
    async with _client(repo) as c:
        promoted = await c.patch(
            f"/api/sessions/{s.id}/questions/oq_threshold",
            json={"priority": "blocking", "expected_revision": 0},
        )
        assert promoted.status_code == 200, promoted.text
        blocked = await c.get(f"/api/sessions/{s.id}/bundle")
        draft = await c.get(f"/api/sessions/{s.id}/artifacts/oir.json")

    assert blocked.status_code == 409
    assert "阻塞问题" in blocked.text
    assert draft.status_code == 200 and draft.content == b"{}"


async def test_nonblocking_open_question_downloads_an_explicit_draft_bundle():
    repo = MemoryRepo()
    s = await _inject(repo, "draft-bundle")
    (s.dir / "oir.json").write_text("{}", encoding="utf-8")
    s.state["release_state"] = "RELEASED"  # Ledger 中未答项应覆盖滞后投影
    async with _client(repo) as c:
        response = await c.get(f"/api/sessions/{s.id}/bundle?materials=false")

    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        manifest = json.loads(zf.read("manifest.json"))
        readme = zf.read("交付说明.md").decode("utf-8")
    assert manifest["release_state"] == "DRAFT"
    assert manifest["session"]["release_state"] == "DRAFT"
    assert "发布状态：DRAFT" in readme and "不得视为正式发布版本" in readme


@pytest.mark.parametrize("closed_status", ["deferred", "cancelled"])
async def test_last_blocker_status_change_executes_release_nodes_and_compiles(
    closed_status: str,
):
    repo = MemoryRepo()
    s = await _inject(repo, "defer-release")
    q = QuestionBacklog.from_dict(s.state["question_backlog"]).questions["oq_threshold"]
    q.priority = QuestionPriority.BLOCKING
    q.blocked_artifacts = ["ontology.package.json"]
    await repo.save_question(s.id, QuestionRow.from_domain(q))
    s.state["question_backlog"] = QuestionBacklog({q.id: q}).to_dict()

    # Establish the real suspended INTERVIEW checkpoint before the HTTP mutation.
    assert await server._resume_engagement_release(s, backlog=QuestionBacklog({q.id: q})) is False
    run_id = s.state["engagement_run_id"]
    before = {
        event.node_id for event in FileJournal(s.dir / "journal").read(run_id)
        if event.kind is EventKind.NODE_COMPLETED
    }
    assert "GAP" in before and "CANONICALIZE" not in before

    async with _client(repo) as c:
        response = await c.patch(
            f"/api/sessions/{s.id}/questions/oq_threshold",
            json={"status": closed_status, "expected_revision": 0},
        )

    assert response.status_code == 200, response.text
    completed = {
        event.node_id for event in FileJournal(s.dir / "journal").read(run_id)
        if event.kind is EventKind.NODE_COMPLETED
    }
    assert {"CANONICALIZE", "REVIEW", "EXPORT"} <= completed
    assert (s.dir / "ontology.package.json").is_file()
    assert s.state["engagement_execution"]["status"] == "completed"


async def test_idempotent_applied_answer_recovers_suspended_release_checkpoint():
    """A retry repairs a crash after Decision/Question commit but before resume."""
    repo = MemoryRepo()
    s = await _inject(repo, "answer-recovery")
    q = QuestionBacklog.from_dict(s.state["question_backlog"]).questions["oq_threshold"]
    q.priority = QuestionPriority.BLOCKING
    q.blocked_artifacts = ["ontology.package.json"]
    await repo.save_question(s.id, QuestionRow.from_domain(q))
    backlog = QuestionBacklog({q.id: q})
    s.state["question_backlog"] = backlog.to_dict()

    assert await server._resume_engagement_release(s, backlog=backlog) is False
    run_id = s.state["engagement_run_id"]

    decision = QuestionDecision(
        id="dec.recovery",
        question_id=q.id,
        answer="是",
        actor="fde",
        affected_ids=[],
        idempotency_key="recover-answer",
    )
    decision.metadata["status"] = "claimed"
    await repo.record_decision_v1(
        s.id, server.DecisionRecordRow.from_domain(decision),
    )
    await repo.finalize_decision_v1(s.id, decision.id, status="applied")
    q.transition(QuestionStatus.ANSWERED)
    await repo.save_question(s.id, QuestionRow.from_domain(q), expected_version=0)

    async with _client(repo) as c:
        response = await c.post(
            f"/api/sessions/{s.id}/questions/oq_threshold/answer",
            json={"answer": "是", "actor": "fde",
                  "idempotencyKey": "recover-answer"},
        )

    assert response.status_code == 200, response.text
    assert response.json()["created"] is False
    assert response.json()["pending"] == 0
    assert (s.dir / "ontology.package.json").is_file()
    completed = {
        event.node_id for event in FileJournal(s.dir / "journal").read(run_id)
        if event.kind is EventKind.NODE_COMPLETED
    }
    assert {"CANONICALIZE", "REVIEW", "EXPORT"} <= completed


async def test_patch_rejects_answered_without_decision_and_accepts_legacy_medium():
    repo = MemoryRepo()
    s = await _inject(repo)
    async with _client(repo) as c:
        answered = await c.patch(
            f"/api/sessions/{s.id}/questions/oq_threshold",
            json={"status": "answered", "expected_revision": 0})
        assert answered.status_code == 400
        assert not await repo.list_decisions_v1(s.id)
        medium = await c.patch(
            f"/api/sessions/{s.id}/questions/oq_threshold",
            json={"priority": "medium", "expected_revision": 0})
        assert medium.status_code == 200
        assert medium.json()["question"]["priority"] == "normal"


async def test_answer_is_idempotent_reopen_and_revision_visible():
    repo = MemoryRepo()
    s = await _inject(repo)
    payload = {"answer": "是", "actor": "fde", "idempotencyKey": "same-answer"}
    async with _client(repo) as c:
        one = await c.post(f"/api/sessions/{s.id}/questions/oq_threshold/answer",
                           json=payload)
        two = await c.post(f"/api/sessions/{s.id}/questions/oq_threshold/answer",
                           json=payload)
        assert one.status_code == 200 and one.json()["created"] is True
        assert two.status_code == 200 and two.json()["created"] is False
        assert len(await repo.list_decisions_v1(s.id)) == 1
        assert len(await repo.list_revisions(s.id)) == 1
        reopened = await c.post(f"/api/sessions/{s.id}/questions/oq_threshold/reopen", json={})
        assert reopened.status_code == 200
        assert reopened.json()["question"]["status"] == "open"
        revisions = (await c.get(f"/api/sessions/{s.id}/revisions")).json()
        assert revisions["count"] == 1 and revisions["current"] == 1


async def test_multiple_blockers_resume_only_after_the_last_http_answer():
    repo = MemoryRepo()
    s = await _inject(repo, "multi-release")
    s.state["_oir"].add_question(OpenQuestion(
        rid="oq_second", text=inferred("审批组织范围是什么？"), options=["集团", "公司"],
    ))
    s.state["oir"] = s.state["_oir"].to_dict()
    backlog = await server._sync_question_backlog(s, oir=s.state["_oir"])
    for q in backlog.questions.values():
        q.priority = QuestionPriority.BLOCKING
        q.blocked_artifacts = ["ontology.package.json"]
        await repo.save_question(s.id, QuestionRow.from_domain(q))
    s.state["question_backlog"] = backlog.to_dict()

    assert await server._resume_engagement_release(s, backlog=backlog) is False
    run_id = s.state["engagement_run_id"]
    async with _client(repo) as c:
        first = await c.post(
            f"/api/sessions/{s.id}/questions/oq_threshold/answer",
            json={"answer": "是", "idempotencyKey": "multi-first"},
        )
        assert first.status_code == 200, first.text
        assert first.json()["pending"] == 1
        assert not (s.dir / "ontology.package.json").exists()
        journal_completed = {
            event.node_id for event in FileJournal(s.dir / "journal").read(run_id)
            if event.kind is EventKind.NODE_COMPLETED
        }
        assert "CANONICALIZE" not in journal_completed

        last = await c.post(
            f"/api/sessions/{s.id}/questions/oq_second/answer",
            json={"answer": "集团", "idempotencyKey": "multi-last"},
        )

    assert last.status_code == 200, last.text
    assert last.json()["pending"] == 0
    completed = {
        event.node_id for event in FileJournal(s.dir / "journal").read(run_id)
        if event.kind is EventKind.NODE_COMPLETED
    }
    assert {"CANONICALIZE", "REVIEW", "EXPORT"} <= completed
    package = json.loads((s.dir / "ontology.package.json").read_text())
    assert package["revision"] >= 1
    assert len(package["decisions"]) == 2
    assert {row["answer"] for row in package["decisions"]} == {"是", "集团"}


async def test_question_exports_are_downloadable_and_saved_as_artifacts():
    repo = MemoryRepo()
    s = await _inject(repo)
    async with _client(repo) as c:
        for fmt, marker in (("json", b"questions"), ("md", "待澄清问题".encode()),
                            ("xlsx", b"PK")):
            r = await c.get(f"/api/sessions/{s.id}/questions/export?format={fmt}")
            assert r.status_code == 200 and marker in r.content[:500]
            assert "filename*=" in r.headers["content-disposition"]
    assert {"问题清单.json", "问题清单.md", "问题清单.xlsx"} <= {
        p.name for p in s.dir.iterdir()}
    data = json.loads((s.dir / "问题清单.json").read_text())
    assert data["summary"]["total"] == 1


async def test_conversation_tools_can_show_and_answer_unified_backlog():
    repo = MemoryRepo()
    s = await _inject(repo, "chat-question")
    tools = server._converse_tools(s)

    class Ctx:
        def __init__(self) -> None:
            self.approved = True
            self.pending: list = []
            self.rec = None
            self.turn_id = "turn.chat.1"

    shown = await tools.call("question.next", {"limit": 5}, Ctx(), scope="converse")
    assert shown["questionIds"] == ["oq_threshold"]
    assert any(e.get("kind") == "ui.table" and e.get("rows") for e in s.events)

    answered = await tools.call(
        "question.answer",
        {"question_id": "oq_threshold", "answer": "是", "answer_text": "包含边界值"},
        Ctx(),
        scope="converse",
    )
    assert answered["created"] is True
    assert answered["pending"] == 0
    assert s.state["_oir"].questions["oq_threshold"].answer.value == "包含边界值"
    assert len(await repo.list_decisions_v1(s.id)) == 1
    assert len(await repo.list_revisions(s.id)) == 1

    doc, err = await server._export_doc(s, "questions", "", "")
    assert not err and doc is not None
    assert doc.title == "待澄清问题"


async def test_repo_backlog_can_be_served_after_memory_state_is_rebuilt():
    repo = MemoryRepo()
    await repo.create_session(SessionRow(id="persisted"))
    q = Question(id="q.persisted", text="发票校验由谁负责？")
    await repo.upsert_questions("persisted", [QuestionRow.from_domain(q)])
    s = Session(id="persisted")
    s.dir.mkdir(parents=True)
    s.state["question_backlog"] = QuestionBacklog({q.id: q}).to_dict()
    SESSIONS[s.id] = s
    async with _client(repo) as c:
        r = await c.get("/api/sessions/persisted/questions")
    assert r.status_code == 200
    assert r.json()["questions"][0]["id"] == "q.persisted"


async def test_legacy_split_answer_replay_applies_effect_and_revision_once():
    repo = MemoryRepo()
    set_repo_for_tests(repo)
    await repo.create_session(SessionRow(id="split"))
    s = Session(id="split")
    s.dir.mkdir(parents=True)
    oir = OIR()
    oir.add_property(PropertyType(
        rid="pt_amount", parent="ot_plan", api_name=inferred("planAmount"),
        display_name=inferred("计划金额"), base_type=inferred(BaseType.DECIMAL),
        definition=inferred("含税，年度累计")))
    conflict = Conflict(
        "cf_split", ConflictKind.SEMANTIC_DIVERGENCE, ["pt_amount"],
        "两处金额口径应拆分",
        options=[Option("split", "拆分为两个属性",
                        effect={"split": ["pt_amount"]})])
    s.state.update({
        "_oir": oir, "oir": oir.to_dict(), "_conflicts": [conflict],
        "questions": [{
            "id": "q_split", "conflict_rid": "cf_split", "title": conflict.summary,
            "options": [{"id": "split", "label": "拆分为两个属性"}],
        }],
    })
    SESSIONS[s.id] = s
    await server._sync_question_backlog(
        s, oir=oir, clarification=s.state["questions"], conflicts=[conflict])

    payload = {"conflict_rid": "cf_split", "option_id": "split", "note": "业务已确认"}
    async with _client(repo) as c:
        first = await c.post(f"/api/sessions/{s.id}/answer", json=payload)
        assert first.status_code == 200, first.text
        after_first = oir.properties["pt_amount"].api_name.value
        second = await c.post(f"/api/sessions/{s.id}/answer", json=payload)
        assert second.status_code == 200, second.text

    assert after_first == "planAmountTaxInclAnnual"
    assert oir.properties["pt_amount"].api_name.value == after_first
    assert first.json()["created"] is True and second.json()["created"] is False
    assert len(await repo.list_decisions_v1(s.id)) == 1
    assert len(await repo.list_revisions(s.id)) == 1


async def test_return_template_preview_is_read_only_then_apply_creates_revision(tmp_path):
    """回传件必须先预审、再明确应用；apply=false 不能偷改活 OIR。"""
    from openpyxl import load_workbook

    repo = MemoryRepo()
    s = await _inject(repo, "audit-api")
    obj = ObjectType(rid="ot_plan", api_name=inferred("purchasePlan"),
                     display_name=inferred("采购计划"), primary_key=inferred([]))
    s.state["_oir"].add_object(obj)
    s.state["oir"] = s.state["_oir"].to_dict()
    spec = compile_template(s.state["_oir"])
    spec.save(s.dir / "template.spec.json")
    returned = write_xlsx(spec, tmp_path / "returned.xlsx", project="T")
    wb = load_workbook(returned)
    changed_cell = next(c for c in spec.cells()
                        if c.rid == "ot_plan" and c.field == "description")
    ws = wb[changed_cell.sheet]
    # 按隐藏 rid + 表头定位，不依赖行号猜测。
    headers = {str(ws.cell(row=2, column=i).value): i
               for i in range(1, ws.max_column + 1)}
    row = next(r for r in range(3, ws.max_row + 1)
               if str(ws.cell(row=r, column=headers["_oir_rid"]).value) == "ot_plan")
    ws.cell(row=row, column=headers["description"], value="由采购部编制的年度需求计划")
    wb.save(returned)
    original = s.state["_oir"].objects["ot_plan"].description.value

    # 能按移动后的锚点读到部分内容也不代表安全：表结构一旦损伤，apply
    # 必须整体 fail closed，不能生成一个静默漏答复的 revision。
    damaged = tmp_path / "returned-damaged.xlsx"
    damaged.write_bytes(returned.read_bytes())
    damaged_wb = load_workbook(damaged)
    damaged_wb[changed_cell.sheet].insert_rows(1)
    damaged_wb.save(damaged)

    async with _client(repo) as c:
        damaged_blob = damaged.read_bytes()
        damaged_preview = await c.post(
            f"/api/sessions/{s.id}/audit?apply=false",
            files={"files": ("returned-damaged.xlsx", damaged_blob,
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert damaged_preview.status_code == 200, damaged_preview.text
        assert damaged_preview.json()["readable"] is False
        rejected = await c.post(
            f"/api/sessions/{s.id}/audit?apply=true",
            files={"files": ("returned-damaged.xlsx", damaged_blob,
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert rejected.status_code == 422, rejected.text
        assert rejected.json()["detail"]["code"] == "RETURN_TEMPLATE_DAMAGED"
        assert s.state["_oir"].objects["ot_plan"].description.value == original
        assert not await repo.list_revisions(s.id)

        blob = returned.read_bytes()
        preview = await c.post(
            f"/api/sessions/{s.id}/audit?apply=false",
            files={"files": ("returned.xlsx", blob,
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert preview.status_code == 200, preview.text
        assert preview.json()["applied"] is False
        assert preview.json()["diffs"]
        assert s.state["_oir"].objects["ot_plan"].description.value == original
        assert not await repo.list_revisions(s.id)

        applied = await c.post(
            f"/api/sessions/{s.id}/audit?apply=true",
            files={"files": ("returned.xlsx", blob,
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert applied.status_code == 200, applied.text
        data = applied.json()
        assert data["applied"] is True and data["created"] is True
        assert "ot_plan.description" in data["changed"]
        assert s.state["_oir"].objects["ot_plan"].description.value.startswith("由采购部")
        assert len(await repo.list_revisions(s.id)) == 1

        duplicate = await c.post(
            f"/api/sessions/{s.id}/audit?apply=true",
            files={"files": ("returned.xlsx", blob,
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert duplicate.status_code == 200
        assert duplicate.json()["created"] is False
        assert len(await repo.list_revisions(s.id)) == 1


async def test_return_template_retry_resumes_proposed_revision_after_compile_failure(
        tmp_path, monkeypatch):
    """编译中断不能把 proposed 冒充 applied；同一文件重试只生成一个版本。"""
    from openpyxl import load_workbook

    repo = MemoryRepo()
    s = await _inject(repo, "audit-recovery")
    s.state["_oir"].add_object(ObjectType(
        rid="ot_plan", api_name=inferred("purchasePlan"),
        display_name=inferred("采购计划"), primary_key=inferred([])))
    s.state["oir"] = s.state["_oir"].to_dict()
    spec = compile_template(s.state["_oir"])
    spec.save(s.dir / "template.spec.json")
    returned = write_xlsx(spec, tmp_path / "recovery.xlsx", project="T")
    wb = load_workbook(returned)
    cell = next(c for c in spec.cells()
                if c.rid == "ot_plan" and c.field == "description")
    ws = wb[cell.sheet]
    headers = {str(ws.cell(row=2, column=i).value): i
               for i in range(1, ws.max_column + 1)}
    row = next(r for r in range(3, ws.max_row + 1)
               if str(ws.cell(row=r, column=headers["_oir_rid"]).value) == "ot_plan")
    ws.cell(row=row, column=headers["description"], value="恢复后仍应只应用一次")
    wb.save(returned)
    blob = returned.read_bytes()

    real_recompile = server._recompile
    calls = 0

    async def fail_once(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise RuntimeError("injected compile failure")
        return await real_recompile(*args, **kwargs)

    monkeypatch.setattr(server, "_recompile", fail_once)
    async with _client(repo) as c:
        with pytest.raises(RuntimeError, match="injected compile failure"):
            await c.post(
                f"/api/sessions/{s.id}/audit?apply=true",
                files={"files": ("recovery.xlsx", blob,
                                  "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        revisions = await repo.list_revisions(s.id)
        assert len(revisions) == 1 and revisions[0].status == "proposed"
        retry = await c.post(
            f"/api/sessions/{s.id}/audit?apply=true",
            files={"files": ("recovery.xlsx", blob,
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})

    assert retry.status_code == 200, retry.text
    assert retry.json()["applied"] is True
    revisions = await repo.list_revisions(s.id)
    assert len(revisions) == 1 and revisions[0].status == "applied"
