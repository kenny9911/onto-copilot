"""FDE 前线旅程的跨层验收。

这些用例不接真实模型：它们把已抽取 OIR / Flow 当作确定性输入，验证 FDE 真正要做的
那条链——分批访谈、记录业务原话、对话修正、生成稳定 JSON、回传预审/应用及下载。
"""

from __future__ import annotations

import asyncio

import httpx
import pytest

from ontocopilot import appconfig, authgate, server
from ontocopilot.kernel.ids import sha256_hex
from ontocopilot.kernel.intent import Intent, IntentMatch
from ontocopilot.kernel.tools import Danger
from ontocopilot.onto import build_package, validate_package
from ontocopilot.onto.flow import EdgeKind, FlowGraph, FlowNode, NodeKind
from ontocopilot.onto.flow_edit import apply_flow_edit
from ontocopilot.onto.oir import (
    OIR,
    BusinessRule,
    ObjectType,
    OpenQuestion,
    Origin,
    RuleKind,
    inferred,
)
from ontocopilot.onto.questions import Question, QuestionBacklog, QuestionPriority
from ontocopilot.onto.template import compile_template, write_xlsx
from ontocopilot.server import SESSIONS, Session, app
from ontocopilot.store.deps import set_repo_for_tests
from ontocopilot.store.repo import FileRow, MemoryRepo, QuestionRow, SessionRow


@pytest.fixture(autouse=True)
def _isolated_app(monkeypatch, tmp_path):
    monkeypatch.delenv("ONTOCOPILOT_AUTH", raising=False)
    monkeypatch.delenv("ONTOCOPILOT_ENABLE_CODEACT", raising=False)
    monkeypatch.setattr(server, "ROOT", tmp_path)
    authgate._ATTEMPTS.clear()
    appconfig._CACHE = {}
    yield
    set_repo_for_tests(None)
    SESSIONS.clear()


def _client(repo: MemoryRepo) -> httpx.AsyncClient:
    set_repo_for_tests(repo)
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t")


async def _session(repo: MemoryRepo, sid: str) -> Session:
    set_repo_for_tests(repo)
    await repo.create_session(SessionRow(id=sid))
    s = Session(id=sid, title="采购前线访谈", project="采购协同")
    s.dir.mkdir(parents=True)
    SESSIONS[sid] = s
    return s


async def test_progressive_interview_keeps_awaiting_and_surfaces_next_batch():
    """回答首批三条后不能假装访谈完成，API 必须给出下一批。"""
    repo = MemoryRepo()
    s = await _session(repo, "progressive")
    oir = OIR()
    for i in range(6):
        oir.add_question(OpenQuestion(
            rid=f"oq_{i}", text=inferred(f"第 {i + 1} 个业务口径是什么？"),
        ))
    s.state.update({"_oir": oir, "oir": oir.to_dict()})
    await server._sync_question_backlog(s, oir=oir)
    s.status = "awaiting_answer"

    async with _client(repo) as c:
        first = (await c.get(f"/api/sessions/{s.id}/questions?limit=3")).json()
        assert first["nextBatch"] == ["oq_0", "oq_1", "oq_2"]
        for i, qid in enumerate(first["nextBatch"]):
            r = await c.post(f"/api/sessions/{s.id}/questions/{qid}/answer", json={
                "answer": f"业务原话-{i}", "actor": "fde",
                "idempotencyKey": f"progressive-{i}",
            })
            assert r.status_code == 200, r.text

        after = (await c.get(f"/api/sessions/{s.id}/questions?limit=3")).json()

    assert s.status == "awaiting_answer"
    assert after["summary"]["answered"] == 3
    assert after["nextBatch"] == ["oq_3", "oq_4", "oq_5"]


async def test_free_text_answer_is_user_provenance_idempotent_and_survives_hydration():
    """业务原话落到 OIR 时必须是 USER/human；重试和进程重载不能重复建版本。"""
    repo = MemoryRepo()
    s = await _session(repo, "durable-answer")
    oir = OIR()
    oir.add_question(OpenQuestion(rid="oq_boundary", text=inferred("50 万边界如何处理？")))
    s.state.update({"_oir": oir, "oir": oir.to_dict()})
    await server._sync_question_backlog(s, oir=oir)
    payload = {
        "answer": "正好 50 万也需要总监审批",
        "answerText": "正好 50 万也需要总监审批，以含税累计金额判断",
        "actor": "fde", "idempotencyKey": "interview-turn-17",
    }

    async with _client(repo) as c:
        first = await c.post(f"/api/sessions/{s.id}/questions/oq_boundary/answer", json=payload)
        replay = await c.post(f"/api/sessions/{s.id}/questions/oq_boundary/answer", json=payload)
    assert first.status_code == 200 and first.json()["created"] is True
    assert replay.status_code == 200 and replay.json()["created"] is False
    assertion = s.state["_oir"].questions["oq_boundary"].answer
    assert assertion.origin is Origin.USER
    assert assertion.evidence[0].extractor == "human"
    assert assertion.evidence[0].snippet.startswith("Question Decision")
    assert len(await repo.list_decisions_v1(s.id)) == 1
    assert len(await repo.list_revisions(s.id)) == 1

    # 模拟 worker/进程内会话丢失；repo + oir.json 应完整恢复工作台状态。
    SESSIONS.clear()
    restored = await server._hydrate(s.id)
    assert restored.state["_oir"].questions["oq_boundary"].answer.origin is Origin.USER
    async with _client(repo) as c:
        row = (await c.get(f"/api/sessions/{s.id}/questions")).json()["questions"][0]
    assert row["status"] == "answered"
    assert row["activeDecision"]["idempotencyKey"] == "interview-turn-17"


def test_content_hash_run_id_is_stable_but_changes_for_same_size_new_content(tmp_path):
    """同名同长度但内容改变不能误用上一轮付费模型 effect。"""
    material = tmp_path / "访谈.csv"
    material.write_bytes(b"abc")
    s = Session(id="runhash", files=[{
        "name": material.name, "path": str(material), "size": 3,
        "sha256": sha256_hex(b"abc"),
    }])
    first = server._run_id_for(s)
    assert server._run_id_for(s) == first
    material.write_bytes(b"xyz")
    s.files[0]["sha256"] = sha256_hex(b"xyz")
    assert server._run_id_for(s) != first


async def test_conversation_action_edit_and_flow_undo_leave_no_replay_ghost():
    """FDE 可在对话里补 Action；撤销人工流程节点时补丁也要一起移除。"""
    repo = MemoryRepo()
    s = await _session(repo, "conversation-edits")
    oir = OIR()
    oir.add_object(ObjectType(
        rid="ot_plan", api_name=inferred("purchasePlan"), display_name=inferred("采购计划"),
    ))
    s.state.update({"_oir": oir, "oir": oir.to_dict()})
    tools = server._converse_tools(s)

    class Ctx:
        def __init__(self) -> None:
            self.approved = True
            self.pending: list = []
            self.rec = None
            self.turn_id = "turn.edit.1"

    out = await tools.call("oir.add", {
        "op": "add_action_type", "api_name": "approvePlan",
        "applies_to": ["采购计划"], "effects": ["状态=已审批"],
    }, Ctx(), scope="converse")
    assert "已改" in out
    action = next(iter(oir.actions.values()))
    assert action.api_name.origin is Origin.USER
    assert action.applies_to == ["ot_plan"]

    g = FlowGraph()
    start = g.add_node(FlowNode("fn_start", NodeKind.ACTION, inferred("提交计划")))
    end = g.add_node(FlowNode("fn_end", NodeKind.EVENT, inferred("计划已提交")))
    g.connect(start.rid, end.rid, kind=EdgeKind.INFERRED)
    s.state["_flow"] = g
    server._push_version(s, "_flow_versions", g.to_dict())
    s.state["_flow_patch_log"] = [{
        "op": "add_node", "args": {"kind": "action", "label": "临时复核"},
    }]
    apply_flow_edit(g, "add_node", {"kind": "action", "label": "临时复核"})
    undone = await tools.call("flow.undo", {}, Ctx(), scope="converse")
    assert undone["已撤销"] is True
    assert not s.state["_flow_patch_log"]
    assert all(n.label.value != "临时复核" for n in s.state["_flow"].nodes.values())
    replay = FlowGraph()
    replay.add_node(FlowNode("fn_start", NodeKind.ACTION, inferred("提交计划")))
    assert server._replay_flow_patches(s, replay) == []
    assert all(n.label.value != "临时复核" for n in replay.nodes.values())


def test_canonical_package_links_action_event_data_object_and_rule():
    """四类 Ontology JSON 不仅都存在，还必须通过跨引用校验。"""
    oir = OIR()
    oir.add_object(ObjectType(
        rid="ot_plan", api_name=inferred("purchasePlan"), display_name=inferred("采购计划"),
    ))
    from ontocopilot.onto.oir_edit import apply_oir_edit
    apply_oir_edit(oir, "add_action_type", {
        "api_name": "approvePlan", "applies_to": ["采购计划"],
        "effects": ["状态=已审批"],
    })
    oir.add_rule(BusinessRule(
        rid="br_threshold", statement=inferred("达到 50 万需要总监审批"),
        kind=inferred(RuleKind.AUTHORITY), applies_to=["ot_plan"],
    ))
    action_rid = next(iter(oir.actions))
    g = FlowGraph()
    g.add_node(FlowNode(action_rid, NodeKind.ACTION, inferred("approvePlan"),
                        objects=["ot_plan"]))
    g.add_node(FlowNode("fn_plan_approved", NodeKind.EVENT, inferred("计划已审批"),
                        objects=["ot_plan"]))
    g.connect(action_rid, "fn_plan_approved")

    data = build_package(oir, g, package_id="fde-journey").to_dict()
    assert data["validation"]["status"] == "passed"
    assert validate_package(data).passed
    assert all(data[k] for k in ("dataObjects", "actions", "events", "rules"))
    action, event, obj, rule = (
        data["actions"][0], data["events"][0], data["dataObjects"][0], data["rules"][0]
    )
    assert action["relatedDataObjects"] == [obj["id"]]
    assert action["emits"] == [event["id"]]
    assert event["producerAction"] == action["id"]
    assert event["payload"]["dataObject"] == obj["id"]
    assert rule["scope"] == [obj["id"]]


async def test_question_download_ui_contract_and_default_tool_boundary():
    """工作台有分批访谈/三格式下载/回传二阶段，默认工具边界不含 code.exec。"""
    repo = MemoryRepo()
    s = await _session(repo, "downloads")
    backlog = QuestionBacklog({
        "q_owner": Question(
            id="q_owner", text="审批责任人是谁？", priority=QuestionPriority.BLOCKING,
            audience_role="business_owner", why="阻塞审批流程图",
        ),
    })
    s.state["question_backlog"] = backlog.to_dict()
    await repo.upsert_questions(s.id, [QuestionRow.from_domain(q)
                                       for q in backlog.questions.values()])
    server._write_question_exports(s, backlog)
    async with _client(repo) as c:
        for fmt, prefix in (("xlsx", b"PK"), ("md", b"# "), ("json", b"{")):
            response = await c.get(f"/api/sessions/{s.id}/questions/export?format={fmt}")
            assert response.status_code == 200
            assert response.content.startswith(prefix)
            assert "filename*=" in response.headers["content-disposition"]

    ui = server.UI.joinpath("index.html").read_text(encoding="utf-8")
    for contract in (
        "nextBatch", "questions/export?format=", "audit?apply=false",
        "audit?apply=true", "上传回传模板", "FDE 问题工作台",
    ):
        assert contract in ui
    catalog = server._converse_tools(s)
    names = {t.spec.name for t in catalog.for_scope("converse")}
    assert {"question.next", "question.answer", "export.file"} <= names
    assert "code.exec" not in names
    assert catalog.get("question.answer", scope="converse").spec.danger is Danger.EXTERNAL


async def test_build_concurrent_guard_claims_status_before_background_task(monkeypatch):
    """两个并发 build 只能有一个取得 queued claim，且不接模型。"""
    repo = MemoryRepo()
    s = await _session(repo, "build-guard")
    material = s.dir / "materials" / "需求.csv"
    material.parent.mkdir()
    material.write_text("name\n采购计划\n", encoding="utf-8")
    s.files = [{
        "name": material.name, "path": str(material), "size": material.stat().st_size,
        "sha256": sha256_hex(material.read_bytes()),
    }]
    await repo.add_files(s.id, [FileRow(
        name=material.name,
        rel_path=str(material.relative_to(server.ROOT)),
        size=material.stat().st_size,
        sha256=s.files[0]["sha256"],
    )])

    async def parked_pipeline(_s, *, tier="full"):
        del _s, tier

    monkeypatch.setattr(server, "_run_pipeline", parked_pipeline)
    async with _client(repo) as c:
        first = await c.post(f"/api/sessions/{s.id}/build")
        second = await c.post(f"/api/sessions/{s.id}/build")
    assert first.status_code == 200
    assert second.status_code == 409
    assert s.status == "queued"
    await s.run_task


async def test_build_claim_never_overwrites_durable_awaiting_answer(monkeypatch):
    """另一个 worker 已把状态推进到 awaiting 时，本 worker 的旧投影也不能重跑。"""
    repo = MemoryRepo()
    s = await _session(repo, "build-awaiting")
    material = s.dir / "materials" / "访谈.csv"
    material.parent.mkdir()
    material.write_text("name\n采购计划\n", encoding="utf-8")
    s.files = [{
        "name": material.name, "path": str(material), "size": material.stat().st_size,
        "sha256": sha256_hex(material.read_bytes()),
    }]
    # 模拟跨 worker：数据库已进入悬挂点，本 worker 内存仍误以为是 done。
    await repo.set_status(s.id, "awaiting_answer", error="还有 3 个口径待确认")
    s.status = "done"

    async def should_not_start(_s, *, tier="full"):
        pytest.fail(f"awaiting_answer 仍启动了 {tier} build: {_s.id}")

    monkeypatch.setattr(server, "_run_pipeline", should_not_start)
    async with _client(repo) as c:
        response = await c.post(f"/api/sessions/{s.id}/build")

    assert response.status_code == 409
    assert "等待业务回答" in response.text
    row = await repo.get_session(s.id)
    assert (row.status, row.error) == ("awaiting_answer", "还有 3 个口径待确认")
    assert s.run_task is None


async def test_chat_start_build_uses_the_same_durable_claim(monkeypatch):
    """自然语言入口不能绕过按钮入口的跨 worker claim。"""
    repo = MemoryRepo()
    s = await _session(repo, "chat-build-claim")
    s.files = [{"name": "访谈.csv", "path": "访谈.csv", "size": 1, "sha256": "a" * 64}]
    await repo.set_status(s.id, "queued")  # 另一个 worker 已抢到
    s.status = "done"                       # 本 worker 的旧投影

    async def should_not_start(_s, *, tier="full"):
        pytest.fail(f"duplicate {tier} build: {_s.id}")

    monkeypatch.setattr(server, "_run_pipeline", should_not_start)
    reply = await server._act(s, IntentMatch(intent=Intent.START_BUILD))

    assert reply == "已经在跑了。"
    assert s.status == "queued"
    assert s.run_task is None


@pytest.mark.parametrize("tool_name, expected_tier", [
    ("build.start", "full"),
    ("flow.preview", "flow_preview"),
])
async def test_two_workers_cannot_duplicate_conversation_tool_build(
    monkeypatch, tool_name, expected_tier,
):
    """ConversationAgent tools must share the same repository lease across workers."""
    repo = MemoryRepo()
    await repo.create_session(SessionRow(id="tool-workers"))
    files = [{"name": "访谈.csv", "path": "访谈.csv", "size": 1,
              "sha256": "a" * 64}]
    material = server.ROOT / "tool-workers" / "materials" / "访谈.csv"
    material.parent.mkdir(parents=True)
    material.write_text("x", encoding="utf-8")
    await repo.add_files("tool-workers", [FileRow(
        name="访谈.csv", rel_path="tool-workers/materials/访谈.csv",
        size=1, sha256="a" * 64,
    )])
    worker_a = Session(id="tool-workers", files=list(files))
    worker_b = Session(id="tool-workers", files=list(files))
    started: list[tuple[Session, str]] = []
    release = asyncio.Event()

    async def parked_pipeline(session, *, tier="full"):
        started.append((session, tier))
        await release.wait()

    class Ctx:
        def __init__(self) -> None:
            self.approved = True
            self.pending: list = []
            self.rec = None
            self.turn_id = "turn.tool.lease"

    monkeypatch.setattr(server, "_run_pipeline", parked_pipeline)
    set_repo_for_tests(repo)
    first, second = await asyncio.gather(
        server._converse_tools(worker_a).call(
            tool_name, {}, Ctx(), scope="converse",
        ),
        server._converse_tools(worker_b).call(
            tool_name, {}, Ctx(), scope="converse",
        ),
    )
    try:
        results = (first, second)
        assert sum("error" not in result for result in results) == 1
        assert sum(result.get("error", "").startswith("已经在跑")
                   for result in results) == 1
        await asyncio.sleep(0)
        assert len(started) == 1
        assert started[0][1] == expected_tier
        assert (await repo.get_session("tool-workers")).status == "queued"
    finally:
        release.set()
        for session in (worker_a, worker_b):
            if session.run_task is not None:
                await session.run_task


async def test_returned_template_preview_then_apply_is_a_single_revision(tmp_path):
    """回传先预审不改 live OIR；明确 apply 后才形成唯一 revision。"""
    from openpyxl import load_workbook

    repo = MemoryRepo()
    s = await _session(repo, "returned")
    oir = OIR()
    oir.add_object(ObjectType(
        rid="ot_plan", api_name=inferred("purchasePlan"), display_name=inferred("采购计划"),
    ))
    s.state.update({"_oir": oir, "oir": oir.to_dict()})
    spec = compile_template(oir)
    spec.save(s.dir / "template.spec.json")
    returned = write_xlsx(spec, tmp_path / "returned.xlsx", project="采购协同")
    wb = load_workbook(returned)
    cell = next(c for c in spec.cells() if c.rid == "ot_plan" and c.field == "description")
    ws = wb[cell.sheet]
    headers = {str(ws.cell(row=2, column=i).value): i for i in range(1, ws.max_column + 1)}
    row = next(r for r in range(3, ws.max_row + 1)
               if str(ws.cell(row=r, column=headers["_oir_rid"]).value) == "ot_plan")
    ws.cell(row=row, column=headers["description"], value="采购部编制的年度需求计划")
    wb.save(returned)
    before = oir.to_dict()
    upload = {"files": ("returned.xlsx", returned.read_bytes(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    async with _client(repo) as c:
        preview = await c.post(f"/api/sessions/{s.id}/audit?apply=false", files=upload)
        assert preview.status_code == 200, preview.text
        assert preview.json()["applied"] is False
        assert oir.to_dict() == before
        assert not await repo.list_revisions(s.id)
        applied = await c.post(f"/api/sessions/{s.id}/audit?apply=true", files=upload)
        replay = await c.post(f"/api/sessions/{s.id}/audit?apply=true", files=upload)

    assert applied.status_code == 200 and applied.json()["created"] is True
    assert replay.status_code == 200 and replay.json()["created"] is False
    assert oir.objects["ot_plan"].description.value.startswith("采购部")
    assert len(await repo.list_revisions(s.id)) == 1
