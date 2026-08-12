"""流程图 ↔ 接口清单的接缝 —— 一份材料从解析走到模板的完整回归。

这两样东西以前是两份互不相干的产物：图上看不出哪一步有系统支撑，接口清单里也
看不出这个接口落在流程的哪一环。这个用例钉住接完之后必须成立的四件事：

  1. 有接口的环节挂上了**正确的**接口（宿主对了、动词也对）；
  2. 没接口的环节变成一个问题，而不是悄悄留白；
  3. 写接口不在流程里也变成一个问题，查询接口不算；
  4. 这些问题真的进了发给业务方的那张表。
"""

from __future__ import annotations

from pathlib import Path

import pytest

from ontocopilot.kernel.memory.evidence import EvidenceIndex
from ontocopilot.onto.flow_extract import build_flow, looks_like_process, parse_steps
from ontocopilot.onto.flow_link import attach_endpoints, coverage_gaps, flow_from_actions
from ontocopilot.onto.gaps import mine_questions
from ontocopilot.onto.parse import default_registry
from ontocopilot.onto.pipeline import build_oir, segment_corpus
from ontocopilot.onto.shape import structural_extract
from ontocopilot.onto.template import compile_template

# 一段真实形态的流程说明：五段式、合并单元格、连错别字都照抄。
PROCESS = """（1）创建采购需求计划：基于已审批集采计划自动注入；
触发条件：集采计划审批完成；
输入：已审批集采计划；
输出：未审批采购需求计划；
执行者：需求申请人
（2）审批采购需求计划：对已编制的采购需求计划进行审批；
触发条件：采购需求计划已编制；
输入：采购需求计划；
输出：已审批采购需求计划；
执行着：部门领导
（3）取消采购需求计划：执行采购需求计划取消流程；
触发条件：无需继续采购；
输入：需求取消；
输出：已取消采购需求计划；
执行者：需求申请人"""

#: 实体登记表 + 接口表，两张表靠中文「业务对象」列对上 —— 真实材料就是这么写的。
ENTITIES = [("采购计划管理", "采购需求计划", "pbpHeader", "采购业务计划头"),
            ("", "", "pbpLine", "采购业务计划行"),
            ("", "", "pbpRel", "采购业务计划头行关系"),
            ("", "", "changeHistory", "变更历史记录"),
            ("", "", "pbpImport", "采购业务计划批量导入表"),
            ("", "", "pbpAllocation", "采购业务计划分摊"),
            ("采购订单管理", "采购订单", "poHeader", "采购订单头"),
            ("", "", "poLine", "采购订单行"),
            ("", "", "poTaxLine", "采购订单税行"),
            ("", "", "poRelation", "采购订单关联"),
            ("", "", "poActionLog", "采购订单操作日志")]

APIS = [("采购计划管理", "采购需求计划", "createPbp", "创建PBP", "/v1/createPbp"),
        ("", "", "queryOpenPbpHeader", "查询PBP头", "/v1/queryPbpHeader"),
        ("", "", "queryOpenPbpLine", "查询PBP行", "/v1/queryPbpLine"),
        ("", "", "cancelOpenPbp", "取消PBP", "/v1/cancelPbp"),
        ("", "", "splitPbp", "拆分PBP", "/v1/splitPbp"),
        ("采购订单管理", "采购订单", "createPo", "创建采购订单", "/v1/createPo"),
        ("", "", "updatePo", "修改采购订单", "/v1/updatePo"),
        ("", "", "queryPoHeader", "查询采购订单头", "/v1/queryPoHeader")]


@pytest.fixture
def material(tmp_path: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("业务对象实体梳理")
    ws.append(["应用模块", "业务对象", "实体编码", "实体名称"])
    for row in ENTITIES:
        ws.append(list(row))
    ws = wb.create_sheet("业务对象API梳理-行动")
    ws.append(["应用模块", "业务对象", "实体编码", "实体名称", "url"])
    for row in APIS:
        ws.append(list(row))
    ws = wb.create_sheet("业务规则")
    ws.append([PROCESS])
    p = tmp_path / "梳理.xlsx"
    wb.save(p)
    return p


def _pipeline(path: Path):
    """跑完确定性那一半：解析 → 切段 → 规则抽取 → OIR → 流程图 → 接接口。"""
    docs = [default_registry().parse(path)]
    index = EvidenceIndex()
    for d in docs:
        for c in d.chunks:
            index.add(c)
    merged: dict[str, list] = {k: [] for k in
                              ("objects", "properties", "links", "actions",
                               "rules", "questions")}
    for seg in segment_corpus(index, docs):
        rows, cites = seg.rows(index)
        pre = structural_extract(rows, cites, seg.shape, carry_in=seg.carry_in)
        for k, bucket in merged.items():
            bucket.extend(pre.get(k) or [])
    oir = build_oir(merged, index)

    steps: list = []
    for d in docs:
        for c in d.chunks:
            for val in list((c.raw or {}).values() if isinstance(c.raw, dict) else ()) \
                    + [c.render]:
                text = str(val or "")
                if looks_like_process(text):
                    known = {x.no for x in steps}
                    steps += [x for x in parse_steps(text, cite=c.cite())
                              if x.no not in known]
    g = build_flow(sorted(steps, key=lambda x: x.no), file_name=path.name)
    return docs, index, oir, g


def test_the_whole_join_holds_on_one_material(material: Path):
    docs, index, oir, g = _pipeline(material)

    # 前提：两张表接上了，接口全挂到宿主对象上
    assert oir.stats()["objects"] == len(ENTITIES)
    assert oir.stats()["actions"] == len(APIS)
    assert all(a.applies_to for a in oir.actions.values()), "接口不该有孤儿"

    report = attach_endpoints(g, oir)

    # 1. 有接口的环节挂上了正确的接口
    created = next(n for n in g.nodes.values() if n.label.value == "创建采购需求计划")
    assert "/v1/createPbp" in created.endpoint
    cancelled = next(n for n in g.nodes.values() if n.label.value == "取消采购需求计划")
    assert "/v1/cancelPbp" in cancelled.endpoint
    assert "/v1/createPbp" not in cancelled.endpoint

    gaps = coverage_gaps(report, oir)
    texts = " ".join(x.text for x in gaps)

    # 2. 没接口的环节被问出来了
    assert "审批采购需求计划" in texts

    # 3. 写接口缺席流程要问；查询接口不算
    assert "splitPbp" in texts and "createPo" in texts
    assert "queryOpenPbpHeader" not in texts and "queryPoHeader" not in texts

    # 4. 这些问题真的进了发给业务方的那张表
    for q in mine_questions(oir, docs=docs, chunks=index.all_chunks(),
                            extra_gaps=gaps):
        oir.add_question(q)
    sheet = next(s for s in compile_template(oir).sheets if "待澄清" in s.name)
    asked = " ".join(c.value for row in sheet.rows for c in row.values())
    assert "审批采购需求计划" in asked and "splitPbp" in asked
    assert "材料出处" in sheet.columns


def test_an_interface_only_material_still_gets_a_workflow(material: Path):
    """把流程说明那张表当作不存在 —— 只靠接口清单也要出得来一张图。"""
    _docs, _index, oir, _g = _pipeline(material)
    g = flow_from_actions(oir, file_name=material.name)

    lanes = {st.title for st in g.stages.values()}
    assert {"采购业务计划头", "采购订单头"} <= lanes
    labels = [n.label.value for n in g.nodes.values()]
    assert "创建PBP" in labels and "拆分PBP" in labels
    assert not any("查询" in x for x in labels), "查询接口不属于流程的任何一步"
    # 生命周期顺序是**推的**：创建排在取消前面，边一律虚线
    order = [n.label.value for n in g.nodes.values() if n.kind.value == "action"]
    assert order.index("创建PBP") < order.index("取消PBP")
    assert all(str(e.kind) == "inferred"
               for e in g.edges.values() if e.label == "推断顺序")
