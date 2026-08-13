"""把对话里的一份内容落成文件 —— md / csv / xlsx / docx / pdf。

FDE 在聊天里拿到了一张 192 行的待澄清问题表，下一句是"把这个转成 excel 给我"。
交付包解决不了这件事（那是把整个会话打成 zip），他要的是**刚刚那一份**。

这个文件盯的主要是**格式本身的坑**，因为它们全都是"跑通了、但打开是坏的"那一类：
中文变方块、Excel 把一格文本当公式、PDF 一页 3.6 MB、导出件混进交付包。
"""

from __future__ import annotations

import io
from typing import ClassVar

import pytest

from ontocopilot.onto.export import (
    Block,
    ExportDoc,
    blocks_from_markdown,
    render,
    resolve_format,
    safe_name,
    table_block,
)

TABLE_MD = """已为您列出 3 条待澄清问题。

#### 1. 核心实体定位
*（单据上的每一个字段）*

| 节点 | 编号 | 澄清问题 | 参考选项 |
| :--- | :--- | :--- | :--- |
| （1）编制集采计划 | 1 | 什么情况下使用集采？ | - |
| （1）编制集采计划 | 3 | 输入栏写了"无系统节点" | ① 全线下 ② 系统里编 |
| （2）审批 | 4 | =SUM(A1) 会不会被当成公式 | - |

- 字段信息缺失：175 个对象没有字段
  - 需要《字段梳理表》
1. 命名规范不统一

---

```json
{"a": 1, "b": [2, 3]}
```
"""


def _doc() -> ExportDoc:
    return ExportDoc(title="采购计划 待澄清问题",
                     blocks=blocks_from_markdown(TABLE_MD),
                     note="共 3 条，由 OntoCopilot 导出")


# ══════════════════════════════════════════════════════════════════
#  markdown → 块
# ══════════════════════════════════════════════════════════════════
def test_markdown_becomes_the_blocks_the_user_saw_on_screen():
    """导出的结构必须和他在屏幕上看到的一致 —— 表还是表，不是一坨文字。"""
    blocks = blocks_from_markdown(TABLE_MD)
    kinds = [b.kind for b in blocks]
    assert "table" in kinds and "heading" in kinds and "code" in kinds

    tbl = next(b for b in blocks if b.kind == "table")
    assert tbl.columns == ["节点", "编号", "澄清问题", "参考选项"]
    assert len(tbl.rows) == 3

    # 行内标记落地时去掉：docx/pdf 里没有对应表达，留着就是正文里裸的 ** 和 `
    assert all("*" not in b.text and "`" not in b.text
               for b in blocks if b.kind in ("heading", "para"))
    assert any("单据上的每一个字段" in b.text for b in blocks)

    # 代码块**不**做行内处理，JSON 要原样留着
    code = next(b for b in blocks if b.kind == "code")
    assert code.text == '{"a": 1, "b": [2, 3]}'

    # 列表保留层级和序号
    items = next(b for b in blocks if b.items).items
    assert [d for d, _, _ in items] == [0, 1, 0]
    assert items[2][1] == "1."          # 有序项记住自己的序号


def test_empty_markers_do_not_become_list_items():
    """模型经常写出只有一个 '-' 的空行；那不是一条，别在 Word 里留个空项目符号。"""
    items = [b for b in blocks_from_markdown("- \n- 真的有内容\n-  \n") if b.items]
    assert len(items) == 1 and len(items[0].items) == 1


# ══════════════════════════════════════════════════════════════════
#  五种格式都得真的能打开
# ══════════════════════════════════════════════════════════════════
@pytest.mark.parametrize("fmt", ["md", "csv", "xlsx", "docx", "pdf"])
def test_every_format_produces_a_non_trivial_file(fmt):
    data, spec = render(_doc(), fmt)
    assert len(data) > 200
    assert spec.ext == fmt
    assert safe_name("采购计划 待澄清问题", spec.ext).endswith("." + fmt)


def test_markdown_keeps_the_table_a_table():
    data, _ = render(_doc(), "md")
    text = data.decode("utf-8")
    assert "| 节点 | 编号 | 澄清问题 | 参考选项 |" in text
    assert "什么情况下使用集采？" in text


def test_csv_carries_a_bom_so_excel_does_not_show_mojibake():
    """无 BOM 的 UTF-8 csv 在 Excel 里打开是乱码 —— 而"导出的中文是乱码"
    会被当成系统坏了，不会被当成编码设置问题。"""
    data, _ = render(_doc(), "csv")
    assert data.startswith(b"\xef\xbb\xbf")
    assert "什么情况下使用集采？" in data.decode("utf-8-sig")


def test_xlsx_is_a_usable_sheet_not_just_a_dump():
    """表要能直接用：表头冻结、能筛选、列宽不至于一屏放不下。"""
    import openpyxl

    data, _ = render(_doc(), "xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.worksheets[0]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref
    assert [c.value for c in ws[1]] == ["节点", "编号", "澄清问题", "参考选项"]
    assert ws.max_row == 4                        # 表头 + 3 行
    assert all(ws.column_dimensions[c].width <= 60 for c in ("A", "B", "C", "D"))

    # 正文和表分开放：正文塞在表上面会让第一行不是表头，筛选和冻结全错位
    assert "说明" in wb.sheetnames


def test_a_cell_that_looks_like_a_formula_stays_text():
    """材料里一格 "=SUM(A1)" 被 openpyxl 当公式写进去，Excel 打开是个错误值。"""
    import openpyxl

    data, _ = render(_doc(), "xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(data)).worksheets[0]
    cell = next(c for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=SUM"))
    assert cell.data_type == "s"                  # 文本，不是公式
    assert cell.value == "=SUM(A1) 会不会被当成公式"   # 也没被加上单引号


def test_sheet_names_survive_excels_hard_rules():
    """Excel 表名 ≤31 字符、不能有 \\/*?:[]、不能重名 —— 踩中就是存不出文件。"""
    import openpyxl

    doc = ExportDoc(title="x", blocks=[
        *table_block(["a"], [["1"]], title="报表/明细[2026]:第一版 " + "长" * 40),
        *table_block(["a"], [["2"]], title="报表/明细[2026]:第一版 " + "长" * 40),
    ])
    data, _ = render(doc, "xlsx")
    names = openpyxl.load_workbook(io.BytesIO(data)).sheetnames
    assert len(names) == len(set(names))
    for n in names:
        assert len(n) <= 31 and not set(n) & set(r"\/*?:[]")


def test_docx_keeps_chinese_readable_in_word():
    """python-docx 只写 w:ascii/w:hAnsi，不设 w:eastAsia 的话 Word 会替换中文字体。"""
    import docx
    from docx.oxml.ns import qn

    data, _ = render(_doc(), "docx")
    d = docx.Document(io.BytesIO(data))
    assert d.styles["Normal"].element.rPr.rFonts.get(qn("w:eastAsia"))
    assert len(d.tables) == 1 and len(d.tables[0].rows) == 4
    assert "什么情况下使用集采？" in d.tables[0].rows[1].cells[2].text


def test_pdf_has_real_chinese_and_is_not_multiple_megabytes():
    """两件事一起验：中文画得出来（能被提取），以及做了字体子集化。

    不子集化的话一页中文 PDF 是 3.6 MB —— 整个回退字体被嵌进去。子集化后是几十 KB。
    这不是优化，是能不能发给客户的分界。
    """
    import pymupdf

    data, _ = render(_doc(), "pdf")
    assert len(data) < 400_000, "忘了 subset_fonts()：整份中文回退字体被嵌进去了"
    text = "".join(p.get_text() for p in pymupdf.open("pdf", data))
    assert "什么情况下使用集采" in text
    assert "采购计划" in text


# ══════════════════════════════════════════════════════════════════
#  格式名 / 文件名
# ══════════════════════════════════════════════════════════════════
@pytest.mark.parametrize("said,want", [
    ("excel", "xlsx"), ("Excel", "xlsx"), ("表格", "xlsx"), (".xlsx", "xlsx"),
    ("word", "docx"), ("文档", "docx"), ("markdown", "md"), ("PDF", "pdf"),
])
def test_user_words_map_to_formats(said, want):
    """用户嘴里说的是"excel"、"word"、"表格"，不是"xlsx"。"""
    assert resolve_format(said) == want


def test_unknown_format_is_rejected_not_guessed():
    assert resolve_format("pptx") == ""
    with pytest.raises(ValueError):
        render(_doc(), "pptx")


@pytest.mark.parametrize("title", ["../../etc/passwd", "a/b\\c:d*e?f", "...", ""])
def test_file_names_cannot_escape_the_export_directory(title):
    """标题是模型给的，而它会变成磁盘上的文件名 —— 这里是唯一的路径穿越入口。"""
    name = safe_name(title, "xlsx")
    assert "/" not in name and "\\" not in name
    assert not name.startswith(".")
    assert name.endswith(".xlsx")


def test_long_titles_are_trimmed_not_rejected():
    name = safe_name("澄" * 300, "pdf")
    assert len(name) < 120 and name.endswith(".pdf")


# ══════════════════════════════════════════════════════════════════
#  服务端：AI 调 export.file 的完整一条路
# ══════════════════════════════════════════════════════════════════
class _Ctx:
    approved = True
    pending: ClassVar[list] = []


async def test_ai_exports_the_table_it_just_showed(tmp_path, monkeypatch):
    """用户："把刚才那个问题清单转成 excel 给我"。

    这是这个功能存在的理由：他不想复制粘贴，也不想被告诉"你可以自己贴进 Excel"。
    """
    from ontocopilot import server

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="ex1")
    s.dir.mkdir(parents=True, exist_ok=True)
    s.state["oir"] = {"questions": [
        {"rid": f"q{i}", "text": {"value": f"第 {i} 个要澄清的问题？"},
         "answer": {"value": ""}, "code": f"Q{i:03d}"} for i in range(192)]}

    reg = server._converse_tools(s)
    listed = await reg.call("ui.table", {"kind": "questions"}, _Ctx(), scope="converse")
    assert listed["已列出"] == 192

    out = await reg.call("export.file", {"format": "excel"}, _Ctx(), scope="converse")
    assert "error" not in out, out
    assert out["表格行数"] == 192
    assert out["已生成"].endswith(".xlsx")
    assert "不要贴链接" in out["说明"]

    # 落在 exports/ 子目录 —— 不能出现在会话根目录，那里是产物
    f = s.dir / "exports" / out["已生成"]
    assert f.exists() and f.stat().st_size > 1000
    assert not any(p.suffix == ".xlsx" for p in s.dir.iterdir() if p.is_file())

    import openpyxl
    ws = openpyxl.load_workbook(f).worksheets[0]
    assert ws.max_row == 193                       # 表头 + 192，一条不少

    ev = [e for e in s.events if e["kind"] == "export.ready"][-1]
    assert ev["rows"] == 192 and ev["label"] == "Excel"


async def test_exports_do_not_pollute_the_delivery_bundle(tmp_path, monkeypatch):
    """产物列表是"会话根目录下所有文件"算出来的。导出件落在根目录，就会混进
    产物 tab、混进交付包 zip，一个叫「问题清单.xlsx」的导出还会被「下载填写模板」
    抓走 —— 那个按钮取的是第一个 .xlsx。"""
    from ontocopilot import server

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="ex2")
    s.dir.mkdir(parents=True, exist_ok=True)
    (s.dir / "模板_v1.xlsx").write_bytes(b"real artifact")
    s.state["oir"] = {"objects": [
        {"rid": "o1", "displayName": {"value": "采购包"}, "apiName": {"value": "pkg"},
         "description": {"value": ""}, "status": "candidate"}]}

    reg = server._converse_tools(s)
    await reg.call("ui.table", {"kind": "objects"}, _Ctx(), scope="converse")
    await reg.call("export.file", {"format": "xlsx"}, _Ctx(), scope="converse")

    flat = sorted(p.name for p in s.dir.iterdir() if p.is_file())
    assert flat == ["模板_v1.xlsx"], f"导出件漏进了产物目录：{flat}"


async def test_export_without_anything_to_export_says_which_step_is_missing(
        tmp_path, monkeypatch):
    """一句"导出失败"会让模型转头跟用户说"系统限制"。真实原因（还没列过表、
    还没跑梳理）是能补的，必须说出来。"""
    from ontocopilot import server

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="ex3")
    s.dir.mkdir(parents=True, exist_ok=True)
    reg = server._converse_tools(s)

    out = await reg.call("export.file", {"format": "xlsx"}, _Ctx(), scope="converse")
    assert "还没有列过表" in out["error"]
    assert "material.rows" in out["下一步"]

    miss = await reg.call("export.file", {"format": "pdf", "source": "questions"},
                          _Ctx(), scope="converse")
    assert "error" in miss and "梳理还没跑过" in miss["error"]

    bad = await reg.call("export.file", {"format": "pptx"}, _Ctx(), scope="converse")
    assert "不支持的格式" in bad["error"]


async def test_export_can_narrow_to_the_rows_he_asked_about(tmp_path, monkeypatch):
    """"只导出临时表相关的那 17 条" —— 导出继承对话上下文才有意义。"""
    from ontocopilot import server

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="ex4")
    s.dir.mkdir(parents=True, exist_ok=True)
    s.state["oir"] = {"rules": [
        {"rid": f"r{i}", "statement": {"value": f"临时表规则 {i}" if i < 17
                                       else f"普通规则 {i}"},
         "ruleKind": {"value": "constraint"}, "actor": {"value": ""}}
        for i in range(40)]}

    reg = server._converse_tools(s)
    out = await reg.call("export.file",
                         {"format": "csv", "source": "rules", "contains": "临时表"},
                         _Ctx(), scope="converse")
    assert out["表格行数"] == 17
    body = (s.dir / "exports" / out["已生成"]).read_bytes().decode("utf-8-sig")
    assert "临时表规则 3" in body and "普通规则 20" not in body


async def test_export_the_last_answer_and_the_conversation(tmp_path, monkeypatch):
    """"把你刚才那段回答存成 word 我发给客户"。"""
    from ontocopilot import server
    from ontocopilot.kernel.memory.dialogue import Speaker

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="ex5")
    s.dir.mkdir(parents=True, exist_ok=True)
    server._publish_turn(s, Speaker.USER, "集采和普通采购怎么分？")
    server._publish_turn(s, Speaker.ASSISTANT, "## 结论\n\n按**金额阈值**分。\n\n- 超 50 万走集采\n- 其余普通采购")

    reg = server._converse_tools(s)
    one = await reg.call("export.file", {"format": "docx", "source": "last_answer"},
                         _Ctx(), scope="converse")
    assert "error" not in one, one

    import docx
    d = docx.Document(s.dir / "exports" / one["已生成"])
    text = "\n".join(p.text for p in d.paragraphs)
    assert "结论" in text and "超 50 万走集采" in text
    assert "**" not in text                        # 行内标记不该落到 Word 正文里

    whole = await reg.call("export.file", {"format": "md", "source": "conversation"},
                           _Ctx(), scope="converse")
    body = (s.dir / "exports" / whole["已生成"]).read_text(encoding="utf-8")
    assert "FDE" in body and "OntoCopilot" in body and "集采和普通采购怎么分？" in body


# ══════════════════════════════════════════════════════════════════
#  评审打出来的坑，逐条钉住
# ══════════════════════════════════════════════════════════════════
def _csv_material(path, n=60, sep=","):
    lines = [sep.join(["节点", "编号", "澄清问题"])]
    lines += [sep.join([f"（{i//10+1}）节点", str(i), f"第 {i} 个问题？"])
              for i in range(1, n + 1)]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


@pytest.mark.parametrize("ext,sep", [(".csv", ","), (".tsv", "\t")])
async def test_material_rows_actually_works_for_csv(tmp_path, monkeypatch, ext, sep):
    """工具描述和后缀白名单都写着支持 csv，但 CsvParser 压根不产 "row" 切片
    （只有一片 schema + 最多 20 片 sample），于是按 tag 过滤把每一行都丢了，
    任何 csv 都回"没读出数据行" —— 正是这个工具被写出来要消灭的那种失败。"""
    from ontocopilot import server

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="csv1")
    (s.dir / "materials").mkdir(parents=True, exist_ok=True)
    p = _csv_material(s.dir / "materials" / f"问题清单{ext}", 60, sep)
    s.files = [{"name": p.name, "size": p.stat().st_size, "path": str(p)}]

    reg = server._converse_tools(s)
    out = await reg.call("material.rows", {"file": p.name}, _Ctx(), scope="converse")
    assert "error" not in out, out
    assert out["已列出"] == 60 and out["总行数"] == 60   # 不是 20，也不是 0
    ev = [e for e in s.events if e["kind"] == "ui.table"][-1]
    assert ev["columns"] == ["节点", "编号", "澄清问题"]
    assert ev["rows"][-1][2] == "第 60 个问题？"


async def test_export_is_full_even_when_the_screen_was_truncated(tmp_path, monkeypatch):
    """界面只画前 500 行，**文件不该跟着截断** —— 尤其不该叫「（900 行）.xlsx」
    却只装 500 行。FDE 会把那份文件当完整清单发给客户。"""
    import openpyxl
    from openpyxl import Workbook

    from ontocopilot import server

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="trunc1")
    (s.dir / "materials").mkdir(parents=True, exist_ok=True)
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("问题")
    ws.append(["编号", "问题"])
    for i in range(1, 901):
        ws.append([i, f"第 {i} 个问题？"])
    p = s.dir / "materials" / "问题清单.xlsx"
    wb.save(p)
    s.files = [{"name": p.name, "size": p.stat().st_size, "path": str(p)}]

    reg = server._converse_tools(s)
    listed = await reg.call("material.rows", {"file": p.name}, _Ctx(), scope="converse")
    assert listed["已列出"] == server._ROWS_MAX and listed["总行数"] == 900
    assert "导出是全量的" in listed["只显示了前几行"]

    out = await reg.call("export.file", {"format": "xlsx"}, _Ctx(), scope="converse")
    assert out["表格行数"] == 900, "导出继承了屏幕的截断"
    assert "注意" not in out
    sheet = openpyxl.load_workbook(s.dir / "exports" / out["已生成"]).worksheets[0]
    assert sheet.max_row == 901                      # 表头 + 900
    assert sheet.cell(row=901, column=2).value == "第 900 个问题？"


async def test_export_says_so_when_it_cannot_rebuild_the_full_table(tmp_path, monkeypatch):
    """重算不出来时（原文件删了）**必须说出来**，标题里也不能留一个骗人的行数。"""
    from ontocopilot import server

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="trunc2")
    s.dir.mkdir(parents=True, exist_ok=True)
    s.emit("ui.table", title="问题清单（900 行）", columns=["问题"],
           rows=[[f"第 {i} 个"] for i in range(500)], total=900,
           src={"kind": "material", "file": "早就没了.xlsx", "sheet": "问题"})

    reg = server._converse_tools(s)
    out = await reg.call("export.file", {"format": "csv"}, _Ctx(), scope="converse")
    assert "不是全量" in out["注意"]
    # 标题里原来那个「（900 行）」要被换掉，不能变成「（900 行）（前 500 行）」
    assert "（前 500 行，原表 900 行）" in out["已生成"]
    assert out["已生成"].count("900") == 1


def test_formula_cells_are_neutralised_everywhere_not_just_data_rows():
    """内容是**别人给的**（客户的表、模型抽出来的口径）。以前只扫了数据行，
    漏掉表头和说明页 —— 而表头正是列名，恰恰来自那份外来表格。"""
    import io

    import openpyxl

    doc = ExportDoc(
        title="=1+1 报表",
        blocks=[Block("para", "=WEBSERVICE(1)"),
                *table_block(['=cmd|\'/c calc\'!A1', "正常列"],
                             [["=HYPERLINK(1)", "+1"]])],
        note="=2+2")
    data, _ = render(doc, "xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", f"{ws.title}!{c.coordinate} 还是公式"

    body = render(doc, "csv")[0].decode("utf-8-sig")
    for line in body.splitlines():
        for cell in line.split(","):
            assert not cell.strip('"').startswith(("=", "+", "@")), line


async def test_conversation_export_admits_what_compaction_dropped(tmp_path, monkeypatch):
    """DialogueMemory 每轮都压缩，旧轮次会被换成一条 system 摘要。把 system 过滤掉，
    导出的"整段对话"就从中间开始，还宣称自己是全部。"""
    from ontocopilot import server
    from ontocopilot.kernel.memory.dialogue import Speaker

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="conv1")
    s.dir.mkdir(parents=True, exist_ok=True)
    for i in range(40):
        server._publish_turn(s, Speaker.USER, f"第 {i} 个问题，" + "细节" * 60)
        server._publish_turn(s, Speaker.ASSISTANT, f"第 {i} 个回答，" + "结论" * 60)

    dm = server._dialogue(s)
    assert any(str(t.speaker) == "system" for t in dm.turns), "这轮没触发压缩，测试无效"

    reg = server._converse_tools(s)
    out = await reg.call("export.file", {"format": "md", "source": "conversation"},
                         _Ctx(), scope="converse")
    body = (s.dir / "exports" / out["已生成"]).read_text(encoding="utf-8")
    assert "早前对话摘要" in body
    assert "已被压缩" in body


async def test_exporting_the_table_the_model_wrote_in_its_answer(tmp_path, monkeypatch):
    """他截图里那张表是模型**直接写在回答正文里**的 markdown 表，不是 ui.table 卡片。

    他说"把这个表转成 excel"，而导出只认 ui.table 事件的话，就会翻出一张**别的**
    表（更早的、甚至是空的）给他 —— 那正是"下载下来是空表格"的由来。
    """

    import openpyxl

    from ontocopilot import server
    from ontocopilot.kernel.memory.dialogue import Speaker

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="mdx1")
    s.dir.mkdir(parents=True, exist_ok=True)

    # 早先有过一张别的（几乎空的）卡片表 —— 就是它被错误地导了出去
    s.state["oir"] = {"questions": [
        {"rid": "q1", "text": {"value": "a"}, "answer": {"value": ""}, "code": ""},
        {"rid": "q2", "text": {"value": "b"}, "answer": {"value": ""}, "code": ""}]}
    reg = server._converse_tools(s)
    await reg.call("ui.table", {"kind": "questions"}, _Ctx(), scope="converse")

    # 然后模型在回答正文里写了真正的那张表
    server._publish_turn(s, Speaker.USER, "请你给我列出一个表格")
    server._publish_turn(s, Speaker.ASSISTANT, """为你梳理了以下问题清单：

**采购计划 Action + Event 业务澄清问题清单**

| 序号 | 业务阶段 | 流程图中的现状 | 建议提问的问题 |
| :--- | :--- | :--- | :--- |
| 1 | 异常与逆向流程 | 当期全流程只建模整单取消 | 是否存在分取消、单行驳回？ |
| 2 | 流程边界与范围定义 | 实线为业务主干，虚线为扩展 | 本次建模边界是否包含虚线部分？ |
| 3 | 需求拆分与聚合逻辑 | 审批已通过 到 已审批需求池 为 1:N | 审批通过后是否自动拆分？ |
""")

    out = await reg.call("export.file", {"format": "excel"}, _Ctx(), scope="converse")
    assert "error" not in out, out
    assert out["表格行数"] == 3, "导的还是那张 a/b 的旧表"

    ws = openpyxl.load_workbook(s.dir / "exports" / out["已生成"]).worksheets[0]
    assert [c.value for c in ws[1]] == ["序号", "业务阶段", "流程图中的现状", "建议提问的问题"]
    assert ws.max_row == 4                       # 表头 + 3 行
    assert "分取消" in str(ws.cell(row=2, column=4).value)
    # 标题取表格前面那个加粗行，而不是旧卡片的「待澄清问题」
    assert "澄清问题清单" in out["已生成"]


async def test_a_ui_table_listed_after_the_answer_still_wins(tmp_path, monkeypatch):
    """反过来也要对：正文里写过表，之后又用 ui.table 列了 192 条，
    「这个表」指的是后出现的那张。"""
    from ontocopilot import server
    from ontocopilot.kernel.memory.dialogue import Speaker

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="mdx2")
    s.dir.mkdir(parents=True, exist_ok=True)
    server._publish_turn(s, Speaker.ASSISTANT,
                         "| 甲 | 乙 |\n| :--- | :--- |\n| 1 | 2 |\n")
    s.state["oir"] = {"rules": [
        {"rid": f"r{i}", "statement": {"value": f"规则 {i}"},
         "ruleKind": {"value": "c"}, "actor": {"value": ""}} for i in range(5)]}

    reg = server._converse_tools(s)
    await reg.call("ui.table", {"kind": "rules"}, _Ctx(), scope="converse")
    out = await reg.call("export.file", {"format": "csv"}, _Ctx(), scope="converse")
    assert out["表格行数"] == 5
    body = (s.dir / "exports" / out["已生成"]).read_bytes().decode("utf-8-sig")
    assert "规则 3" in body and "甲" not in body


async def test_exporting_a_table_the_user_named_not_just_the_last_one(tmp_path, monkeypatch):
    """他截图里的场景：先要了一张《AI 招聘业务流程梳理及访谈提问框架》，聊了两轮
    别的，然后说"把这个列表导出成 Excel：AI 招聘业务流程梳理及访谈提问框架"。

    导出取的是**最后一张**表 —— 也就是后面那轮的主数据/事务数据对比表。他点了名，
    代码没听，下载下来是另一张表。
    """
    import openpyxl

    from ontocopilot import server
    from ontocopilot.kernel.memory.dialogue import Speaker

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="named1")
    s.dir.mkdir(parents=True, exist_ok=True)

    server._publish_turn(s, Speaker.USER, "AI 招聘流程该怎么梳理？")
    server._publish_turn(s, Speaker.ASSISTANT, """**AI 招聘业务流程梳理及访谈提问框架**

| 阶段 / 步骤 | 梳理重点 | 核心业务访谈提问 |
| :--- | :--- | :--- |
| 1. 岗位画像与JD定义 | 岗位实体 | 招聘需求提出时怎么写 JD？ |
| 2. 简历解析与人岗匹配 | 候选人/简历实体 | AI 解析出哪些核心字段？ |
""")
    server._publish_turn(s, Speaker.USER, "主数据和事务数据怎么划分？")
    server._publish_turn(s, Speaker.ASSISTANT, """**主数据与事务数据判定**

| 判定维度 | 主数据 | 事务数据 |
| :--- | :--- | :--- |
| 生命周期 | 独立于流程 | 强依赖流程 |
| 变化频率 | 低频更新 | 高频追加 |
""")

    reg = server._converse_tools(s)
    out = await reg.call("export.file",
                         {"format": "excel", "name": "AI 招聘业务流程梳理及访谈提问框架"},
                         _Ctx(), scope="converse")
    assert "error" not in out, out
    ws = openpyxl.load_workbook(s.dir / "exports" / out["已生成"]).worksheets[0]
    assert [c.value for c in ws[1]] == ["阶段 / 步骤", "梳理重点", "核心业务访谈提问"]
    assert "岗位画像" in str(ws.cell(row=2, column=1).value)
    assert "判定维度" not in [c.value for c in ws[1]], "导的还是最后那张对比表"

    # 不给名字仍然是"最后一张"
    last = await reg.call("export.file", {"format": "csv"}, _Ctx(), scope="converse")
    body = (s.dir / "exports" / last["已生成"]).read_bytes().decode("utf-8-sig")
    assert "判定维度" in body


async def test_naming_a_table_that_does_not_exist_lists_what_does(tmp_path, monkeypatch):
    """点了名却找不到时，**不能默默导另一张给他** —— 那正是他遇到的事。"""
    from ontocopilot import server
    from ontocopilot.kernel.memory.dialogue import Speaker

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="named2")
    s.dir.mkdir(parents=True, exist_ok=True)
    server._publish_turn(s, Speaker.ASSISTANT,
                         "**采购问题清单**\n\n| a | b |\n| :--- | :--- |\n| 1 | 2 |\n")

    reg = server._converse_tools(s)
    out = await reg.call("export.file", {"format": "md", "name": "根本没有这张表"},
                         _Ctx(), scope="converse")
    assert "error" in out
    assert "采购问题清单" in str(out["现有的表"])


async def test_a_table_survives_compaction_and_is_still_exportable(tmp_path, monkeypatch):
    """DialogueMemory 会把最老的几轮压成一句摘要。用户过几轮回头要那张表时，
    只看 DialogueMemory 就找不到了 —— 得从耐久事件表里翻。"""
    from ontocopilot import server
    from ontocopilot.kernel.memory.dialogue import Speaker

    monkeypatch.setattr(server, "ROOT", tmp_path)
    s = server.Session(id="named3")
    s.dir.mkdir(parents=True, exist_ok=True)

    server._publish_turn(s, Speaker.ASSISTANT,
                         "**访谈提纲**\n\n| 阶段 | 提问 |\n| :--- | :--- |\n| 岗位 | 怎么写 JD？ |\n")
    # 灌到超预算，把那一轮压掉
    for i in range(30):
        server._publish_turn(s, Speaker.USER, f"第 {i} 个问题，" + "细节" * 80)
        server._publish_turn(s, Speaker.ASSISTANT, f"第 {i} 个回答，" + "结论" * 80)

    dm = server._dialogue(s)
    assert any(str(t.speaker) == "system" for t in dm.turns), "没触发压缩，测试无效"
    assert not any("访谈提纲" in (t.text or "") and "|" in (t.text or "")
                   for t in dm.turns), "那一轮还没被压掉，测试无效"

    tables = await server._conversation_tables(s)
    assert any("访谈提纲" in str(t.get("title")) for t in tables), \
        "压缩之后就再也找不到那张表了"
