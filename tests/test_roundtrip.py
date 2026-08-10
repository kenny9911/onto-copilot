"""往返闭环：模板编译 → 写 xlsx → 业务方填写 → 回传审核。

这里验的是 ADR-4（隐藏锚点列）到底成不成立 —— 业务方把行打乱、插行、删行之后，
系统还能不能精确对齐并算出可信的完成度。
"""

from __future__ import annotations

import random

import pytest
from openpyxl import load_workbook

from ontocopilot.onto.audit import ReturnAuditor, merge_into_oir, read_returned
from ontocopilot.onto.conflict import ConflictKind, detect_all
from ontocopilot.onto.oir import Origin, Status
from ontocopilot.onto.template import (
    ANCHOR_HASH,
    ANCHOR_RID,
    Role,
    compile_template,
    write_xlsx,
)

pytest_plugins = []


@pytest.fixture
def oir():
    from test_onto import _ddl, _xlsx  # 复用「计划金额」双口径场景

    import test_onto

    return test_onto.oir.__wrapped__() if hasattr(test_onto.oir, "__wrapped__") else _build()


def _build():
    """独立构造，避免依赖另一个模块的 fixture 内部结构。"""
    from ontocopilot.onto.oir import (
        OIR,
        ActionType,
        BaseType,
        Cardinality,
        LinkType,
        ObjectType,
        PropertyType,
        Provenance,
        extracted,
        inferred,
        make_rid,
    )

    def px(row, snip):
        return Provenance("f3", "实体梳理.xlsx",
                          {"kind": "cell", "sheet": "业务对象实体梳理", "row": row, "col": "F"},
                          snippet=snip, extractor="docling", confidence=0.94)

    def pd(obj, snip):
        return Provenance("f5", "schema.ddl", {"kind": "ddl", "object": obj},
                          snippet=snip, extractor="sqlglot", confidence=0.9)

    o = OIR()
    plan = o.add_object(ObjectType(
        rid=make_rid("ot", "purchase_plan_header"),
        api_name=extracted("purchasePlanHeader", pd("pbp_header", "CREATE TABLE pbp_header")),
        display_name=extracted("采购业务计划头", px(2, "采购业务计划头")),
        primary_key=extracted(["pt_plan_id"], pd("pbp_header", "PRIMARY KEY")),
        aliases=["pbpHeader", "采购需求计划"], owner="王明"))
    contract = o.add_object(ObjectType(
        rid=make_rid("ot", "clm_contract"),
        api_name=extracted("clmContract", pd("clm_contract", "CREATE TABLE clm_contract")),
        display_name=extracted("采购合同", px(44, "采购合同")),
        primary_key=extracted(["pt_contract_id"], pd("clm_contract", "PRIMARY KEY")),
        owner="李强"))

    for rid, parent, api, disp, owner in [
        ("pt_plan_id", plan.rid, "planId", "计划编号", "王明"),
        ("pt_contract_id", contract.rid, "contractId", "合同编号", "李强"),
    ]:
        o.add_property(PropertyType(
            rid=rid, parent=parent, api_name=extracted(api, px(2, api)),
            display_name=extracted(disp, px(2, disp)),
            base_type=extracted(BaseType.STRING, px(2, "STRING")),
            definition=extracted("主键", px(2, "主键")), owner=owner))

    o.add_property(PropertyType(
        rid="pt_amount_budget", parent=plan.rid,
        api_name=extracted("planAmount", px(44, "planAmount")),
        display_name=extracted("计划金额", px(44, "计划金额")),
        base_type=extracted(BaseType.DECIMAL, px(44, "DECIMAL(18,2)")),
        definition=extracted("含税，年度累计，CNY", px(44, "计划金额（含税，年度累计）")),
        owner="王明"))
    o.add_property(PropertyType(
        rid="pt_amount_contract", parent=contract.rid,
        api_name=extracted("planAmount", pd("clm_contract", "plan_amount")),
        display_name=extracted("计划金额", pd("clm_contract", "plan_amount")),
        base_type=extracted(BaseType.DECIMAL, pd("clm_contract", "DECIMAL(18,2)")),
        definition=extracted("不含税，单次，CNY", pd("clm_contract", "-- 不含税·单次")),
        owner="李强"))

    ev = pd("clm_contract", "FOREIGN KEY (plan_id)")
    o.add_link(LinkType(rid="lt_plan_contract", api_name=extracted("planContracts", ev),
                        source=plan.rid, target=contract.rid,
                        cardinality=extracted(Cardinality.ONE_TO_MANY, ev),
                        join_key=extracted({"fromProp": "pt_plan_id",
                                            "toProp": "pt_contract_id"}, ev)))
    o.add_action(ActionType(
        rid="at_submit_plan",
        api_name=inferred("submitPurchasePlan"), applies_to=[plan.rid],
        source_endpoint=inferred({"method": "post", "path": "/purchase-plans/{id}/submit"}),
        status=Status.DRAFT_FROM_API))
    return o


@pytest.fixture
def built():
    return _build()


# ══════════════════════════════════════════════════════════════════
#  编译
# ══════════════════════════════════════════════════════════════════
def test_definition_is_always_business_required(built):
    """口径永远是业务必填 —— 系统抽出来的只是候选，必须有人认领。"""
    spec = compile_template(built)
    defs = [c for c in spec.cells() if c.field == "definition"]
    assert defs and all(c.role is Role.REQUIRED for c in defs)


def test_conflicted_cells_are_flagged_with_the_reason(built):
    conflicts = detect_all(built)
    spec = compile_template(built, conflicts)
    flagged = [c for c in spec.cells() if c.conflict]
    assert flagged
    amount = next(c for c in flagged if c.rid == "pt_amount_budget")
    assert "口径不一致" in amount.comment


def test_prefilled_cells_carry_their_evidence(built):
    """业务方对某个预填值有疑问时，能直接从批注看到系统凭什么这么填。"""
    spec = compile_template(built)
    api = next(c for c in spec.cells() if c.rid == "pt_amount_budget" and c.field == "apiName")
    assert "实体梳理.xlsx!业务对象实体梳理!R44CF" in api.comment


def test_stats_report_prefill_rate_and_owners(built):
    s = compile_template(built).stats()
    assert 0 < s["prefill_rate"] < 1
    assert s["business_required"] > 0
    assert set(s["owners"]) == {"王明", "李强"}


# ══════════════════════════════════════════════════════════════════
#  写 xlsx
# ══════════════════════════════════════════════════════════════════
def test_xlsx_has_hidden_anchor_columns_and_frozen_guide(built, tmp_path):
    spec = compile_template(built, detect_all(built))
    path = write_xlsx(spec, tmp_path / "t.xlsx", project="ONT-112")
    wb = load_workbook(path)

    assert wb.sheetnames[0] == "00_填写指引"
    ws = wb["02_属性明细"]
    assert ws.cell(row=2, column=1).value == ANCHOR_RID
    assert ws.cell(row=2, column=2).value == ANCHOR_HASH
    assert ws.column_dimensions["A"].hidden and ws.column_dimensions["B"].hidden
    assert ws.freeze_panes == "C3"
    assert ws.cell(row=1, column=1).value.startswith("口径定义是本轮的重点")


def test_xlsx_marks_required_cells_yellow_and_adds_dropdowns(built, tmp_path):
    path = write_xlsx(compile_template(built), tmp_path / "t.xlsx")
    ws = load_workbook(path)["02_属性明细"]
    headers = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    col = headers.index("definition") + 1
    assert ws.cell(row=3, column=col).fill.fgColor.rgb.endswith("FDF6E3")  # 黄
    assert ws.cell(row=3, column=col).comment is not None
    assert any(dv.type == "list" for dv in ws.data_validations.dataValidation)


# ══════════════════════════════════════════════════════════════════
#  往返对齐 —— ADR-4 的核心验证
# ══════════════════════════════════════════════════════════════════
def _fill(path, answers: dict[tuple[str, str], str], *, shuffle=False, drop=None, insert=0):
    """模拟业务方填写：填内容、打乱行序、删行、插入新行。"""
    wb = load_workbook(path)
    for ws in wb.worksheets:
        if ws.title.startswith("00_") or ws.cell(row=2, column=1).value != ANCHOR_RID:
            continue
        headers = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
        rows = []
        for r in range(3, ws.max_row + 1):
            rid = ws.cell(row=r, column=1).value
            if not rid:
                continue
            if drop and rid in drop:
                continue
            vals = [ws.cell(row=r, column=c).value for c in range(1, len(headers) + 1)]
            for c, name in enumerate(headers):
                if (str(rid), str(name)) in answers:
                    vals[c] = answers[(str(rid), str(name))]
            rows.append(vals)
        if shuffle:
            random.Random(7).shuffle(rows)
        for _ in range(insert):
            rows.insert(1, [None] * len(headers))  # 业务方顺手插的空行
        for r in range(3, ws.max_row + 1):
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).value = None
        for i, vals in enumerate(rows):
            for c, v in enumerate(vals, start=1):
                ws.cell(row=3 + i, column=c, value=v)
    wb.save(path)
    return path


def test_alignment_survives_reorder_insert_and_delete(built, tmp_path):
    """行被打乱、插空行、删行之后，锚点仍然精确对齐。"""
    spec = compile_template(built)
    path = write_xlsx(spec, tmp_path / "t.xlsx")

    answers = {
        ("pt_amount_budget", "definition"): "含税，年度累计，CNY，口径由预算科维护",
        ("pt_amount_contract", "definition"): "不含税，单次，CNY",
    }
    _fill(path, answers, shuffle=True, insert=2, drop={"pt_plan_id"})

    returned = read_returned(path)
    assert returned[("pt_amount_budget", "definition")].startswith("含税")
    assert returned[("pt_amount_contract", "definition")] == "不含税，单次，CNY"

    result = ReturnAuditor().audit(spec, returned, oir=built)
    assert "pt_plan_id" in result.unmatched_rows, "被删的行要报出来，不能静默丢失"
    assert result.new_rows == []


def test_unchanged_prefill_is_caught_and_does_not_count_as_filled(built, tmp_path):
    """AI 预填被原样交回 —— 这格没被审过，不能算完成。"""
    spec = compile_template(built)
    path = write_xlsx(spec, tmp_path / "t.xlsx")
    prefilled = next(c for c in spec.cells()
                     if c.rid == "pt_amount_budget" and c.field == "definition")
    _fill(path, {("pt_amount_budget", "definition"): prefilled.value})

    result = ReturnAuditor().audit(spec, read_returned(path), oir=built)
    perf = [c for c in result.findings if c.kind is ConflictKind.PERFUNCTORY]
    assert any("原样交回" in c.summary for c in perf)
    d = next(x for x in result.diffs
             if x.rid == "pt_amount_budget" and x.field == "definition")
    assert d.untouched_prefill and not d.changed


def test_completeness_is_weighted_and_excludes_perfunctory(built, tmp_path):
    spec = compile_template(built)
    path = write_xlsx(spec, tmp_path / "t.xlsx")

    empty = ReturnAuditor().audit(spec, read_returned(path), oir=built)
    assert empty.completeness < 0.5

    # 自由文本列每行给不同内容 —— 全填一样的正是 BULK_FILLED 想抓的敷衍
    answers = {}
    for i, c in enumerate(spec.cells()):
        if c.role is not Role.REQUIRED:
            continue
        answers[(c.rid, c.field)] = {
            "definition": f"不含税、单次结算、CNY，由财务共享中心维护税率表（{c.rid}）",
            "description": f"{c.rid} 在采购业务中承载的实体，口径见属性明细表",
            "owner": "王明" if i % 2 else "李强", "cardinality": "MANY_TO_MANY",
            "primaryKey": "planId", "effects": "CREATE", "confirmed": "已确认",
            "correct": "正确",
        }.get(c.field, "已确认")
    _fill(path, answers)

    full = ReturnAuditor().audit(spec, read_returned(path), oir=built)
    assert full.completeness > empty.completeness
    assert full.completeness >= 0.95, f"未达标的项: {full.counts}"


def test_enum_columns_are_not_flagged_as_bulk_filled(built, tmp_path):
    """枚举列和责任人列本来取值就少 —— 按重复率判会把正常填写全部误报成敷衍。"""
    spec = compile_template(built)
    path = write_xlsx(spec, tmp_path / "t.xlsx")
    _fill(path, {(c.rid, c.field): "已确认"
                 for c in spec.cells() if c.field == "confirmed"})

    result = ReturnAuditor().audit(spec, read_returned(path), oir=built)
    perf = [c for c in result.findings if c.kind is ConflictKind.PERFUNCTORY]
    assert not any("confirmed" in c.summary for c in perf)


def test_enum_violation_is_caught_by_rule_not_model(built, tmp_path):
    spec = compile_template(built)
    path = write_xlsx(spec, tmp_path / "t.xlsx")
    _fill(path, {("lt_plan_contract", "cardinality"): "一对多"})  # 中文，不在枚举里

    result = ReturnAuditor().audit(spec, read_returned(path), oir=built)
    bad = [c for c in result.findings if c.kind is ConflictKind.TYPE_MISMATCH]
    assert bad and "不在允许集合" in bad[0].summary
    assert bad[0].detector == "rule:ENUM-01"


def test_naming_violation_from_return_is_auto_repaired(built, tmp_path):
    spec = compile_template(built)
    path = write_xlsx(spec, tmp_path / "t.xlsx")
    _fill(path, {("pt_amount_budget", "apiName"): "plan_amount_incl_tax"})

    result = ReturnAuditor().audit(spec, read_returned(path), oir=built)
    assert result.auto_repaired
    entry = result.auto_repaired[0]
    assert entry["to"] == "planAmountInclTax" and entry["reversible"]
    assert built.properties["pt_amount_budget"].api_name.origin is Origin.AUTO_REPAIRED
    # 自动修掉的不该再出现在打回单里 —— 那是给人做的事，系统已经做完了
    assert not any(c.kind is ConflictKind.NAMING_VIOLATION for c in result.findings)


def test_return_slips_are_grouped_by_owner(built, tmp_path):
    """打回给"团队"等于打回给没有人。"""
    spec = compile_template(built)
    path = write_xlsx(spec, tmp_path / "t.xlsx")
    result = ReturnAuditor().audit(spec, read_returned(path), oir=built)

    owners = {s.owner for s in result.slips}
    assert {"王明", "李强"} <= owners
    assert all(s.items for s in result.slips)
    assert result.slips == sorted(result.slips, key=lambda s: -len(s.items))


def test_post_return_divergence_is_detected(built, tmp_path):
    """两个人各填各的，轴上又冲突了 —— 第二轮还得对齐。"""
    spec = compile_template(built)
    path = write_xlsx(spec, tmp_path / "t.xlsx")
    _fill(path, {
        ("pt_amount_budget", "definition"): "含税，年度累计",
        ("pt_amount_contract", "definition"): "不含税，单次",
    })
    result = ReturnAuditor().audit(spec, read_returned(path), oir=built)
    div = [c for c in result.findings if c.kind is ConflictKind.SEMANTIC_DIVERGENCE]
    assert div and "需 " in div[0].summary and "对齐" in div[0].summary


def test_merge_writes_back_only_what_humans_actually_filled(built, tmp_path):
    spec = compile_template(built)
    path = write_xlsx(spec, tmp_path / "t.xlsx")
    prefilled = next(c for c in spec.cells()
                     if c.rid == "pt_amount_contract" and c.field == "definition")
    _fill(path, {
        ("pt_amount_budget", "definition"): "含税，年度累计，由预算科认领",
        ("pt_amount_contract", "definition"): prefilled.value,  # 原样交回
    })
    result = ReturnAuditor().audit(spec, read_returned(path), oir=built)
    changed, _dropped = merge_into_oir(built, result.diffs)

    assert "pt_amount_budget.definition" in changed
    assert "pt_amount_contract.definition" not in changed, "原样交回不算填写"
    assert built.properties["pt_amount_budget"].definition.origin is Origin.USER
    assert built.properties["pt_amount_contract"].definition.origin is Origin.EXTRACTED


# ══════════════════════════════════════════════════════════════════
#  回传件的结构损伤
# ══════════════════════════════════════════════════════════════════
# 业务方最常做的三个动作 —— 顶上插一行标题、最左边插一列做批注、删掉看不懂的
# 隐藏列 —— 以前各自会让整张表静默消失：不抛异常、不告警，只有 FDE 自己去数
# unmatched_rows 才能发现。**一整张表的人白填了，而他们收不到任何提示。**
import shutil as _shutil  # noqa: E402

from openpyxl import load_workbook as _load  # noqa: E402


@pytest.fixture
def spec_fixture(built):
    """一份编译好的模板规格。结构损伤用例共用它。"""
    return compile_template(built)


def _mutated(tmp_path, spec, mutate, name="m"):
    """写一份模板、按 mutate 改动、再读回来。"""
    from ontocopilot.onto.audit import read_returned
    from ontocopilot.onto.template import write_xlsx

    src = tmp_path / "tpl.xlsx"
    if not src.exists():
        write_xlsx(spec, src, project="T")
    dst = tmp_path / f"{name}.xlsx"
    _shutil.copy(src, dst)
    if mutate is not None:
        wb = _load(dst)
        mutate(wb)
        wb.save(dst)
    return read_returned(dst)


def _first_sheet(spec):
    return next(sh.name for sh in spec.sheets if sh.rows)


def test_untouched_return_reports_no_damage(tmp_path, spec_fixture):
    """原样交回来必须零损伤 —— 假阳性会让真损伤被忽略。"""
    from ontocopilot.onto.audit import ReturnAuditor

    r = ReturnAuditor().audit(spec_fixture, _mutated(tmp_path, spec_fixture, None))
    assert r.readable, r.damage


def test_inserted_top_row_is_still_read(tmp_path, spec_fixture):
    """插一行标题以前会让整表消失。现在按实际位置读到，并告知结构变了。"""
    from ontocopilot.onto.audit import ReturnAuditor

    name = _first_sheet(spec_fixture)
    got = _mutated(tmp_path, spec_fixture,
                   lambda wb: wb[name].insert_rows(1), "toprow")
    r = ReturnAuditor().audit(spec_fixture, got)
    assert not r.unmatched_rows, f"插一行导致 {len(r.unmatched_rows)} 行读不到"
    assert not r.readable and any("表头在第 3 行" in d for d in r.damage), r.damage


def test_inserted_left_column_is_still_read(tmp_path, spec_fixture):
    from ontocopilot.onto.audit import ReturnAuditor

    name = _first_sheet(spec_fixture)
    got = _mutated(tmp_path, spec_fixture,
                   lambda wb: wb[name].insert_cols(1), "leftcol")
    r = ReturnAuditor().audit(spec_fixture, got)
    assert not r.unmatched_rows, f"插一列导致 {len(r.unmatched_rows)} 行读不到"
    assert not r.readable


def test_deleted_anchor_columns_are_reported_not_swallowed(tmp_path, spec_fixture):
    """锚点删了就是真的读不了 —— 但必须**说出来**，而不是静默返回空。"""
    from ontocopilot.onto.audit import ReturnAuditor

    name = _first_sheet(spec_fixture)
    got = _mutated(tmp_path, spec_fixture,
                   lambda wb: wb[name].delete_cols(1, 2), "noanchor")
    r = ReturnAuditor().audit(spec_fixture, got)
    assert not r.readable
    assert any("找不到锚点列" in d for d in r.damage), r.damage
    assert r.unmatched_rows, "整表读不到却没有任何一行被标为未匹配"


def test_duplicate_rid_is_reported(tmp_path, spec_fixture):
    """复制粘贴出来的重复行会后写覆盖先写，静默替换掉原来的答案。"""
    from ontocopilot.onto.audit import ReturnAuditor

    name = _first_sheet(spec_fixture)

    def dup(wb):
        ws = wb[name]
        ws.append([ws.cell(row=3, column=c).value
                   for c in range(1, ws.max_column + 1)])

    r = ReturnAuditor().audit(spec_fixture, _mutated(tmp_path, spec_fixture, dup, "dup"))
    assert any("与前面重复" in d for d in r.damage), r.damage


def test_a_renamed_sheet_is_still_read(tmp_path, spec_fixture):
    """以前靠 00_ 前缀判断该不该读 —— 把数据表改名成 00_xxx 就能让它静默消失。"""
    from ontocopilot.onto.audit import ReturnAuditor

    name = _first_sheet(spec_fixture)

    def rename(wb):
        wb[name].title = "00_我自己改的名字"

    r = ReturnAuditor().audit(spec_fixture,
                              _mutated(tmp_path, spec_fixture, rename, "renamed"))
    assert not r.unmatched_rows, f"改名导致 {len(r.unmatched_rows)} 行读不到"


def test_structural_damage_becomes_a_finding(tmp_path, spec_fixture):
    """损伤要和别的问题一样进 findings —— 只放在 summary 的一个列表里会被淹没，
    而它的后果比任何一条填错都严重。"""
    from ontocopilot.onto.audit import ReturnAuditor

    name = _first_sheet(spec_fixture)
    got = _mutated(tmp_path, spec_fixture,
                   lambda wb: wb[name].delete_cols(1, 2), "finding")
    r = ReturnAuditor().audit(spec_fixture, got)
    assert any(c.rid == "cf_return_damaged" for c in r.findings), \
        [c.rid for c in r.findings]
    assert r.summary()["readable"] is False


# ══════════════════════════════════════════════════════════════════
#  回写黑洞
# ══════════════════════════════════════════════════════════════════
# 实测：627 格被判为真正填写，最后只有 172 条落回 OIR。primaryKey / owner /
# effects 共 455 格被读进来、被算进完成度、然后在 `case _: continue` 里静默丢弃。
# 填表的人以为答案生效了，FDE 看到完成度上升，而模型里什么都没变。
def test_answers_to_open_questions_are_written_back():
    """一个待澄清问题变成一条事实 —— 整轮往返里最有价值的那次写入。"""
    from ontocopilot.onto.audit import CellDiff, merge_into_oir
    from ontocopilot.onto.oir import OIR, OpenQuestion, Status, inferred
    from ontocopilot.onto.template import Role

    oir = OIR()
    oir.add_question(OpenQuestion(rid="oq_1", text=inferred("集采还是普采？")))
    changed, dropped = merge_into_oir(oir, [
        CellDiff(rid="oq_1", sheet="02_待澄清问题", field="答复",
                 before="", after="按金额分档，50 万以上走集采", role=Role.REQUIRED)])
    assert changed == ["oq_1.答复"] and not dropped
    q = oir.questions["oq_1"]
    assert q.answered and q.status is Status.CONFIRMED
    assert str(q.answer.origin) == "user"


def test_rule_confirmation_is_written_back():
    from ontocopilot.onto.audit import CellDiff, merge_into_oir
    from ontocopilot.onto.oir import BusinessRule, OIR, Status, inferred
    from ontocopilot.onto.template import Role

    oir = OIR()
    for rid in ("br_1", "br_2"):
        oir.add_rule(BusinessRule(rid=rid, statement=inferred("某条规则")))
    merge_into_oir(oir, [
        CellDiff(rid="br_1", sheet="03_业务规则确认", field="这条对吗",
                 before="", after="对", role=Role.REQUIRED),
        CellDiff(rid="br_2", sheet="03_业务规则确认", field="这条对吗",
                 before="", after="不对", role=Role.REQUIRED)])
    assert oir.rules["br_1"].status is Status.CONFIRMED
    assert oir.rules["br_2"].status is Status.REJECTED


def test_rule_host_is_resolved_by_display_name():
    """业务方填的是中文单据名，OIR 里是 rid —— 对不上就必须报，不能默默丢。"""
    from ontocopilot.onto.audit import CellDiff, merge_into_oir
    from ontocopilot.onto.oir import BusinessRule, OIR, ObjectType, inferred
    from ontocopilot.onto.template import Role

    oir = OIR()
    oir.add_object(ObjectType(rid="ot_a", api_name=inferred("pbpHeader"),
                              display_name=inferred("采购需求计划"),
                              primary_key=inferred([])))
    oir.add_rule(BusinessRule(rid="br_1", statement=inferred("某条规则")))
    oir.add_rule(BusinessRule(rid="br_2", statement=inferred("另一条")))
    changed, dropped = merge_into_oir(oir, [
        CellDiff(rid="br_1", sheet="03_业务规则确认", field="管哪个单据",
                 before="", after="采购需求计划", role=Role.REQUIRED),
        CellDiff(rid="br_2", sheet="03_业务规则确认", field="管哪个单据",
                 before="", after="根本不存在的单据", role=Role.REQUIRED)])
    assert oir.rules["br_1"].applies_to == ["ot_a"]
    assert "br_1.管哪个单据" in changed
    assert any("对不上任何对象" in d for d in dropped), dropped
    assert oir.rules["br_2"].applies_to == []


def test_dropped_writes_are_reported_not_swallowed():
    """没有回写路径的列必须出现在 dropped 里。静默丢弃是这条链最危险的失败 ——
    完成度照样上升，没人会发现。"""
    from ontocopilot.onto.audit import CellDiff, merge_into_oir
    from ontocopilot.onto.oir import OIR, ObjectType, inferred
    from ontocopilot.onto.template import Role

    oir = OIR()
    oir.add_object(ObjectType(rid="ot_a", api_name=inferred("x"),
                              display_name=inferred("X"), primary_key=inferred([])))
    changed, dropped = merge_into_oir(oir, [
        CellDiff(rid="ot_a", sheet="01_对象清单", field="某个没人接的列",
                 before="", after="填了很认真的内容", role=Role.REQUIRED),
        CellDiff(rid="ot_不存在", sheet="01_对象清单", field="description",
                 before="", after="也很认真", role=Role.REQUIRED)])
    assert not changed
    assert len(dropped) == 2
    assert any("没有回写路径" in d for d in dropped)
    assert any("找不到这个条目" in d for d in dropped)


def test_previously_dropped_columns_now_land():
    """primaryKey / owner / effects —— 实测被丢掉的那 455 格。"""
    from ontocopilot.onto.audit import CellDiff, merge_into_oir
    from ontocopilot.onto.oir import ActionType, OIR, ObjectType, inferred
    from ontocopilot.onto.template import Role

    oir = OIR()
    oir.add_object(ObjectType(rid="ot_a", api_name=inferred("x"),
                              display_name=inferred("X"), primary_key=inferred([])))
    oir.add_action(ActionType(rid="at_a", api_name=inferred("createX")))
    changed, dropped = merge_into_oir(oir, [
        CellDiff(rid="ot_a", sheet="s", field="primaryKey", before="",
                 after="计划编号、版本号", role=Role.REQUIRED),
        CellDiff(rid="ot_a", sheet="s", field="owner", before="",
                 after="采购计划员", role=Role.REQUIRED),
        CellDiff(rid="at_a", sheet="s", field="effects", before="",
                 after="创建单据、写入台账", role=Role.REQUIRED)])
    assert len(changed) == 3 and not dropped
    assert oir.objects["ot_a"].primary_key.value == ["计划编号", "版本号"]
    assert oir.objects["ot_a"].owner == "采购计划员"
    assert oir.actions["at_a"].effects.value == ["创建单据", "写入台账"]
