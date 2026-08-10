"""回传审核器 —— 业务方填完交回来之后的这一关。

流程：**锚点对齐 → 单元格级 diff → 规则审核 → 语义审核 → 自动修 → 按责任人打回**。

锚点是全部精度的来源。业务方会增删行、重排序、改列宽、复制粘贴 —— 按行号对齐
必然错位，按内容模糊匹配会把改过名的行判成新增。隐藏的 ``_oir_rid`` 让这一切
都不成问题，也让「这格有没有被动过」变成一次哈希比对。

分工照旧（架构文档 ADR-5）：必填缺失、枚举越界、命名规范走确定性规则，秒级、
零成本、零方差；只有口径矛盾和疑似敷衍交模型，且先用启发式收窄候选。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .conflict import (
    Conflict,
    ConflictKind,
    Handling,
    _cid,
    axis_diff,
    perfunctory_signals,
    to_camel,
)
from .oir import OIR, Origin, Status, by_user
from .template import ANCHOR_HASH, ANCHOR_RID, WEIGHTS, Cell, Role, TemplateSpec


@dataclass(slots=True)
class CellDiff:
    """一格的前后对比。"""

    rid: str
    sheet: str
    field: str
    before: str
    after: str
    role: Role
    owner: str | None = None

    @property
    def changed(self) -> bool:
        return self.before.strip() != self.after.strip()

    @property
    def filled(self) -> bool:
        return bool(self.after.strip())

    @property
    def untouched_prefill(self) -> bool:
        """AI 预填了、业务方原样交回 —— 说明这格他没审。"""
        return bool(self.before.strip()) and not self.changed


@dataclass(slots=True)
class ReturnSlip:
    """一份打回单。按责任人分组 —— 打回给"团队"等于打回给没有人。"""

    owner: str
    items: list[Conflict] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        by_kind: dict[str, list[str]] = {}
        for c in self.items:
            by_kind.setdefault(str(c.kind), []).append(c.summary)
        return {"owner": self.owner, "count": len(self.items), "by_kind": by_kind}


@dataclass(slots=True)
class AuditResult:
    completeness: float
    diffs: list[CellDiff] = field(default_factory=list)
    findings: list[Conflict] = field(default_factory=list)
    auto_repaired: list[dict[str, Any]] = field(default_factory=list)
    slips: list[ReturnSlip] = field(default_factory=list)
    unmatched_rows: list[str] = field(default_factory=list)
    new_rows: list[str] = field(default_factory=list)
    #: 回传件结构被破坏的地方。**空列表以外的一切都要拦住流程** ——
    #: 一张读不了的表意味着那些人的工作白做了，而他们不知道。
    damage: list[str] = field(default_factory=list)

    @property
    def readable(self) -> bool:
        return not self.damage

    @property
    def counts(self) -> dict[str, int]:
        out: dict[str, int] = {}
        for c in self.findings:
            out[str(c.kind)] = out.get(str(c.kind), 0) + 1
        return out

    def summary(self) -> dict[str, Any]:
        return {
            "completeness": round(self.completeness, 4),
            "cells_compared": len(self.diffs),
            "cells_changed": sum(1 for d in self.diffs if d.changed),
            "findings": self.counts,
            "auto_repaired": len(self.auto_repaired),
            "slips": [s.to_dict() for s in self.slips],
            "unmatched_rows": self.unmatched_rows,
            "new_rows": self.new_rows,
            "damage": self.damage,
            "readable": self.readable,
        }


# ══════════════════════════════════════════════════════════════════
#  读回传件
# ══════════════════════════════════════════════════════════════════
def read_returned(path: Path | str) -> dict[tuple[str, str], str]:
    """读回传的 xlsx，按 ``(rid, 字段)`` 索引取值。

    行序、行数、列宽的任何变化都不影响结果 —— 全靠隐藏的 ``_oir_rid`` 对齐。
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    out: dict[tuple[str, str], str] = {}
    damage: list[str] = []

    for ws in wb.worksheets:
        # 重复检测必须**按表**做：同一个对象本来就会同时出现在对象清单和术语表里，
        # 跨表去重会把正常结构报成损伤。
        seen_rids: set[str] = set()
        loc = _locate_anchor(ws)
        if loc is None:
            # 这张表读不了。以前这里是 `continue` —— 静默丢掉整张 172 行的表，
            # 不抛异常、不告警，FDE 只能靠自己去数 unmatched_rows 才发现。
            # **读不了必须说出来**：业务方在顶上插一行标题、在最左边插一列做批注，
            # 都是极常见的动作，而它们的代价是一整张表的工作白做。
            if _looks_like_data(ws):
                damage.append(
                    f"「{ws.title}」找不到锚点列 {ANCHOR_RID} —— 这张表的 "
                    f"{max(0, ws.max_row - 2)} 行没有被读取。"
                    f"通常是因为顶部插了行、最左边插了列，或者隐藏列被删掉了")
            continue
        hrow, rcol, headers = loc
        if hrow != 2 or rcol != 1:
            # 能读，但结构被动过 —— 也要说，因为下次回传可能就读不了了
            damage.append(f"「{ws.title}」表头在第 {hrow} 行、锚点在第 {rcol} 列"
                          f"（原本是第 2 行第 1 列），已按实际位置读取")

        for r in range(hrow + 1, ws.max_row + 1):
            rid = ws.cell(row=r, column=rcol).value
            if not rid:
                continue
            key = str(rid)
            if key in seen_rids:
                # 复制粘贴出来的重复行会后写覆盖先写，静默替换掉原来的答案
                damage.append(f"「{ws.title}」第 {r} 行的 {key} 与前面重复，后者已覆盖前者")
            seen_rids.add(key)
            for c, name in headers:
                v = ws.cell(row=r, column=c).value
                out[(key, str(name))] = "" if v is None else str(v).strip()

    if damage:
        out[("__damage__", "__damage__")] = "\n".join(damage)
    return out


#: 扫描表头的深度。业务方在顶上加标题、加说明，通常不超过这个行数。
_ANCHOR_SCAN_ROWS = 8
_ANCHOR_SCAN_COLS = 6


def _locate_anchor(ws: Any) -> tuple[int, int, list[tuple[int, str]]] | None:
    """找出表头在第几行、锚点在第几列，以及每个数据列的位置。

    以前这三件事全是硬编码（第 2 行、第 1 列、``00_`` 前缀）。硬编码的代价不是
    读错，是**读不到而且不说** —— 三种最常见的用户动作（顶部插行、左侧插列、
    删掉看不懂的隐藏列）各自会让整张表凭空消失。

    返回 ``None`` 表示这张表里没有锚点。
    """
    for row in range(1, min(_ANCHOR_SCAN_ROWS, ws.max_row) + 1):
        for col in range(1, min(_ANCHOR_SCAN_COLS, ws.max_column) + 1):
            if str(ws.cell(row=row, column=col).value or "").strip() != ANCHOR_RID:
                continue
            headers: list[tuple[int, str]] = []
            for c in range(col, ws.max_column + 1):
                name = ws.cell(row=row, column=c).value
                name = str(name).strip() if name is not None else ""
                # 锚点两列不是数据列；空列名跳过（业务方插的空白列）
                if name and name not in (ANCHOR_RID, ANCHOR_HASH):
                    headers.append((c, name))
            return row, col, headers
    return None


def _looks_like_data(ws: Any) -> bool:
    """这张表看起来是不是本该有数据的。

    用来区分"业务方自己加的说明页"和"锚点被破坏的数据表"—— 前者读不了是正常的，
    后者读不了是事故。判据是行列规模，不是表名：以前靠 ``00_`` 前缀判断，
    于是把 ``03_关系清单`` 改名成 ``00_关系清单`` 就能让 19 行静默消失。
    """
    return ws.max_row > 3 and ws.max_column >= 3


# ══════════════════════════════════════════════════════════════════
#  审核
# ══════════════════════════════════════════════════════════════════
class ReturnAuditor:
    """把回传件审成一份可执行的处置清单。

    Args:
        target: 达标线。低于它就打回，不发布。
    """

    def __init__(self, *, target: float = 0.95) -> None:
        self.target = target


    def _damage_findings(self, damage: list[str]) -> list[Conflict]:
        """把结构损伤变成一条要人处理的冲突。

        只放进 summary 的一个列表里是不够的 —— 那一行会被淹没在完成度、打回单
        之间。它得和别的问题一样进 findings，因为它的后果比任何一条填错都严重：
        **一整张表的人白填了，而他们收不到任何提示。**
        """
        if not damage:
            return []
        return [Conflict(
            rid="cf_return_damaged",
            kind=ConflictKind.PERFUNCTORY,   # 复用"这份回传不可信"这一档
            summary="回传件结构被改动，有内容没能读回来：\n" + "\n".join(damage[:6]),
            subjects=[], evidence=[], detector="rule:return_structure",
            options=[])]

    def audit(
        self,
        spec: TemplateSpec,
        returned: dict[tuple[str, str], str],
        *,
        oir: OIR | None = None,
    ) -> AuditResult:
        expected = spec.by_rid()
        # 损伤信息由 read_returned 塞在一个保留 key 里。取出来之后必须从
        # returned 里拿掉，否则它会被当成一个 rid 进 new_rows。
        damage = returned.pop(("__damage__", "__damage__"), "")
        diffs = [
            CellDiff(rid=cell.rid, sheet=cell.sheet, field=cell.field,
                     before=cell.value, after=returned.get(key, ""),
                     role=cell.role, owner=cell.owner)
            for key, cell in expected.items()
        ]
        result = AuditResult(completeness=0.0, diffs=diffs,
                             damage=[x for x in damage.split("\n") if x])
        result.unmatched_rows = sorted(
            {c.rid for k, c in expected.items() if k not in returned}
        )
        result.new_rows = sorted({rid for rid, _ in returned} - {c.rid for c in expected.values()})

        result.findings = [
            *self._damage_findings(result.damage),
            *self._missing(diffs),
            *self._enum_violations(diffs, expected),
            *self._naming(diffs),
            *self._perfunctory(diffs, expected),
            *self._divergence(diffs),
        ]
        result.auto_repaired = self._auto_repair(result.findings, oir)
        result.findings = [c for c in result.findings
                           if c.rid not in {r["conflict"] for r in result.auto_repaired}]
        result.completeness = self.completeness(diffs, expected)
        result.slips = self._slips(result.findings)
        return result

    # ── 敷衍信号（唯一入口）────────────────────────────────────
    @staticmethod
    def signals(d: CellDiff, cell: Cell | None, ratio: float) -> list[str]:
        """一格的敷衍信号。

        **完成度计算和打回单必须走同一个入口** —— 之前两处各自调用、参数还不同，
        结果就是修好一处另一处照旧误报。判「太短/整列一个值」需要知道这格是不是
        自由文本，而只有编译期的 :class:`Cell` 知道。
        """
        prose = bool(cell and cell.expects_prose)
        return perfunctory_signals(
            value=d.after, column_header=d.field, ai_prefill=d.before,
            expects_definition=prose,
            column_distinct_ratio=ratio if prose else 1.0,
        )

    @staticmethod
    def _ratios(diffs: list[CellDiff]) -> dict[str, float]:
        """各字段的取值离散度。整列一个值是敷衍信号（仅对自由文本成立）。"""
        cols: dict[str, list[str]] = {}
        for d in diffs:
            cols.setdefault(d.field, []).append(d.after)
        out: dict[str, float] = {}
        for f, vals in cols.items():
            nonempty = [v for v in vals if v]
            out[f] = len(set(nonempty)) / len(nonempty) if nonempty else 1.0
        return out

    # ── 完成度 ──────────────────────────────────────────────────
    def completeness(self, diffs: list[CellDiff], expected: dict | None = None) -> float:
        """加权完成度。

        只统计**业务必填**的格子 —— 系统自己填好的部分算进去会把数字虚高，
        而这个数字是 FDE 决定要不要再发一轮的依据。原样交回的预填不算「已填」，
        命中敷衍信号的也不算。
        """
        target = [d for d in diffs if d.role is Role.REQUIRED]
        if not target:
            return 1.0
        expected = expected or {}
        ratios = self._ratios(diffs)
        total = sum(WEIGHTS.get(d.field, 1.0) for d in target)
        got = sum(
            WEIGHTS.get(d.field, 1.0) for d in target
            if d.filled and not d.untouched_prefill
            and not self.signals(d, expected.get((d.rid, d.field)), ratios.get(d.field, 1.0))
        )
        return got / total if total else 1.0

    # ── 规则审核 ────────────────────────────────────────────────
    @staticmethod
    def _missing(diffs: list[CellDiff]) -> list[Conflict]:
        return [
            Conflict(
                rid=_cid(ConflictKind.MISSING_REQUIRED, d.rid, d.field),
                kind=ConflictKind.MISSING_REQUIRED, subjects=[d.rid],
                summary=f"{d.sheet} · {d.field} 未填",
                owner=d.owner, detector="rule:REQ-01",
            )
            for d in diffs if d.role is Role.REQUIRED and not d.filled
        ]

    @staticmethod
    def _enum_violations(diffs: list[CellDiff], expected: dict) -> list[Conflict]:
        out: list[Conflict] = []
        for d in diffs:
            cell: Cell | None = expected.get((d.rid, d.field))
            if cell is None or not cell.options or not d.filled:
                continue
            if d.after not in cell.options:
                out.append(Conflict(
                    rid=_cid(ConflictKind.TYPE_MISMATCH, d.rid, d.field),
                    kind=ConflictKind.TYPE_MISMATCH, subjects=[d.rid],
                    summary=f"{d.sheet} · {d.field} 取值「{d.after}」不在允许集合 "
                            f"{'、'.join(cell.options)}",
                    owner=d.owner, detector="rule:ENUM-01",
                ))
        return out

    @staticmethod
    def _naming(diffs: list[CellDiff]) -> list[Conflict]:
        from .conflict import _CAMEL, _HAS_CJK, Option

        out: list[Conflict] = []
        for d in diffs:
            if d.field != "apiName" or not d.filled or not d.changed:
                continue  # 只审业务方改过的；没动过的在编译期已经合规
            if _HAS_CJK.search(d.after) or not _CAMEL.match(d.after):
                out.append(Conflict(
                    rid=_cid(ConflictKind.NAMING_VIOLATION, d.rid, d.field),
                    kind=ConflictKind.NAMING_VIOLATION, subjects=[d.rid],
                    summary=f"{d.sheet} · apiName「{d.after}」不符合 lowerCamelCase",
                    owner=d.owner, detector="rule:NAME-01",
                    options=[Option("apply", f"改为 {to_camel(d.after)}", "可逆、零语义损失",
                                    effect={"set_api_name": to_camel(d.after)})],
                ))
        return out

    def _perfunctory(self, diffs: list[CellDiff], expected: dict) -> list[Conflict]:
        """启发式收窄候选。命中信号的才进 LLM 终判队列，没命中的连模型都不用调。"""
        out: list[Conflict] = []
        ratios = self._ratios(diffs)
        for d in diffs:
            if d.role is not Role.REQUIRED or not d.filled:
                continue
            sig = self.signals(d, expected.get((d.rid, d.field)), ratios.get(d.field, 1.0))
            if not sig:
                continue
            out.append(Conflict(
                rid=_cid(ConflictKind.PERFUNCTORY, d.rid, d.field),
                kind=ConflictKind.PERFUNCTORY, subjects=[d.rid],
                summary=f"{d.sheet} · {d.field} 疑似敷衍（{'、'.join(sig)}）"
                        + ("：AI 预填值被原样交回，这格没有被真正审过"
                           if "UNCHANGED_PREFILL" in sig else ""),
                owner=d.owner, detector="heuristic:PERF-01",
            ))
        return out

    @staticmethod
    def _divergence(diffs: list[CellDiff]) -> list[Conflict]:
        """回传后新产生的口径矛盾 —— 两个人各填各的，轴上冲突。"""
        by_field: dict[str, list[CellDiff]] = {}
        for d in diffs:
            if d.field == "definition" and d.filled:
                by_field.setdefault(d.rid.rsplit("_", 1)[0], []).append(d)

        out: list[Conflict] = []
        for group in by_field.values():
            for i, a in enumerate(group):
                for b in group[i + 1 :]:
                    diff = axis_diff(a.after, b.after)
                    if not diff:
                        continue
                    axes = "、".join(f"{k}（{va} vs {vb}）" for k, (va, vb) in diff.items())
                    owners = sorted({o for o in (a.owner, b.owner) if o})
                    out.append(Conflict(
                        rid=_cid(ConflictKind.SEMANTIC_DIVERGENCE, a.rid, b.rid, "returned"),
                        kind=ConflictKind.SEMANTIC_DIVERGENCE, subjects=[a.rid, b.rid],
                        summary=f"回传后仍存在口径矛盾：{axes}"
                                + (f"，需 {'、'.join(owners)} 对齐" if len(owners) > 1 else ""),
                        owner=owners[0] if owners else None, detector="rule:AXIS-01",
                    ))
        return out

    # ── 自动修 ──────────────────────────────────────────────────
    @staticmethod
    def _auto_repair(findings: list[Conflict], oir: OIR | None) -> list[dict[str, Any]]:
        """只修可逆、零语义损失、可完整记账的。命名满足；口径统一不满足。"""
        log: list[dict[str, Any]] = []
        for c in findings:
            if c.handling is not Handling.AUTO_REPAIR or not c.options:
                continue
            new_name = c.options[0].effect.get("set_api_name")
            if not new_name:
                continue
            entry = {"conflict": c.rid, "rid": c.subjects[0], "field": "apiName",
                     "to": new_name, "reversible": True}
            if oir is not None:
                for bucket in (oir.objects, oir.properties, oir.links, oir.actions):
                    if (e := bucket.get(c.subjects[0])) is not None:
                        entry["from"] = e.api_name.value
                        e.api_name.value = new_name
                        e.api_name.origin = Origin.AUTO_REPAIRED
            log.append(entry)
        return log

    # ── 打回单 ──────────────────────────────────────────────────
    @staticmethod
    def _slips(findings: list[Conflict]) -> list[ReturnSlip]:
        grouped: dict[str, ReturnSlip] = {}
        for c in findings:
            if c.handling not in (Handling.ROUND_TRIP, Handling.ASK_USER):
                continue
            owner = c.owner or "未分派"
            grouped.setdefault(owner, ReturnSlip(owner)).items.append(c)
        return sorted(grouped.values(), key=lambda s: -len(s.items))


# ══════════════════════════════════════════════════════════════════
#  回写
# ══════════════════════════════════════════════════════════════════
def merge_into_oir(oir: OIR, diffs: list[CellDiff]) -> tuple[list[str], list[str]]:
    """把业务方真正填的内容写回 OIR，标 ``Origin.USER``。

    返回 ``(写回了什么, 读到了却没地方放的)``。

    第二个返回值是新加的，因为它对应一个实测过的黑洞：**627 格被判为真正填写，
    最后只有 172 条落回 OIR**。primaryKey / owner / effects 共 455 格被读进来、
    被算进完成度、然后在这个 ``match`` 的 ``case _: continue`` 里静默丢弃 ——
    填表的人以为答案生效了，FDE 看到完成度上升，而模型里什么都没变。

    丢弃有时是对的（有些列本来就只是给人核对的），但**丢了必须说**。

    **原样交回的预填不算填写** —— 那格没有被人认领过，不该获得 USER 的可信度。
    """
    changed: list[str] = []
    dropped: list[str] = []
    for d in diffs:
        if not d.filled or d.untouched_prefill:
            continue
        entity = (oir.properties.get(d.rid) or oir.objects.get(d.rid)
                  or oir.links.get(d.rid) or oir.actions.get(d.rid)
                  # 规则和问题以前不在这个链上 —— 即使模板里有了它们的 sheet，
                  # 回写这一步也接不住。
                  or oir.rules.get(d.rid) or oir.questions.get(d.rid))
        if entity is None:
            dropped.append(f"{d.rid}.{d.field}（找不到这个条目）")
            continue
        note = f"业务方回传（{d.owner or '未分派'}）"
        match d.field:
            case "definition" if hasattr(entity, "definition"):
                entity.definition = by_user(d.after, note=note)
            case "description" if hasattr(entity, "description"):
                entity.description = by_user(d.after, note=note)
            case "displayName":
                entity.display_name = by_user(d.after, note=note)
            case "cardinality" if hasattr(entity, "cardinality"):
                entity.cardinality = by_user(d.after, note=note)
            case "owner":
                entity.owner = d.after
            case "primaryKey" if hasattr(entity, "primary_key"):
                entity.primary_key = by_user(
                    [x.strip() for x in d.after.replace("，", "、").split("、") if x.strip()],
                    note=note)
            case "effects" if hasattr(entity, "effects"):
                entity.effects = by_user(
                    [x.strip() for x in d.after.replace("，", "、").split("、") if x.strip()],
                    note=note)
            # ── 新表的回写 ──────────────────────────────────────
            case "答复" if hasattr(entity, "answer"):
                # 答了就不再是待办。**这是整轮往返里最有价值的一次写入** ——
                # 一个待澄清问题变成了一条事实。
                entity.answer = by_user(d.after, note=note)
                entity.status = Status.CONFIRMED
            case "执行角色" if hasattr(entity, "actor"):
                entity.actor = by_user(d.after, note=note)
            case "这条对吗" if hasattr(entity, "statement"):
                entity.status = (Status.CONFIRMED if d.after.strip() == "对"
                                 else Status.REJECTED if d.after.strip() == "不对"
                                 else Status.PROPOSED)
            case "管哪个单据":
                by_name = {o.display_name.value: rid for rid, o in oir.objects.items()}
                hosts = [by_name[x.strip()] for x in d.after.replace("，", "、").split("、")
                         if x.strip() in by_name]
                if not hosts:
                    dropped.append(f"{d.rid}.{d.field}（「{d.after[:20]}」对不上任何对象）")
                    continue
                entity.applies_to = hosts
            case _:
                dropped.append(f"{d.rid}.{d.field}（没有回写路径）")
                continue
        changed.append(f"{d.rid}.{d.field}")
    return changed, dropped
