"""把对话里的一段内容落成可下载的文件（md / csv / xlsx / docx / pdf）。

FDE 在聊天里问出了一张 192 行的待澄清问题表，下一句是"把这个转成 excel 给我"。
交付包（``bundle.py``）解决不了这件事 —— 那是把**整个会话的产物**打成一个 zip，
而他要的是**刚刚那一份东西**，按他指定的格式。

所以这里的输入不是会话，是一份 :class:`ExportDoc`：一个标题 + 一串有序的块
（标题/正文/表格/代码）。谁去组装这个 doc 是调用方的事（服务端从产物、从材料、
从上一张表、从上一条回答都能组装），这个模块只负责把它写成字节。

和 ``bundle.py`` 同一条规矩：**纯函数、不 import server**，所以能脱离服务单测。
第三方库一律**函数内惰性 import** —— 与 ``template.py`` 写 xlsx 的手法一致，
装不上某个库时只是那一种格式不可用，不会让整个模块 import 失败。

三个踩过的坑，写在这里免得下次再踩：

1. **PDF 的中文不是字体路径问题。** PyMuPDF 自带 MuPDF 的 Droid Sans Fallback，
   ``Story`` 在空 CSS 下就会自动回退，中文能画也能被复制出来 —— 不需要装字体、
   不需要 pymupdf-fonts（``css_for_pymupdf_font()`` 在没装它时直接抛异常），更不要
   硬编码 ``/System/Library/Fonts/PingFang.ttc``（这台 macOS 上根本没有这个路径）。
2. **但一定要 ``subset_fonts()``。** 不做子集化，一页中文 PDF 是 3.6 MB（整个回退
   字体被嵌进去）；做了是 27 KB，中文照样能提取。差 130 倍。
3. **docx 的中文要设 ``w:eastAsia``。** python-docx 只写 ``w:ascii``/``w:hAnsi``，
   Word 打开时中文会被替换成别的字体。设在 Normal 样式上一次搞定，比逐个 run
   打补丁可靠（新建 run 的 ``rPr`` 是 None，得先碰一下 ``font.name`` 才存在）。
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from typing import Any

__all__ = [
    "FORMATS",
    "SPECS",
    "Block",
    "ExportDoc",
    "ExportSpec",
    "blocks_from_markdown",
    "render",
    "safe_name",
    "table_block",
]


# ══════════════════════════════════════════════════════════════════
#  文档模型
# ══════════════════════════════════════════════════════════════════
@dataclass
class Block:
    """一个内容块。五种写法共用同一份输入，只有 ``to_*`` 不同。"""

    kind: str                                   # heading | para | table | code | rule
    text: str = ""
    level: int = 1                              # heading 用
    #: 列表项：(缩进层级, 序号或"", 文本)。para 为空表示是普通段落。
    items: list[tuple[int, str, str]] = field(default_factory=list)
    columns: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


@dataclass
class ExportDoc:
    """要导出的一份东西。"""

    title: str
    blocks: list[Block] = field(default_factory=list)
    #: 一句话说清这份文件是从哪来的 —— 导出件脱离了对话，没有出处就是孤证。
    note: str = ""

    @property
    def tables(self) -> list[Block]:
        return [b for b in self.blocks if b.kind == "table"]


def table_block(columns: list[str], rows: list[list[str]], title: str = "") -> list[Block]:
    """一张表 → 块列表（带可选小标题）。服务端最常用的就是这个。"""
    out = [Block("heading", title, level=2)] if title else []
    out.append(Block("table", columns=list(columns),
                     rows=[[str(c) if c is not None else "" for c in r] for r in rows]))
    return out


# ══════════════════════════════════════════════════════════════════
#  markdown → 块
# ══════════════════════════════════════════════════════════════════
#: 认的语法和前端 `md()` 保持一致：代码围栏 / #..###### / 竖线表格 / 有序无序列表 /
#: 分隔线 / 段落。**不是完整 CommonMark** —— 覆盖的是模型实际会写的那几种。
_FENCE = re.compile(r"^\s*```(\w*)\s*$")
_HEAD = re.compile(r"^\s*(#{1,6})\s+(.+?)\s*#*$")
_RULE = re.compile(r"^\s*(?:---+|\*\*\*+|___+)\s*$")
_ITEM = re.compile(r"^(\s*)(\d+[.)]|[-*+])\s+(.*)$")
_ROW = re.compile(r"^\s*\|.*\|\s*$")
_SEP = re.compile(r"^\s*\|[\s:|-]+\|\s*$")
#: 行内标记在 docx/pdf 里没有对应表达（不做富文本），落地前先去掉，免得正文里
#: 出现裸的 ** 和 `。
_INLINE = re.compile(r"\*\*(.+?)\*\*|`(.+?)`|(?<![*\w])\*([^*\n]+)\*(?![*\w])")


def _plain(s: str) -> str:
    return _INLINE.sub(lambda m: m.group(1) or m.group(2) or m.group(3) or "", s).strip()


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def blocks_from_markdown(text: str) -> list[Block]:
    """把一段 markdown 切成块。回答（模型写的 markdown）走这条路。"""
    lines = (text or "").split("\n")
    out: list[Block] = []
    buf: list[str] = []
    items: list[tuple[int, str, str]] = []

    def flush() -> None:
        nonlocal buf, items
        if items:
            out.append(Block("para", items=items))
            items = []
        if buf:
            body = "\n".join(buf).strip()
            if body:
                out.append(Block("para", _plain(body)))
            buf = []

    i = 0
    while i < len(lines):
        ln = lines[i]

        m = _FENCE.match(ln)
        if m:                                        # 代码块：原样保留，不做行内处理
            flush()
            body, i = [], i + 1
            while i < len(lines) and not _FENCE.match(lines[i]):
                body.append(lines[i])
                i += 1
            out.append(Block("code", "\n".join(body)))
            i += 1
            continue

        if _ROW.match(ln) and i + 1 < len(lines) and _SEP.match(lines[i + 1]):
            flush()
            head = _cells(ln)
            body: list[list[str]] = []
            i += 2
            while i < len(lines) and _ROW.match(lines[i]):
                body.append([_plain(c) for c in _cells(lines[i])])
                i += 1
            out.append(Block("table", columns=[_plain(c) for c in head], rows=body))
            continue

        m = _HEAD.match(ln)
        if m:
            flush()
            out.append(Block("heading", _plain(m.group(2)), level=len(m.group(1))))
            i += 1
            continue

        if _RULE.match(ln):
            flush()
            out.append(Block("rule"))
            i += 1
            continue

        m = _ITEM.match(ln)
        if m:
            if buf:                                  # 段落先收尾，列表另起一块
                body = "\n".join(buf).strip()
                if body:
                    out.append(Block("para", _plain(body)))
                buf = []
            body = _plain(m.group(3))
            if body:                                 # 只有标记没内容的行不算一条
                depth = len(m.group(1).replace("\t", "    ")) // 2
                marker = m.group(2) if m.group(2)[0].isdigit() else ""
                items.append((depth, marker, body))
            i += 1
            continue

        if not ln.strip():
            if buf:
                flush()
            i += 1
            continue

        if items:                                    # 列表被普通正文打断
            flush()
        buf.append(ln)
        i += 1

    flush()
    return out


# ══════════════════════════════════════════════════════════════════
#  markdown
# ══════════════════════════════════════════════════════════════════
def to_markdown(doc: ExportDoc) -> bytes:
    out = [f"# {doc.title}", ""]
    if doc.note:
        out += [f"> {doc.note}", ""]
    for b in doc.blocks:
        if b.kind == "heading":
            out += ["#" * max(2, min(6, b.level + 1)) + " " + b.text, ""]
        elif b.kind == "rule":
            out += ["---", ""]
        elif b.kind == "code":
            out += ["```", b.text, "```", ""]
        elif b.kind == "table":
            out.append("| " + " | ".join(b.columns) + " |")
            out.append("| " + " | ".join("---" for _ in b.columns) + " |")
            for r in b.rows:
                # 竖线和换行会把表格结构撑破，落地时替换掉
                out.append("| " + " | ".join(
                    str(c).replace("|", "丨").replace("\n", "<br>") for c in r) + " |")
            out.append("")
        elif b.items:
            for depth, marker, txt in b.items:
                out.append("  " * depth + (f"{marker} " if marker else "- ") + txt)
            out.append("")
        elif b.text:
            out += [b.text, ""]
    return "\n".join(out).rstrip().encode("utf-8")


# ══════════════════════════════════════════════════════════════════
#  csv
# ══════════════════════════════════════════════════════════════════
def to_csv(doc: ExportDoc) -> bytes:
    """只导表格。没有表就退化成一列正文 —— 空文件比一句解释更让人困惑。"""
    buf = io.StringIO()
    w = csv.writer(buf)
    tables = doc.tables
    if tables:
        for n, b in enumerate(tables):
            if n:
                w.writerow([])
            w.writerow([_csv_safe(c) for c in b.columns])
            w.writerows([_csv_safe(c) for c in r] for r in b.rows)
    else:
        w.writerow([_csv_safe(doc.title)])
        for b in doc.blocks:
            for line in _flatten(b):
                w.writerow([_csv_safe(line)])
    # BOM：Excel 打开无 BOM 的 UTF-8 csv 会把中文显示成乱码，而"导出的中文是乱码"
    # 是最容易被当成系统坏了的一种失败。
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


#: csv 里以这些字符开头的格子，Excel/Sheets 打开时会当成公式求值。
_CSV_TRIGGER = ("=", "+", "-", "@", "\t", "\r")


def _csv_safe(v: Any) -> str:
    """csv 没有 xlsx 那种"这格是文本"的类型位，只能靠前缀。

    加一个单引号 —— 在 csv 这条路上它是**必要**的（xlsx 那边靠 data_type 就够，
    所以那边反而不该加，加了 Excel 会原样显示出来）。不做这一步的后果：一份客户
    给的表里一格 ``=HYPERLINK("http://x?"&A1)``，FDE 导成 csv、双击打开，
    Excel 就把隔壁格子的内容发出去了；而代码还特意加了 BOM 让 Excel 乐意打开它。
    """
    s = "" if v is None else str(v)
    return "'" + s if s[:1] in _CSV_TRIGGER else s


def _flatten(b: Block) -> list[str]:
    """块 → 纯文本行。csv/兜底路径用。"""
    if b.kind == "table":
        return [" | ".join(b.columns)] + [" | ".join(str(c) for c in r) for r in b.rows]
    if b.items:
        return ["  " * d + (f"{m} " if m else "· ") + t for d, m, t in b.items]
    if b.kind == "rule":
        return ["—" * 20]
    return [b.text] if b.text else []


# ══════════════════════════════════════════════════════════════════
#  xlsx
# ══════════════════════════════════════════════════════════════════
def to_xlsx(doc: ExportDoc) -> bytes:
    """表格一张一个 sheet；正文另开一个「说明」sheet。

    表和正文分开放，是因为它们的用法不同：表要被筛选/透视/粘进别的表，正文只是
    看一眼。把正文塞在表上面，会让第一行不是表头，Excel 的筛选和冻结全都不对。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="EFEDE8")
    head_font = Font(bold=True)
    thin = Side(style="thin", color="D8D5CE")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(vertical="top", wrap_text=True)

    tables = doc.tables
    used: set[str] = set()
    for n, b in enumerate(tables):
        title = _sheet_title(_table_title(doc, b) or doc.title or "表", n, used)
        ws = wb.create_sheet(title)
        ws.append(list(b.columns))
        for c in ws[1]:
            c.fill, c.font, c.border, c.alignment = head_fill, head_font, box, wrap
        for r in b.rows:
            ws.append([_cell(v) for v in r])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.border, c.alignment = box, wrap
        _detext(ws)
        # 列宽按内容估，但封顶 —— 一条 200 字的澄清问题会把列拉到屏幕外
        for i, name in enumerate(b.columns, start=1):
            width = max([len(str(name))] + [len(str(r[i - 1])) for r in b.rows
                                            if i - 1 < len(r)] or [8])
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    prose = [ln for b in doc.blocks if b.kind != "table" for ln in _flatten(b)]
    if prose or not tables:
        ws = wb.create_sheet("说明" if tables else (doc.title[:28] or "内容"))
        ws.column_dimensions["A"].width = 100
        ws.append([doc.title])
        ws["A1"].font = Font(bold=True, size=13)
        if doc.note:
            ws.append([doc.note])
        ws.append([])
        for ln in prose:
            ws.append([ln])
        for row in ws.iter_rows(min_col=1, max_col=1):
            row[0].alignment = wrap
        _detext(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _detext(ws: Any) -> None:
    """把这一页里所有被当成公式的格子降回文本。

    openpyxl 看到以 ``=`` 开头的字符串就写成公式（``data_type='f'``）。内容是
    **别人给的**（客户的表、模型抽出来的口径），所以一格 ``=HYPERLINK(...)``
    会在 FDE 双击打开时被 Excel 求值。以前只扫了数据行，漏掉表头和说明页 ——
    而表头正是列名，恰恰来自那份外来表格。整页扫，一格不漏。
    """
    for row in ws.iter_rows():
        for c in row:
            if c.data_type == "f":
                c.data_type = "s"


def _cell(v: Any) -> Any:
    """openpyxl 只接受标量；别的原样 str 化（公式那一层在写完之后按 data_type 纠）。"""
    return v if isinstance(v, (int, float, bool)) or v is None else str(v)


def _table_title(doc: ExportDoc, b: Block) -> str:
    """表格前面紧挨着的那个小标题就是它的名字。"""
    prev = ""
    for x in doc.blocks:
        if x is b:
            return prev
        prev = x.text if x.kind == "heading" else prev
    return ""


_BAD_SHEET = re.compile(r"[\\/*?:\[\]]")


def _sheet_title(name: str, n: int, used: set[str]) -> str:
    """Excel 的表名规矩很硬：≤31 字符、不能有 \\/*?:[]、不能重名。踩中就存不出文件。"""
    t = _BAD_SHEET.sub("-", (name or "").strip()) or f"表{n + 1}"
    t = t[:31] or f"表{n + 1}"
    base, i = t, 2
    while t.lower() in used:
        suffix = f"({i})"
        t = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(t.lower())
    return t


# ══════════════════════════════════════════════════════════════════
#  docx
# ══════════════════════════════════════════════════════════════════
#: Word 里靠谱的中文字体：优先微软雅黑（Windows/Office 必有），macOS 上 Word 会
#: 自己回退。设在 Normal 样式上，全篇继承。
_EAST_ASIA = "Microsoft YaHei"


def to_docx(doc: ExportDoc) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    d = docx.Document()
    normal = d.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = "Calibri"
    # 关键的一行：不设 w:eastAsia，Word 打开时中文会被换成别的字体
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), _EAST_ASIA)

    d.add_heading(doc.title, level=0)
    if doc.note:
        p = d.add_paragraph(doc.note)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.size = Pt(9)
            run.italic = True

    for b in doc.blocks:
        if b.kind == "heading":
            d.add_heading(b.text, level=max(1, min(4, b.level)))
        elif b.kind == "rule":
            d.add_paragraph("—" * 30)
        elif b.kind == "code":
            p = d.add_paragraph()
            run = p.add_run(b.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif b.kind == "table":
            if not b.columns:
                continue
            t = d.add_table(rows=1, cols=len(b.columns))
            t.style = "Table Grid"
            for i, name in enumerate(b.columns):
                cell = t.rows[0].cells[i]
                cell.text = str(name)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
            for r in b.rows:
                cells = t.add_row().cells
                for i in range(len(b.columns)):
                    cells[i].text = str(r[i]) if i < len(r) else ""
        elif b.items:
            for depth, marker, txt in b.items:
                style = "List Number" if marker else "List Bullet"
                # Word 内置样式只到三级；再深就退回同一级，别让它抛 KeyError
                lvl = min(depth, 2)
                name = style if not lvl else f"{style} {lvl + 1}"
                try:
                    d.add_paragraph(txt, style=name)
                except KeyError:
                    d.add_paragraph(("    " * depth) + "· " + txt)
        elif b.text:
            d.add_paragraph(b.text)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════
#  pdf
# ══════════════════════════════════════════════════════════════════
_PDF_CSS = """
* { font-family: sans-serif; }
h1 { font-size: 17px; margin: 0 0 10px 0; }
h2 { font-size: 13.5px; margin: 14px 0 5px 0; }
h3, h4, h5, h6 { font-size: 12px; margin: 11px 0 4px 0; }
p  { font-size: 10px; margin: 0 0 7px 0; line-height: 1.5; }
.note { font-size: 8.5px; color: #666; margin-bottom: 12px; }
li { font-size: 10px; margin-bottom: 3px; }
pre { font-family: monospace; font-size: 8.5px; background: #f4f2ee; padding: 6px; }
table { width: 100%; border-collapse: collapse; margin: 4px 0 12px 0; }
th { font-size: 9px; text-align: left; background: #efede8;
     border: 1px solid #d8d5ce; padding: 3px 5px; }
td { font-size: 9px; border: 1px solid #d8d5ce; padding: 3px 5px; vertical-align: top; }
"""


def _esc(s: Any) -> str:
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def doc_to_html(doc: ExportDoc) -> str:
    """块 → HTML。PDF 走这条路（Story 吃 HTML），也方便单测看结构。"""
    out = [f"<h1>{_esc(doc.title)}</h1>"]
    if doc.note:
        out.append(f'<p class="note">{_esc(doc.note)}</p>')
    for b in doc.blocks:
        if b.kind == "heading":
            lvl = max(2, min(6, b.level + 1))
            out.append(f"<h{lvl}>{_esc(b.text)}</h{lvl}>")
        elif b.kind == "rule":
            out.append("<hr/>")
        elif b.kind == "code":
            out.append(f"<pre>{_esc(b.text)}</pre>")
        elif b.kind == "table":
            out.append("<table><tr>"
                       + "".join(f"<th>{_esc(c)}</th>" for c in b.columns) + "</tr>")
            for r in b.rows:
                out.append("<tr>" + "".join(f"<td>{_esc(c)}</td>" for c in r) + "</tr>")
            out.append("</table>")
        elif b.items:
            # 缩进用 margin 表达：Story 对嵌套 ul/ol 的支持不稳，扁平列表更可靠
            out.append("<ul>")
            for depth, marker, txt in b.items:
                pad = depth * 16
                lead = f"{_esc(marker)} " if marker else ""
                out.append(f'<li style="margin-left:{pad}px">{lead}{_esc(txt)}</li>')
            out.append("</ul>")
        elif b.text:
            out.append(f"<p>{_esc(b.text)}</p>")
    return "\n".join(out)


def to_pdf(doc: ExportDoc) -> bytes:
    import pymupdf

    page = pymupdf.paper_rect("a4")
    frame = page + (40, 40, -40, -50)
    buf = io.BytesIO()
    story = pymupdf.Story(html=doc_to_html(doc), user_css=_PDF_CSS)
    writer = pymupdf.DocumentWriter(buf)
    more, guard = 1, 0
    while more:
        dev = writer.begin_page(page)
        more, _ = story.place(frame)
        story.draw(dev)
        writer.end_page()
        guard += 1
        if guard > 400:          # 排不下的单个元素会让 place 永远返回 1
            break
    writer.close()

    # **不做子集化就是 3.6 MB。** 整个中文回退字体会被原样嵌进去；子集化后 27 KB，
    # 中文照样能提取。这一行不是优化，是能不能用的分界。
    pdf = pymupdf.open("pdf", buf.getvalue())
    try:
        pdf.subset_fonts()
    except Exception as exc:  # noqa: BLE001 — 子集化失败只是文件大，不该导不出来
        import warnings

        warnings.warn(f"PDF 字体子集化失败，文件会偏大：{exc}", stacklevel=2)
    return pdf.tobytes(deflate=True, garbage=4)


# ══════════════════════════════════════════════════════════════════
#  格式表
# ══════════════════════════════════════════════════════════════════
@dataclass(frozen=True)
class ExportSpec:
    ext: str
    media_type: str
    label: str
    write: Callable[[ExportDoc], bytes]


SPECS: dict[str, ExportSpec] = {
    "md": ExportSpec("md", "text/markdown; charset=utf-8", "Markdown", to_markdown),
    "csv": ExportSpec("csv", "text/csv; charset=utf-8", "CSV", to_csv),
    "xlsx": ExportSpec(
        "xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "Excel", to_xlsx),
    "docx": ExportSpec(
        "docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "Word", to_docx),
    "pdf": ExportSpec("pdf", "application/pdf", "PDF", to_pdf),
}

#: 别名：用户嘴里说的和格式名往往不是一个词（"excel"、"表格"、"word"、"文档"）。
_ALIAS = {"excel": "xlsx", "xls": "xlsx", "表格": "xlsx", "spreadsheet": "xlsx",
          "word": "docx", "doc": "docx", "文档": "docx",
          "markdown": "md", "文本": "md", "txt": "md"}

FORMATS = tuple(SPECS)


def resolve_format(fmt: str) -> str:
    f = (fmt or "").strip().lower().lstrip(".")
    return f if f in SPECS else _ALIAS.get(f, "")


def render(doc: ExportDoc, fmt: str) -> tuple[bytes, ExportSpec]:
    """写成字节。未知格式抛 ValueError（调用方转成给模型的回执）。"""
    key = resolve_format(fmt)
    if not key:
        raise ValueError(f"不支持的格式「{fmt}」。可用：{'/'.join(SPECS)}")
    spec = SPECS[key]
    return spec.write(doc), spec


#: 文件名里不能出现的字符（跨 Windows/macOS 取并集），外加控制字符。
_BAD_NAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]')


def safe_name(title: str, ext: str, *, fallback: str = "导出") -> str:
    """标题 → 文件名。**这也是路径穿越的唯一入口**，所以在这里就把分隔符干掉。"""
    t = _BAD_NAME.sub("_", (title or "").strip()).strip(". ")
    t = re.sub(r"\s+", " ", t)[:80].strip() or fallback
    if t in (".", ".."):
        t = fallback
    return f"{t}.{ext}"
