"""文本类材料 —— docx / md / txt。

流程说明文档的价值集中在**规则句**里："一个执行计划可拆入多个采购包"这一句
直接决定 LinkType 的基数。所以切片按**标题层级**分段而不是按固定长度切 ——
固定长度会把一条规则切成两半，两半都失去意义。

docx 里的表格单独抽出来，它们通常是字段清单，比正文更结构化。
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from .base import Finding, ParsedDoc, Parser, make_chunk

#: 一个切片的软上限（字符）。超了就在句号处切开，不硬切。
SECTION_SOFT_LIMIT = 1200

#: 建模相关的规则句特征。命中的切片打 rule 标签，检索时优先。
_RULE_HINTS = re.compile(
    r"(一个|每个|多个|至少|最多|必须|不得|应当|由.{0,8}(生成|下达|拆分|合并)"
    r"|对应|关联|一对多|多对多|口径|含税|不含税|唯一|主键)")

_HEADING = re.compile(r"^\s*(#{1,6}\s+|第[一二三四五六七八九十百]+[章节条]|"
                      r"\d+(\.\d+)*[、.\s]|[一二三四五六七八九十]+[、.])\s*")


class TextParser(Parser):
    """md / txt / 兜底。"""

    kind = "text"
    extensions = (".md", ".txt", ".markdown", ".rst")

    def parse(self, path: Path, *, file_id: str) -> ParsedDoc:
        doc = ParsedDoc(file_id=file_id, file_name=path.name, kind=self.kind)
        from .tabular import _read_text

        text, enc = _read_text(path)
        if enc != "utf-8":
            doc.findings.append(Finding("encoding_guess", f"按 {enc} 解码", {}))
        blocks = [(_heading_of(b), b) for b in _split_sections(text)]
        _emit(doc, blocks, file_id, path.name, kind="para")
        doc.structured = {"sections": len(blocks), "chars": len(text)}
        return doc


class DocxParser(Parser):
    kind = "docx"
    extensions = (".docx",)

    def parse(self, path: Path, *, file_id: str) -> ParsedDoc:
        import docx

        doc = ParsedDoc(file_id=file_id, file_name=path.name, kind=self.kind)
        d = docx.Document(str(path))

        core = d.core_properties
        meta = {k: v for k, v in (
            ("creator", core.author), ("last_modified_by", core.last_modified_by),
            ("modified", str(core.modified) if core.modified else None),
            ("title", core.title)) if v}
        doc.meta.update(meta)
        if leaked := {k: v for k, v in meta.items()
                      if k in ("creator", "last_modified_by")}:
            doc.findings.append(Finding(
                "metadata_leak",
                "文档元数据里带着作者信息（" +
                "；".join(f"{k}={v}" for k, v in leaked.items()) +
                "），对外发布前建议清理。",
                {"kind": "meta", "field": "docProps/core.xml"}, severity="warn"))

        # 正文：按标题层级聚段
        current = ""
        buf: list[str] = []
        blocks: list[tuple[str, str]] = []
        for p in d.paragraphs:
            t = (p.text or "").strip()
            if not t:
                continue
            if (p.style.name or "").lower().startswith("heading") or _HEADING.match(t):
                if buf:
                    blocks.append((current, "\n".join(buf)))
                    buf = []
                current = t
            else:
                buf.append(t)
                if sum(len(x) for x in buf) > SECTION_SOFT_LIMIT:
                    blocks.append((current, "\n".join(buf)))
                    buf = []
        if buf:
            blocks.append((current, "\n".join(buf)))
        order = _emit(doc, blocks, file_id, path.name, kind="para")

        # 表格单独抽 —— 通常是字段清单，比正文更结构化
        tables: list[dict[str, Any]] = []
        for ti, tbl in enumerate(d.tables):
            rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
            rows = [r for r in rows if any(r)]
            if not rows:
                continue
            header, body = rows[0], rows[1:]
            tables.append({"index": ti, "columns": header, "rows": len(body)})
            for ri, row in enumerate(body):
                doc.chunks.append(make_chunk(
                    doc_id=f"t{ti}r{ri}", file_id=file_id, file_name=path.name,
                    locator={"kind": "page", "page": ti + 1,
                             "bbox": [0, 0, 0, 0], "table": ti, "row": ri + 1},
                    render=" | ".join(f"{header[i]}={v}" for i, v in enumerate(row)
                                      if i < len(header) and v),
                    order=order, tags=["table"]))
                order += 1

        doc.structured = {"sections": len(blocks), "tables": tables}
        return doc


# ══════════════════════════════════════════════════════════════════
def _split_sections(text: str) -> list[str]:
    out: list[str] = []
    buf: list[str] = []
    for line in text.splitlines():
        if _HEADING.match(line) and buf:
            out.append("\n".join(buf).strip())
            buf = []
        buf.append(line)
        if sum(len(x) for x in buf) > SECTION_SOFT_LIMIT:
            joined = "\n".join(buf)
            cut = max(joined.rfind("。"), joined.rfind("\n\n"))
            if cut > len(joined) // 2:  # 在句号处切，不把规则句劈成两半
                out.append(joined[: cut + 1].strip())
                buf = [joined[cut + 1 :]]
            else:
                out.append(joined.strip())
                buf = []
    if buf and "\n".join(buf).strip():
        out.append("\n".join(buf).strip())
    return [b for b in out if b]


def _heading_of(block: str) -> str:
    first = block.strip().splitlines()[0] if block.strip() else ""
    return first[:60] if _HEADING.match(first) else ""


def _emit(doc: ParsedDoc, blocks: list[tuple[str, str]], file_id: str,
          file_name: str, *, kind: str) -> int:
    order = 0
    for heading, body in blocks:
        if not body.strip():
            continue
        is_rule = bool(_RULE_HINTS.search(body))
        doc.chunks.append(make_chunk(
            doc_id=f"s{order}", file_id=file_id, file_name=file_name,
            locator={"kind": "page", "page": 1, "bbox": [0, 0, 0, 0],
                     "section": heading or f"§{order + 1}"},
            render=(f"〔{heading}〕\n" if heading else "") + body,
            order=order, tags=[kind] + (["rule"] if is_rule else [])))
        order += 1
    return order
