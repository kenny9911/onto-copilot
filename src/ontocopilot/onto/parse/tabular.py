"""表格解析 —— xlsx / csv。

真实梳理表不长成规整的二维数组。这个解析器处理的都是实际会遇到的情况：

* **表头不在第一行。** 前面往往有标题、说明、空行。靠启发式定位表头行而不是
  假设它在第 1 行。
* **合并单元格伪装成分组。** 合并区只有左上角有值，直接读会得到一堆空格。
  向下/向右填充还原成真实的行列结构。
* **口径藏在批注里。** 单元格批注常常是最关键的信息（"这里指含税"），必须抽出来。
* **元数据里有客户名。** Office 文件在 ``workbook.xml`` 里保存绝对路径、在
  ``docProps/core.xml`` 里保存作者。这既是理解材料来源的线索，也是对外发布前
  必须清理的泄漏点。

列画像（唯一率、空值率、推断类型）由**确定性统计**产出，不交给模型 ——
LLM 无法可靠发现需要跨行分布理解的问题（arXiv:2503.06664）。
"""

from __future__ import annotations

import csv
import re
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

from .base import Finding, ParsedDoc, Parser, make_chunk

#: 表头探测的扫描深度。再深就不是表头了，是数据。
HEADER_SCAN_ROWS = 12
#: 一行要被当成表头，非空单元格至少要占这个比例。
HEADER_FILL_RATIO = 0.5


# ══════════════════════════════════════════════════════════════════
#  列画像
# ══════════════════════════════════════════════════════════════════
_INT = re.compile(r"^-?\d+$")
_DEC = re.compile(r"^-?\d+\.\d+$")
_DATE = re.compile(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}")
_BOOL = frozenset({"是", "否", "true", "false", "y", "n", "yes", "no"})


def infer_type(values: list[str]) -> str:
    """按实际取值推断类型。全空返回 STRING。"""
    vals = [v.strip() for v in values if v and v.strip()]
    if not vals:
        return "STRING"
    if all(_INT.match(v) for v in vals):
        return "INTEGER"
    if all(_INT.match(v) or _DEC.match(v) for v in vals):
        return "DECIMAL"
    if all(_DATE.match(v) for v in vals):
        return "DATE"
    if all(v.lower() in _BOOL for v in vals):
        return "BOOLEAN"
    # 取值集合小且重复率高 → 枚举
    if len(set(vals)) <= max(2, len(vals) // 8) and len(vals) >= 8:
        return "ENUM"
    return "STRING"


def profile_column(name: str, values: list[str]) -> dict[str, Any]:
    """一列的确定性画像。冲突检测的 TYPE_MISMATCH 用它。"""
    nonempty = [v for v in values if v and str(v).strip()]
    distinct = len(set(nonempty))
    return {
        "name": name,
        "count": len(values),
        "non_null": len(nonempty),
        "null_rate": round(1 - len(nonempty) / len(values), 4) if values else 1.0,
        "distinct": distinct,
        "distinct_ratio": round(distinct / len(nonempty), 4) if nonempty else 0.0,
        "unique": distinct == len(nonempty) and bool(nonempty),
        "inferred_type": infer_type(nonempty),
        "samples": list(dict.fromkeys(nonempty))[:5],
    }


# ══════════════════════════════════════════════════════════════════
#  表头探测
# ══════════════════════════════════════════════════════════════════

def _render_pairs(header: list[str], row: list[str],
                  comments: dict[int, str] | None = None) -> list[str]:
    """把一行渲染成 ``列名=值`` 列表，**同值的相邻列合并**。

    合并单元格被填充后，同一个值会出现在连续的好几列里。一行 22 列全是同一段
    职责说明时，不合并的话这一行的 render 就是那段话重复 22 遍 —— 一个 45 行的
    sheet 渲染出 27 万字符，模型只看得到开头，而开头全是重复。它会合理地推断
    "这些列是重复的"然后整段放弃。**真实材料上就是这么丢掉一整张业务规则表的。**

    ``comments`` 是 ``{列下标: 批注}``。批注**就地绑在被批注的那一列后面**，而不是
    甩到行尾 —— 口径几乎都写在批注里，"这条口径是哪一列的"必须跟着列走。值为空、
    没能落到任何 pair 的批注不丢，挂到末尾并标出所属列。
    """
    comments = comments or {}
    out: list[str] = []
    used: set[int] = set()
    i = 0
    n = min(len(header), len(row))
    while i < n:
        v = str(row[i]).strip()
        if not v:
            i += 1
            continue
        j = i + 1
        while j < n and str(row[j]).strip() == v:
            j += 1
        name = header[i] if j == i + 1 else f"{header[i]}~{header[j - 1]}"
        tail = "".join(f"　〔批注〕{comments[c]}" for c in range(i, j) if c in comments)
        used.update(range(i, j))
        out.append(f"{name}={v}{tail}")
        i = j
    for c, txt in comments.items():
        if c not in used and 0 <= c < len(header):
            out.append(f"〔批注·{header[c]}〕{txt}")
    return out


def _fill_hierarchy(header: list[str], body: list[list[str]]) -> list[list[str]]:
    """把"写一次、下面留空继承"的分组列向下补全 —— **只给 render 用，不动 raw**。

    合并单元格已在 :func:`_read_sheet` 里按物理区域填过；这里补的是没做物理合并、
    纯靠视觉缩进表达层级的分组列。``实体名称`` 最典型：每个实体写一次，其下的字段
    行留空,视觉上归属上面那个实体。不还原,``plan_amount`` 这样的字段行进了检索
    索引就不知道自己属于哪个实体,跨实体口径冲突永远归不了因。

    安全边界（宁可少补也不能误填"可选属性列"的空）：
    * 必须存在一个"密集锚列"（逐行都填、逐行变化的记录键，如 ``字段``）才动手；
      没有锚列说明这不是一行一记录的表,直接原样返回、不猜。
    * 只补锚列**左侧**、比锚列**更粗**（重复更多）的短文本列 —— 那才是层级分组；
      锚列右侧的属性列留空往往是"没有值",不能继承。
    * 只在"真数据行"（至少一个密集列有值）上继承,空白分隔行不碰。
    """
    n = len(body)
    if n < 2 or not header:
        return body
    ncol = len(header)

    def col(i: int) -> list[str]:
        return [(r[i].strip() if i < len(r) else "") for r in body]

    fill, dratio, maxlen = [], [], []
    for i in range(ncol):
        vals = col(i)
        nonblank = [v for v in vals if v]
        fill.append(len(nonblank) / n)
        dratio.append(len(set(nonblank)) / n if nonblank else 0.0)
        maxlen.append(max((len(v) for v in vals), default=0))

    # 记录键列：从左到右第一个"密集且逐行变化"的列（字段/编码）。
    key = next((i for i in range(ncol) if fill[i] >= 0.85 and dratio[i] >= 0.5), None)
    if key is None:
        return body
    # 分组列：键列左侧、比键列更粗、短文本、至少写过一次。
    carry_cols = [i for i in range(key)
                  if fill[i] > 0 and dratio[i] < dratio[key] and maxlen[i] <= 24]
    if not carry_cols:
        return body
    dense_cols = [i for i in range(ncol) if fill[i] >= 0.85]

    out = [list(r) + [""] * (ncol - len(r)) for r in body]
    last = dict.fromkeys(carry_cols, "")
    for r in out:
        is_data = any(r[d].strip() for d in dense_cols)
        for i in carry_cols:
            if r[i].strip():
                last[i] = r[i].strip()
            elif is_data and last[i]:
                r[i] = last[i]
    return out


def detect_header_row(rows: list[list[str]]) -> int:
    """找出表头在第几行（0-based）；**没有表头时返回 -1**。

    判据：填充率够高、几乎全是短文本、且**下一行的类型构成与它不同**。最后一条
    是关键 —— 数据区内部相邻两行的类型构成是一致的，表头与首行数据则不然。

    "没有表头"必须是一个可返回的答案。真实材料里有大量"第一行就是正文"的表
    （合并单元格里塞一整段职责说明、规则清单）。以前这里无论如何都会选一行当
    表头，于是那一整段散文变成了列名，下游每一行的 render 都背上它，切片被这段
    重复文本撑爆、真正的内容被挤掉。返回 -1 让调用方改用位置列名、正文从第 0
    行开始。
    """
    best, best_score = -1, 0.0
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        cells = [c.strip() for c in row]
        filled = [c for c in cells if c]
        if len(cells) == 0 or len(filled) / len(cells) < HEADER_FILL_RATIO:
            continue
        # 表头单元格通常短、非数值
        texty = sum(1 for c in filled if not _INT.match(c) and not _DEC.match(c) and len(c) <= 24)
        if texty / len(filled) < 0.5:
            continue  # 过半格子是长散文/数值 —— 这是正文，不是表头
        score = texty / len(filled)
        nxt = rows[i + 1] if i + 1 < len(rows) else []
        if nxt:
            nxt_filled = [c.strip() for c in nxt if str(c).strip()]
            nxt_numeric = sum(1 for c in nxt_filled
                              if _INT.match(c) or _DEC.match(c) or _DATE.match(c))
            if nxt_filled and nxt_numeric / len(nxt_filled) > 0.3:
                score += 0.5  # 下一行明显是数据 → 更像表头
            if len(set(nxt_filled)) == len(nxt_filled) and len(nxt_filled) > 1:
                score += 0.1
        score += max(0.0, 0.3 - i * 0.05)  # 靠前略加权，但不是决定性的
        if score > best_score:
            best, best_score = i, score
    return best


def dedupe_headers(header: list[str]) -> list[str]:
    """表头重名会让下游按名取列时静默取错，必须消歧。"""
    seen: Counter[str] = Counter()
    out: list[str] = []
    for i, h in enumerate(header):
        name = (h or "").strip() or f"col{i + 1}"
        seen[name] += 1
        out.append(name if seen[name] == 1 else f"{name}#{seen[name]}")
    return out


# ══════════════════════════════════════════════════════════════════
#  xlsx
# ══════════════════════════════════════════════════════════════════
class XlsxParser(Parser):
    kind = "xlsx"
    extensions = (".xlsx", ".xlsm", ".xltx")

    def parse(self, path: Path, *, file_id: str) -> ParsedDoc:
        from openpyxl import load_workbook

        doc = ParsedDoc(file_id=file_id, file_name=path.name, kind=self.kind)
        doc.findings.extend(_office_metadata(path, doc))

        wb = load_workbook(path, data_only=True)
        sheets: list[dict[str, Any]] = []
        order = 0

        for ws in wb.worksheets:
            grid, comments = _read_sheet(ws)
            if not grid:
                doc.findings.append(Finding(
                    "empty_sheet", f"工作表「{ws.title}」是空的，已跳过",
                    {"sheet": ws.title}))
                continue

            h = detect_header_row(grid)
            width = max(len(r) for r in grid)
            header = (dedupe_headers(grid[h]) if h >= 0
                      else [f"col{i + 1}" for i in range(width)])
            body = grid[h + 1 :]
            cols = {
                name: profile_column(name, [r[i] if i < len(r) else "" for r in body])
                for i, name in enumerate(header)
            }
            sheets.append({
                "name": ws.title, "header_row": h + 1, "columns": header,
                "rows": len(body), "profile": cols,
            })
            if h > 0:
                doc.findings.append(Finding(
                    "header_offset",
                    f"「{ws.title}」表头在第 {h + 1} 行，不是第 1 行",
                    {"sheet": ws.title, "row": h + 1}))
            elif h < 0:
                doc.findings.append(Finding(
                    "no_header",
                    f"「{ws.title}」没有可识别的表头，已按位置列名处理、正文从第 1 行起",
                    {"sheet": ws.title}))

            # 每行一个切片：行是表格里语义完整的最小单元。render 用补全了分组列的
            # 视图（切片才能自证归属），raw 保持原样（下游 shape 靠留空判分组列）。
            render_body = _fill_hierarchy(header, body)
            for ri, row in enumerate(body):
                excel_row = h + 2 + ri
                row_comments = {c0: t for (r0, c0), t in comments.items()
                                if r0 == excel_row}
                pairs = _render_pairs(header, render_body[ri], row_comments)
                if not pairs:
                    continue
                render = " | ".join(pairs)
                doc.chunks.append(make_chunk(
                    doc_id=f"{ws.title}:r{excel_row}", file_id=file_id, file_name=path.name,
                    locator={"kind": "range", "sheet": ws.title,
                             "rows": [excel_row, excel_row]},
                    render=render,
                    raw={header[i]: row[i] for i in range(min(len(header), len(row)))},
                    order=order, tags=["row"],
                ))
                order += 1

        doc.structured = {"sheets": sheets}
        return doc


def _read_sheet(ws: Any) -> tuple[list[list[str]], dict[tuple[int, int], str]]:
    """读成字符串网格，顺带还原合并单元格、抽出批注（按 (行, 列下标) 定位）。"""
    grid: list[list[str]] = []
    for row in ws.iter_rows():
        grid.append(["" if c.value is None else str(c.value).strip() for c in row])

    # 合并区只有左上角有值 —— 填充回去，否则下游看到一堆空格
    for rng in ws.merged_cells.ranges:
        r0, c0, r1, c1 = rng.min_row - 1, rng.min_col - 1, rng.max_row - 1, rng.max_col - 1
        if r0 >= len(grid) or c0 >= len(grid[r0]):
            continue
        v = grid[r0][c0]
        if not v:
            continue
        for r in range(r0, min(r1 + 1, len(grid))):
            for c in range(c0, min(c1 + 1, len(grid[r]))):
                if not grid[r][c]:
                    grid[r][c] = v

    comments: dict[tuple[int, int], str] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.comment and c.comment.text:
                txt = c.comment.text.strip()
                key = (c.row, c.column - 1)  # 0-based 列下标，与 header 对齐
                comments[key] = f"{comments[key]}；{txt}" if key in comments else txt

    while grid and not any(x.strip() for x in grid[-1]):
        grid.pop()
    return grid, comments


def _office_metadata(path: Path, doc: ParsedDoc) -> list[Finding]:
    """抽 Office 文档元数据。

    ``workbook.xml`` 里的 ``absPath`` 保存着这份文件最后一次保存时的绝对路径，
    ``docProps/core.xml`` 里有作者名。两者都不在表格内容里，但常常泄漏客户名、
    项目代号和人名 —— 既是理解材料来源的线索，也是对外发布前必须清的东西。
    """
    out: list[Finding] = []
    meta: dict[str, Any] = {}
    try:
        with zipfile.ZipFile(path) as z:
            names = set(z.namelist())
            if "xl/workbook.xml" in names:
                xml = z.read("xl/workbook.xml").decode("utf-8", "ignore")
                if m := re.search(r'absPath[^>]*url="([^"]+)"', xml):
                    meta["abs_path"] = m.group(1)
            if "docProps/core.xml" in names:
                xml = z.read("docProps/core.xml").decode("utf-8", "ignore")
                for tag, key in (("dc:creator", "creator"),
                                 ("cp:lastModifiedBy", "last_modified_by"),
                                 ("dcterms:modified", "modified")):
                    if m := re.search(rf"<{tag}[^>]*>([^<]+)</{tag}>", xml):
                        meta[key] = m.group(1)
    except (zipfile.BadZipFile, KeyError, OSError):
        return out

    doc.meta.update(meta)
    leaked = {k: v for k, v in meta.items()
              if k in ("abs_path", "creator", "last_modified_by") and v}
    if leaked:
        detail = "；".join(f"{k}={v}" for k, v in leaked.items())
        out.append(Finding(
            "metadata_leak",
            f"文档元数据里带着作者与保存路径（{detail}）。这些不在表格内容里，"
            "但对外发布前建议清理，或把项目名换成中性代号。",
            {"kind": "meta", "field": "workbook.xml/docProps"}, severity="warn"))
        # 元数据本身也是可检索的证据 —— FDE 问"你怎么知道客户是谁"时要答得出来
        doc.chunks.append(make_chunk(
            doc_id="meta", file_id=doc.file_id, file_name=path.name,
            locator={"kind": "meta", "field": "workbook.xml/absPath"},
            render=f"〔文档元数据〕{detail}", raw=meta, order=-1, tags=["meta"]))
    return out


# ══════════════════════════════════════════════════════════════════
#  csv
# ══════════════════════════════════════════════════════════════════
class CsvParser(Parser):
    kind = "csv"
    extensions = (".csv", ".tsv")

    def parse(self, path: Path, *, file_id: str) -> ParsedDoc:
        doc = ParsedDoc(file_id=file_id, file_name=path.name, kind=self.kind)
        text, enc = _read_text(path)
        if enc != "utf-8":
            doc.findings.append(Finding(
                "encoding_guess", f"文件不是 UTF-8，按 {enc} 解码（中文 CSV 常见 GBK）",
                {}, severity="warn"))

        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            delim = dialect.delimiter
        except csv.Error:
            delim = "\t" if path.suffix.lower() == ".tsv" else ","
        rows = [list(r) for r in csv.reader(text.splitlines(), delimiter=delim)]
        rows = [r for r in rows if any(str(x).strip() for x in r)]
        if not rows:
            doc.structured = {"columns": [], "rows": 0, "profile": {}}
            return doc

        h = detect_header_row([[str(x) for x in r] for r in rows])
        header = (dedupe_headers([str(x) for x in rows[h]]) if h >= 0
                  else [f"col{i + 1}" for i in range(max(len(r) for r in rows))])
        body = rows[h + 1 :]
        profile = {
            name: profile_column(name, [str(r[i]) if i < len(r) else "" for r in body])
            for i, name in enumerate(header)
        }
        doc.structured = {"columns": header, "rows": len(body), "delimiter": delim,
                          "profile": profile}

        # 大表不逐行切片 —— 几万行会把索引撑爆且毫无检索价值。
        # 给表头 + 画像 + 少量样本行，这才是抽取需要看的东西。
        doc.chunks.append(make_chunk(
            doc_id="schema", file_id=file_id, file_name=path.name,
            locator={"kind": "range", "sheet": path.stem, "rows": [h + 1, h + 1]},
            render="列：" + " | ".join(
                f"{n}({p['inferred_type']}，唯一率{p['distinct_ratio']:.0%}，"
                f"空值率{p['null_rate']:.0%})" for n, p in profile.items()),
            raw=profile, order=0, tags=["schema"]))
        for ri, row in enumerate(body[:20]):
            doc.chunks.append(make_chunk(
                doc_id=f"r{h + 2 + ri}", file_id=file_id, file_name=path.name,
                locator={"kind": "range", "sheet": path.stem,
                         "rows": [h + 2 + ri, h + 2 + ri]},
                render=" | ".join(_render_pairs(header, row)),
                order=ri + 1, tags=["sample"]))
        if len(body) > 20:
            doc.findings.append(Finding(
                "sampled", f"共 {len(body)} 行，索引里只放了前 20 行样本 + 全量列画像",
                {}))
        return doc


def _read_text(path: Path) -> tuple[str, str]:
    """按常见编码依次尝试。中文 CSV 里 GBK 极常见，直接用 utf-8 会炸。"""
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030", "big5", "latin-1"):
        try:
            return raw.decode(enc), enc.replace("-sig", "")
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace"), "utf-8(replace)"
