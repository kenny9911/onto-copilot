"""PowerPoint ``.pptx`` -> 按页的正文、表格、备注与可追溯证据。

优先使用 ``python-pptx`` 读取；依赖不可用或遇到它尚不支持的包时，退回到受限的
OOXML ZIP/XML 读取器。降级是显式 finding，不会把“只读到一部分”伪装成成功。
"""

from __future__ import annotations

import posixpath
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from .base import Finding, ParsedDoc, Parser, make_chunk
from .text import _RULE_HINTS

_MAX_MEMBER_BYTES = 20 * 1024 * 1024
_MAX_TOTAL_BYTES = 100 * 1024 * 1024
_SLIDE_RE = re.compile(r"^ppt/slides/slide(\d+)\.xml$")


class PptxParser(Parser):
    """提取每页文本框、表格行、演讲者备注和 Office 元数据。"""

    kind = "pptx"
    extensions = (".pptx", ".pptm", ".ppsx")

    def parse(self, path: Path, *, file_id: str) -> ParsedDoc:
        try:
            import pptx  # type: ignore[import-not-found]
        except ImportError:
            doc = _parse_ooxml(path, file_id=file_id)
            doc.findings.insert(0, Finding(
                "parser_fallback",
                "未安装 python-pptx，已使用内置 OOXML 兜底解析正文、表格和备注；"
                "复杂 SmartArt/图表不会被当作已解析内容。",
                {"kind": "package", "path": "/"}))
            return doc

        try:
            return _parse_with_python_pptx(path, file_id=file_id, pptx_module=pptx)
        except Exception as exc:  # noqa: BLE001 - 第三方库可能拒绝合法但少见的 OOXML
            doc = _parse_ooxml(path, file_id=file_id)
            doc.findings.insert(0, Finding(
                "parser_fallback",
                f"python-pptx 解析失败（{type(exc).__name__}: {exc}），"
                "已改用内置 OOXML 兜底解析。",
                {"kind": "package", "path": "/"}, severity="warn"))
            return doc


def _parse_with_python_pptx(path: Path, *, file_id: str, pptx_module: Any) -> ParsedDoc:
    doc = ParsedDoc(file_id=file_id, file_name=path.name, kind="pptx")
    presentation = pptx_module.Presentation(str(path))
    _pptx_metadata(doc, presentation.core_properties)

    slides: list[dict[str, Any]] = []
    order = 0
    for page, slide in enumerate(presentation.slides, start=1):
        slide_data: dict[str, Any] = {
            "number": page, "title": "", "texts": [], "tables": [], "notes": "",
        }
        title_shape = getattr(slide.shapes, "title", None)
        slide_data["title"] = _clean_text(getattr(title_shape, "text", ""))

        table_no = 0
        for shape_no, shape in enumerate(slide.shapes, start=1):
            shape_id = str(getattr(shape, "shape_id", shape_no))
            shape_name = str(getattr(shape, "name", f"shape-{shape_no}"))
            bbox = [
                int(getattr(shape, "left", 0)), int(getattr(shape, "top", 0)),
                int(getattr(shape, "width", 0)), int(getattr(shape, "height", 0)),
            ]
            if bool(getattr(shape, "has_table", False)):
                table_no += 1
                rows = [[_clean_text(cell.text) for cell in row.cells]
                        for row in shape.table.rows]
                table_data = _table_data(table_no, rows)
                slide_data["tables"].append(table_data)
                order = _emit_table(
                    doc, file_id=file_id, file_name=path.name, page=page,
                    table=table_data, bbox=bbox, shape_id=shape_id, order=order)
                continue

            if not bool(getattr(shape, "has_text_frame", False)):
                continue
            text = _clean_text(getattr(shape, "text", ""))
            if not text:
                continue
            shape_data = {"id": shape_id, "name": shape_name, "text": text, "bbox": bbox}
            slide_data["texts"].append(shape_data)
            doc.chunks.append(make_chunk(
                doc_id=f"slide:{page}:shape:{shape_id}", file_id=file_id,
                file_name=path.name,
                locator={"kind": "page", "page": page, "bbox": bbox,
                         "shape_id": shape_id, "shape": shape_name},
                render=_slide_render(page, slide_data["title"], text), raw=shape_data,
                order=order, tags=_text_tags("slide_text", text)))
            order += 1

        notes = ""
        try:
            notes = _clean_text(slide.notes_slide.notes_text_frame.text)
        except (AttributeError, KeyError, ValueError):
            # 没有 notes part 是正常情况；不能因此把整页解析标红。
            notes = ""
        if notes:
            slide_data["notes"] = notes
            doc.chunks.append(make_chunk(
                doc_id=f"slide:{page}:notes", file_id=file_id, file_name=path.name,
                locator={"kind": "page", "page": page, "bbox": [0, 0, 0, 0],
                         "notes": True},
                render=f"PPT 第 {page} 页演讲者备注：{notes}", raw={"notes": notes},
                order=order, tags=_text_tags("speaker_notes", notes)))
            order += 1

        slides.append(slide_data)

    doc.structured = {"slides": slides, "slide_count": len(slides)}
    if not slides:
        doc.findings.append(Finding(
            "empty_presentation", "PPTX 中没有幻灯片", {"kind": "package", "path": "/"},
            severity="warn"))
    return doc


def _parse_ooxml(path: Path, *, file_id: str) -> ParsedDoc:
    doc = ParsedDoc(file_id=file_id, file_name=path.name, kind="pptx")
    try:
        archive = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile) as exc:
        doc.findings.append(Finding(
            "parse_failed", f"不是可读取的 PPTX OOXML 包：{exc}",
            {"kind": "package", "path": "/"}, severity="warn"))
        return doc

    with archive:
        infos = archive.infolist()
        if (any(info.file_size > _MAX_MEMBER_BYTES for info in infos)
                or sum(info.file_size for info in infos) > _MAX_TOTAL_BYTES):
            doc.findings.append(Finding(
                "file_too_large", "PPTX 解压后体积超过安全上限，已拒绝解析",
                {"kind": "package", "path": "/"}, severity="warn"))
            return doc

        names = {info.filename for info in infos}
        _ooxml_metadata(doc, archive, names)
        slide_parts = _ordered_slide_parts(archive, names)
        slides: list[dict[str, Any]] = []
        order = 0
        for page, part in enumerate(slide_parts, start=1):
            root = _read_xml(archive, part, doc)
            if root is None:
                continue
            slide_data: dict[str, Any] = {
                "number": page, "title": "", "texts": [], "tables": [], "notes": "",
            }
            title_candidates: list[str] = []
            for shape_no, shape in enumerate(_descendants(root, "sp"), start=1):
                texts = [_clean_text("".join(node.itertext()))
                         for node in _descendants(shape, "t")]
                text = _clean_text("\n".join(t for t in texts if t))
                if not text:
                    continue
                shape_id, shape_name = _shape_identity(shape, shape_no)
                bbox = _shape_bbox(shape)
                placeholder = _first_descendant(shape, "ph")
                placeholder_type = placeholder.attrib.get("type", "") if placeholder is not None else ""
                if placeholder_type in {"title", "ctrTitle"}:
                    title_candidates.append(text)
                shape_data = {"id": shape_id, "name": shape_name, "text": text, "bbox": bbox}
                slide_data["texts"].append(shape_data)

            slide_data["title"] = title_candidates[0] if title_candidates else (
                slide_data["texts"][0]["text"] if slide_data["texts"] else "")
            for shape_data in slide_data["texts"]:
                doc.chunks.append(make_chunk(
                    doc_id=f"slide:{page}:shape:{shape_data['id']}", file_id=file_id,
                    file_name=path.name,
                    locator={"kind": "page", "page": page, "bbox": shape_data["bbox"],
                             "shape_id": shape_data["id"], "shape": shape_data["name"]},
                    render=_slide_render(page, slide_data["title"], shape_data["text"]),
                    raw=shape_data, order=order,
                    tags=_text_tags("slide_text", shape_data["text"])))
                order += 1

            for table_no, table in enumerate(_descendants(root, "tbl"), start=1):
                rows: list[list[str]] = []
                for row in [child for child in table if _local(child.tag) == "tr"]:
                    cells: list[str] = []
                    for cell in [child for child in row if _local(child.tag) == "tc"]:
                        cells.append(_clean_text(" ".join(
                            "".join(t.itertext()) for t in _descendants(cell, "t"))))
                    rows.append(cells)
                table_data = _table_data(table_no, rows)
                slide_data["tables"].append(table_data)
                order = _emit_table(
                    doc, file_id=file_id, file_name=path.name, page=page,
                    table=table_data, bbox=[0, 0, 0, 0], shape_id=f"table-{table_no}",
                    order=order)

            notes_part = _notes_part(archive, names, part)
            notes_root = _read_xml(archive, notes_part, doc) if notes_part else None
            if notes_root is not None:
                notes = _notes_text(notes_root)
                if notes:
                    slide_data["notes"] = notes
                    doc.chunks.append(make_chunk(
                        doc_id=f"slide:{page}:notes", file_id=file_id, file_name=path.name,
                        locator={"kind": "page", "page": page, "bbox": [0, 0, 0, 0],
                                 "notes": True, "part": notes_part},
                        render=f"PPT 第 {page} 页演讲者备注：{notes}", raw={"notes": notes},
                        order=order, tags=_text_tags("speaker_notes", notes)))
                    order += 1
            slides.append(slide_data)

        doc.structured = {"slides": slides, "slide_count": len(slides)}
        if not slide_parts:
            doc.findings.append(Finding(
                "empty_presentation", "PPTX 包中没有 slide part",
                {"kind": "package", "path": "/ppt/slides"}, severity="warn"))
    return doc


def _pptx_metadata(doc: ParsedDoc, core: Any) -> None:
    meta = {key: value for key, value in (
        ("creator", getattr(core, "author", None)),
        ("last_modified_by", getattr(core, "last_modified_by", None)),
        ("modified", str(getattr(core, "modified", "") or "")),
        ("title", getattr(core, "title", None)),
    ) if value}
    doc.meta.update(meta)
    _metadata_finding(doc, meta)


def _ooxml_metadata(doc: ParsedDoc, archive: zipfile.ZipFile, names: set[str]) -> None:
    part = "docProps/core.xml"
    if part not in names:
        return
    root = _read_xml(archive, part, doc)
    if root is None:
        return
    wanted = {"creator", "lastModifiedBy", "modified", "title"}
    raw = {_local(node.tag): _clean_text("".join(node.itertext()))
           for node in root.iter() if _local(node.tag) in wanted}
    meta = {
        "creator": raw.get("creator", ""),
        "last_modified_by": raw.get("lastModifiedBy", ""),
        "modified": raw.get("modified", ""),
        "title": raw.get("title", ""),
    }
    doc.meta.update({key: value for key, value in meta.items() if value})
    _metadata_finding(doc, meta)


def _metadata_finding(doc: ParsedDoc, meta: dict[str, Any]) -> None:
    leaked = {key: value for key, value in meta.items()
              if key in {"creator", "last_modified_by"} and value}
    if leaked:
        doc.findings.append(Finding(
            "metadata_leak",
            "演示文稿元数据里带着作者信息（"
            + "；".join(f"{key}={value}" for key, value in leaked.items())
            + "），对外发布前建议清理。",
            {"kind": "meta", "field": "docProps/core.xml"}, severity="warn"))


def _emit_table(
    doc: ParsedDoc, *, file_id: str, file_name: str, page: int,
    table: dict[str, Any], bbox: list[int], shape_id: str, order: int,
) -> int:
    columns = table["columns"]
    body = table["data"]
    if columns:
        doc.chunks.append(make_chunk(
            doc_id=f"slide:{page}:table:{table['index']}:schema", file_id=file_id,
            file_name=file_name,
            locator={"kind": "page", "page": page, "bbox": bbox,
                     "shape_id": shape_id, "table": table["index"], "row": 1},
            render=f"PPT 第 {page} 页表格 {table['index']} 列："
                   + "、".join(value for value in columns if value),
            raw={"columns": columns}, order=order, tags=["pptx", "table", "schema"]))
        order += 1
    for row_no, row in enumerate(body, start=2):
        pairs = [f"{columns[index] if index < len(columns) and columns[index] else f'C{index + 1}'}"
                 f"={value}" for index, value in enumerate(row) if value]
        if not pairs:
            continue
        render = f"PPT 第 {page} 页表格 {table['index']} 第 {row_no} 行：" + " | ".join(pairs)
        doc.chunks.append(make_chunk(
            doc_id=f"slide:{page}:table:{table['index']}:row:{row_no}", file_id=file_id,
            file_name=file_name,
            locator={"kind": "page", "page": page, "bbox": bbox,
                     "shape_id": shape_id, "table": table["index"], "row": row_no},
            render=render, raw={"row": row, "columns": columns}, order=order,
            tags=_text_tags("table", render)))
        order += 1
    return order


def _table_data(index: int, rows: list[list[str]]) -> dict[str, Any]:
    clean_rows = [[_clean_text(value) for value in row] for row in rows]
    clean_rows = [row for row in clean_rows if any(row)]
    columns = clean_rows[0] if clean_rows else []
    return {
        "index": index,
        "columns": columns,
        "rows": max(0, len(clean_rows) - 1),
        "data": clean_rows[1:] if clean_rows else [],
    }


def _slide_render(page: int, title: str, text: str) -> str:
    heading = f"〔{title}〕\n" if title and title != text else ""
    return f"PPT 第 {page} 页\n{heading}{text}"


def _text_tags(kind: str, text: str) -> list[str]:
    return ["pptx", kind] + (["rule"] if _RULE_HINTS.search(text) else [])


def _read_xml(
    archive: zipfile.ZipFile, part: str, doc: ParsedDoc,
) -> ET.Element | None:
    try:
        data = archive.read(part)
        upper = data[: min(len(data), 64 * 1024)].upper()
        if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
            raise ValueError("含 DTD/ENTITY 声明")
        return ET.fromstring(data)
    except (KeyError, ET.ParseError, OSError, ValueError) as exc:
        doc.findings.append(Finding(
            "unparsed_part", f"无法解析 OOXML part {part}：{exc}",
            {"kind": "package", "path": f"/{part}"}, severity="warn"))
        return None


def _notes_part(archive: zipfile.ZipFile, names: set[str], slide_part: str) -> str:
    folder, name = posixpath.split(slide_part)
    rels_part = posixpath.join(folder, "_rels", name + ".rels")
    if rels_part in names:
        try:
            root = ET.fromstring(archive.read(rels_part))
            for rel in root:
                if rel.attrib.get("Type", "").endswith("/notesSlide"):
                    target = rel.attrib.get("Target", "")
                    resolved = posixpath.normpath(posixpath.join(folder, target))
                    if resolved in names:
                        return resolved
        except (ET.ParseError, OSError, KeyError):
            pass
    match = _SLIDE_RE.match(slide_part)
    candidate = f"ppt/notesSlides/notesSlide{match.group(1)}.xml" if match else ""
    return candidate if candidate in names else ""


def _ordered_slide_parts(archive: zipfile.ZipFile, names: set[str]) -> list[str]:
    """按演示文稿关系顺序返回 slide part，而不是猜 ``slideN.xml`` 的 N。

    OOXML 允许 ``slide2.xml`` 排在第一页；人看的页码必须跟 ``p:sldIdLst`` 一致。
    缺少关系文件的非标准/极简包才退回按数字文件名排序。
    """
    presentation = "ppt/presentation.xml"
    relationships = "ppt/_rels/presentation.xml.rels"
    if presentation in names and relationships in names:
        try:
            presentation_root = _safe_fromstring(archive.read(presentation))
            relationships_root = _safe_fromstring(archive.read(relationships))
            targets = {
                rel.attrib.get("Id", ""): rel.attrib.get("Target", "")
                for rel in relationships_root
                if rel.attrib.get("Type", "").endswith("/slide")
            }
            ordered: list[str] = []
            for slide_id in _descendants(presentation_root, "sldId"):
                relation_id = next(
                    (value for key, value in slide_id.attrib.items()
                     if key.startswith("{") and _local(key) == "id"),
                    "",
                )
                target = targets.get(relation_id, "")
                part = posixpath.normpath(posixpath.join("ppt", target))
                if part in names:
                    ordered.append(part)
            if ordered:
                return ordered
        except (ET.ParseError, OSError, KeyError, ValueError):
            pass
    numbered: list[tuple[int, str]] = []
    for name in names:
        if match := _SLIDE_RE.match(name):
            numbered.append((int(match.group(1)), name))
    return [name for _, name in sorted(numbered, key=lambda item: item[0])]


def _notes_text(root: ET.Element) -> str:
    texts: list[str] = []
    for shape in _descendants(root, "sp"):
        placeholder = _first_descendant(shape, "ph")
        placeholder_type = placeholder.attrib.get("type", "") if placeholder is not None else ""
        if placeholder_type in {"sldImg", "sldNum", "hdr", "ftr", "dt"}:
            continue
        text = _clean_text("\n".join(
            "".join(node.itertext()) for node in _descendants(shape, "t")))
        if text:
            texts.append(text)
    return _clean_text("\n".join(texts))


def _shape_identity(shape: ET.Element, fallback: int) -> tuple[str, str]:
    c_nv_pr = _first_descendant(shape, "cNvPr")
    if c_nv_pr is None:
        return str(fallback), f"shape-{fallback}"
    return c_nv_pr.attrib.get("id", str(fallback)), c_nv_pr.attrib.get("name", f"shape-{fallback}")


def _shape_bbox(shape: ET.Element) -> list[int]:
    xfrm = _first_descendant(shape, "xfrm")
    if xfrm is None:
        return [0, 0, 0, 0]
    off = next((child for child in xfrm if _local(child.tag) == "off"), None)
    ext = next((child for child in xfrm if _local(child.tag) == "ext"), None)
    return [
        int((off.attrib if off is not None else {}).get("x", 0)),
        int((off.attrib if off is not None else {}).get("y", 0)),
        int((ext.attrib if ext is not None else {}).get("cx", 0)),
        int((ext.attrib if ext is not None else {}).get("cy", 0)),
    ]


def _descendants(element: ET.Element, local_name: str) -> list[ET.Element]:
    return [node for node in element.iter() if _local(node.tag) == local_name]


def _first_descendant(element: ET.Element, local_name: str) -> ET.Element | None:
    return next((node for node in element.iter() if _local(node.tag) == local_name), None)


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1].rsplit(":", 1)[-1]


def _safe_fromstring(data: bytes) -> ET.Element:
    upper = data[: min(len(data), 64 * 1024)].upper()
    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
        raise ValueError("含 DTD/ENTITY 声明")
    return ET.fromstring(data)


def _clean_text(value: Any) -> str:
    lines = [" ".join(line.split()) for line in str(value or "").splitlines()]
    return "\n".join(line for line in lines if line).strip()
