"""把 Python 侧的行为导成 golden fixtures，给 TS 迁移当安全网。

那 1182 个 pytest 用例不会跟着迁移（它们断言的是 Python 结构），但 golden 会：
TS 侧每写完一个模块，就拿**同样的输入**去比对**同样的输出**。字节级一致的模块
（ids 的 canonical_json / fingerprint）尤其重要 —— 盘点确认过 Python 与 JS 的
JSON 序列化在数字上不等价（1.0 → "1.0" vs "1"），rid 漂一个字节，重放和跨版本
diff 就全断了。

跑法::

    .venv/bin/python tools/export_golden.py

产物落在 golden/ 下，按模块分文件。**这个目录进 git** —— 它是迁移期间两边共同的
真相，不是临时产物。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

OUT = Path(__file__).resolve().parent.parent / "golden"


def w(name: str, obj: object) -> None:
    OUT.mkdir(exist_ok=True)
    p = OUT / name
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"  {name:36} {p.stat().st_size:>8} B")


# ══════════════════════════════════════════════════════════════════
#  1. ids —— 必须字节级一致
# ══════════════════════════════════════════════════════════════════
def export_ids() -> None:
    from ontocopilot.kernel.ids import (
        canonical_json,
        child_id,
        content_ref,
        fingerprint,
        rid,
        sha256_hex,
        slug,
    )

    # 覆盖盘点点名的雷区：整数/浮点/负零/大整数、None/bool、嵌套排序、
    # CJK、代理对 emoji、混合键、空容器
    objs = [
        {"b": 1, "a": 2},
        {"n": 1.0, "m": 1, "z": -0.0, "big": 2**60},
        {"s": "中文键", "emoji": "🐍🦀", "empty": [], "null": None, "t": True},
        [3, 1, 2, {"nested": {"y": [1.5, "1.5"], "x": None}}],
        {"金额": 10.25, "税率": 0.13, "备注": "含税；口径=年度累计"},
        "纯字符串",
        [],
        {},
    ]
    slugs = ["采购计划管理", "PO Header 采购订单头", "  spaces  ", "ALL-CAPS_and.dots",
             "日本語テスト", "emoji🐍mixed", "a" * 80, ""]
    w("ids.json", {
        "canonical_json": [{"in": o, "out": canonical_json(o)} for o in objs],
        "fingerprint": [{"in": o, "out": fingerprint(o)} for o in objs],
        "sha256_hex": [{"in": s, "out": sha256_hex(s)}
                       for s in ["", "abc", "中文", "🐍"]],
        "slug": [{"in": s, "out": slug(s)} for s in slugs],
        "rid": [{"kind": k, "name": n, "out": rid(k, n)}
                for k, n in [("ot", "采购订单头"), ("pt", "poHeader_金额_3"),
                             ("at", "createPbp"), ("oq", "slot:数值未定_「超出XX金额」")]],
        "content_ref": [{"in": s, "out": content_ref(s)} for s in ["abc", "中文内容"]],
        "child_id": [{"parent": "EXTRACT", "parts": ["s0", 1],
                      "out": child_id("EXTRACT", "s0", 1)}],
    })


# ══════════════════════════════════════════════════════════════════
#  2. 事件编解码 —— jsonl 的字节形状
# ══════════════════════════════════════════════════════════════════
def export_events() -> None:
    from ontocopilot.kernel.events import Event, EventKind

    samples = [
        Event(run_id="r1", seq=0, kind=EventKind.RUN_STARTED, ts_ms=1723500000000,
              payload={"dag": "onto_extract"}),
        Event(run_id="r1", seq=1, kind=EventKind.NODE_ENTERED, ts_ms=1723500000100,
              node_id="EXTRACT.s0", payload={"mode": "plan_execute", "attempt": 1}),
        Event(run_id="r1", seq=2, kind=EventKind.EFFECT_COMPLETED, ts_ms=1723500000200,
              node_id="EXTRACT.s0", ref="fx_abc123", payload={}),
        Event(run_id="r1", seq=3, kind=EventKind.THOUGHT, ts_ms=1723500000300,
              node_id="EXTRACT.s0",
              payload={"text": "中文思考内容，带「引号」和 emoji 🐍"}),
    ]
    w("events.json", {
        "kinds": [str(k) for k in EventKind],
        "roundtrip": [{"dict": e.to_dict(),
                       "line": json.dumps(e.to_dict(), ensure_ascii=False,
                                          separators=(",", ":"), sort_keys=True)}
                      for e in samples],
    })


# ══════════════════════════════════════════════════════════════════
#  3. 预算降级阶梯
# ══════════════════════════════════════════════════════════════════
def export_budget() -> None:
    from ontocopilot.kernel.budget import Budget

    rows = []
    for frac in (0.0, 0.5, 0.61, 0.76, 0.86, 0.9601, 1.0):
        b = Budget(tokens=1000, usd=10.0)
        b.spend(tokens=1000 * frac)
        rows.append({"spent_frac": frac, "level": int(b.level),
                     "critic_rounds_of_2": b.critic_rounds(2),
                     "allow_self_consistency": b.allow_self_consistency(),
                     "allow_llm_critic": b.allow_llm_critic(),
                     "must_halt": b.must_halt()})
    # 单调性：level 只升不降 —— 花钱后又"退款"也不许降级回去
    b = Budget(tokens=1000)
    b.spend(tokens=900)
    lvl_high = int(b.level)
    b._spent["tokens"] = 100  # 人为回退，level 必须保持
    rows.append({"case": "latch", "after_spend_900": lvl_high,
                 "after_rollback_to_100": int(b.level)})
    w("budget.json", {"ladder": rows})


# ══════════════════════════════════════════════════════════════════
#  4. 形状推断 + 规则抽取 —— 领域层的核心判据
# ══════════════════════════════════════════════════════════════════
def export_shape() -> None:
    from ontocopilot.onto.shape import infer_shape, split_options, structural_extract

    registry = [
        {"应用模块": "采购计划管理", "业务对象": "采购需求计划",
         "实体编码": "pbpHeader", "实体名称": "采购业务计划头"},
        {"应用模块": "采购计划管理", "业务对象": "", "实体编码": "pbpLine",
         "实体名称": "采购业务计划行"},
    ] + [{"应用模块": "供应商关系管理", "业务对象": "", "实体编码": f"clmDoc{i}",
          "实体名称": f"合同附件{i}"} for i in range(12)]
    actions = [
        {"应用模块": "采购计划管理", "业务对象": "采购需求计划",
         "实体编码": "createPbp", "实体名称": "创建PBP",
         "url": "/msourcing/openapi/v1/createPbp"},
    ] + [{"应用模块": "采购计划管理", "业务对象": "", "实体编码": f"queryPbp{i}",
          "实体名称": f"查询PBP{i}", "url": f"/msourcing/openapi/v1/queryPbp{i}"}
         for i in range(11)]
    types = ["varchar(64)", "decimal(18,2)", "date", "int", "varchar(32)", "timestamp"]
    fields = [{"所属对象": "pbpHeader", "字段名": f"field{i}", "类型": types[i % 6],
               "口径": f"第 {i} 个字段的口径说明", "必填": "是" if i % 2 else "否"}
              for i in range(12)]
    survey = [{"节点": "（1）编制集采计划" if i == 0 else "", "编号": str(i + 1),
               "澄清问题": "集采计划是在系统里编的，还是线下编好只录结果？",
               "参考选项": "① 全线下 ② 系统里编 ③ 线下编、系统里审",
               "答复": ""} for i in range(8)]

    def cites(rows):
        return [f"x.xlsx!Sheet1!R{i + 2}-{i + 2}" for i in range(len(rows))]

    out = {}
    for name, rows in [("registry", registry), ("actions", actions),
                       ("fields", fields), ("survey", survey)]:
        shape = infer_shape(rows)
        out[name] = {
            "shape": shape.to_dict(),
            "extract": structural_extract(rows, cites(rows), shape),
        }
    out["split_options"] = [
        {"in": s, "out": split_options(s)}
        for s in ["① 全线下 ② 系统里编 ③ 线下编、系统里审",
                  "1. 固定周期 2. 到期前 3. 人工发起",
                  "一句不该被拆开的完整的话"]]
    w("shape.json", out)


# ══════════════════════════════════════════════════════════════════
#  5. 端到端 —— 确定性半边的全部产物
# ══════════════════════════════════════════════════════════════════
#: 合成材料的三段流程说明 —— 形态照抄真实材料：合并单元格里塞多个五段式节点、
#: 错别字「执行着」、并列产物、占位符、取值清单，一样不少。
_PROCESS = """（1）创建采购需求计划：基于已审批集采计划自动注入；
触发条件：集采计划审批完成；
输入：已审批集采计划；
输出：未审批采购需求计划；
执行者：需求申请人
（2）审批采购需求计划：对已编制的采购需求计划进行审批；
触发条件：采购需求计划已编制；
输入：采购需求计划；
输出：已审批采购需求计划；
执行着：部门领导
（3）取消采购需求计划：执行采购需求计划取消流程；
触发条件：无需继续采购；
输入：需求取消；
输出：已取消采购需求计划；
执行者：需求申请人"""

_RULES_TEXT = (
    "业务规则：\n"
    "1、需进行采购包总金额上线控制，如超出XX金额，则采购包创建失败；\n"
    "2、实际与计划基线对比，如提前XX天需要进行预警提醒；\n"
    "3、进度状态标准：未开始、执行中、部分完成、已完成、已暂停、已取消；\n"
)


def _make_material() -> Path:
    """合成一份全形态材料，落进 golden/。它既是 TS 侧的输入，也是行为的锚点。

    现成的 materials/实体梳理.xlsx 太小（2 个对象、流程为空），护不住任何东西；
    真实的 326 行材料随会话工作区被清理了。合成的这份是**确定性**的 —— 重跑
    导出器字节不变，golden diff 才有意义。
    """
    from openpyxl import Workbook

    OUT.mkdir(exist_ok=True)
    p = OUT / "材料.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("业务对象实体梳理")
    ws.append(["应用模块", "业务对象", "实体编码", "实体名称"])
    ents = [("采购计划管理", "采购需求计划", "pbpHeader", "采购业务计划头"),
            ("", "", "pbpLine", "采购业务计划行"),
            ("", "", "pbpRel", "采购业务计划头行关系"),
            ("采购订单管理", "采购订单", "poHeader", "采购订单头"),
            ("", "", "poLine", "采购订单行"),
            ("", "", "poTaxLine", "采购订单税行"),
            ("", "", "poActionLog", "采购订单操作日志")]
    for r in ents:
        ws.append(list(r))
    ws = wb.create_sheet("业务对象API梳理-行动")
    ws.append(["应用模块", "业务对象", "实体编码", "实体名称", "url"])
    apis = [("采购计划管理", "采购需求计划", "createPbp", "创建PBP", "/v1/createPbp"),
            ("", "", "cancelOpenPbp", "取消PBP", "/v1/cancelPbp"),
            ("", "", "queryOpenPbpHeader", "查询PBP头", "/v1/queryPbpHeader"),
            ("采购订单管理", "采购订单", "createPo", "创建采购订单", "/v1/createPo"),
            ("", "", "updatePo", "修改采购订单", "/v1/updatePo"),
            ("", "", "queryPoHeader", "查询采购订单头", "/v1/queryPoHeader")]
    for r in apis:
        ws.append(list(r))
    ws = wb.create_sheet("业务规则")
    ws.append([_PROCESS])
    ws.append([_RULES_TEXT])
    wb.create_sheet("实体间关系-待梳理")     # 空表 —— empty_containers 的判据
    wb.save(p)
    _freeze_xlsx(p)
    return p


def _freeze_xlsx(p: Path) -> None:
    """把 xlsx 重写成字节确定的形态。

    openpyxl 有两层挡不住的 now()：``save()`` 内部会覆盖 ``properties.modified``
    （构造时钉的值直接被冲掉），zip 成员的 mtime 也取当前时间。而 ``file_id`` 是
    对**文件字节**内容寻址的 —— 字节漂一次，pipeline.oir.json 里每条 provenance
    的 file_id 跟着漂，golden diff 全是噪声。所以 save 之后整包重写：
    core.xml 里的时间戳替换成常量，每个成员的 date_time 钉在同一天。
    """
    import re
    import zipfile

    stamp = "2026-01-01T00:00:00Z"
    with zipfile.ZipFile(p) as zin:
        members = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    for idx, (name, data) in enumerate(members):
        if name == "docProps/core.xml":
            text = data.decode("utf-8")
            text = re.sub(r"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</)",
                          rf"\g<1>{stamp}\g<2>", text)
            members[idx] = (name, text.encode("utf-8"))
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in members:
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(info, data)


def export_pipeline() -> None:
    from ontocopilot.kernel.memory.evidence import EvidenceIndex
    from ontocopilot.onto.flow_extract import build_flow, looks_like_process, parse_steps
    from ontocopilot.onto.gaps import mine_questions
    from ontocopilot.onto.parse import default_registry
    from ontocopilot.onto.pipeline import build_oir, finish, segment_corpus
    from ontocopilot.onto.shape import structural_extract

    mat = _make_material()
    docs = [default_registry().parse(mat)]
    index = EvidenceIndex()
    for d in docs:
        for c in d.chunks:
            index.add(c)
    segs = segment_corpus(index, docs)
    merged: dict[str, list] = {k: [] for k in
                               ("objects", "properties", "links", "actions",
                                "rules", "questions")}
    for s in segs:
        rows, cs = s.rows(index)
        pre = structural_extract(rows, cs, s.shape, carry_in=s.carry_in)
        for k in merged:
            merged[k].extend(pre.get(k) or [])
    oir = build_oir(merged, index)
    for q in mine_questions(oir, docs=docs, chunks=index.all_chunks()):
        if q.rid not in oir.questions:
            oir.add_question(q)

    steps = []
    for d in docs:
        for c in d.chunks:
            vals = list(c.raw.values()) if isinstance(c.raw, dict) else []
            for v in [*vals, c.render]:
                text = str(v or "")
                if looks_like_process(text):
                    known = {x.no for x in steps}
                    steps += [x for x in parse_steps(text, cite=c.cite())
                              if x.no not in known]
    g = build_flow(sorted(steps, key=lambda x: x.no), file_name=mat.name)
    res = finish(oir, project="golden")
    spec = res["template_spec"]

    w("pipeline.segments.json", [
        {"key": s.key, "label": s.label, "chunks": len(s.chunk_ids),
         "shape": s.shape.to_dict(), "carry_in": s.carry_in} for s in segs])
    w("pipeline.oir.json", oir.to_dict())
    w("pipeline.flow.json", g.to_dict())
    w("pipeline.template.json", spec.to_dict())
    w("pipeline.stats.json", {"oir": oir.stats(), "flow": g.stats(),
                              "template": spec.stats()})


# ══════════════════════════════════════════════════════════════════
#  6. xlsx 逐格 dump —— 样式即语义那一层
# ══════════════════════════════════════════════════════════════════
def export_xlsx_cells() -> None:
    import tempfile

    from openpyxl import load_workbook

    from ontocopilot.onto.oir import OIR, ObjectType, OpenQuestion, inferred
    from ontocopilot.onto.template import compile_template, write_xlsx

    oir = OIR()
    oir.add_object(ObjectType(rid="ot_po", api_name=inferred("poHeader"),
                              display_name=inferred("采购订单头"),
                              description=inferred(""), primary_key=inferred([])))
    oir.add_question(OpenQuestion(
        rid="oq1", text=inferred("金额含不含税？"),
        options=["含税", "不含税"], group="口径", code="1", asked_by="customer"))
    spec = compile_template(oir)
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "t.xlsx"
        write_xlsx(spec, p, project="golden")
        wb = load_workbook(p)
        sheets = {}
        for ws in wb.worksheets:
            cells = []
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None and not c.comment:
                        continue
                    cells.append({
                        "ref": c.coordinate, "value": c.value,
                        "fill": c.fill.fgColor.rgb if c.fill and c.fill.fgColor else None,
                        "bold": bool(c.font and c.font.bold),
                        "comment": c.comment.text if c.comment else None,
                    })
            sheets[ws.title] = {
                "cells": cells,
                "hidden_cols": [k for k, v in ws.column_dimensions.items() if v.hidden],
                "freeze": ws.freeze_panes,
                "validations": [{"sqref": str(dv.sqref), "formula1": dv.formula1}
                                for dv in ws.data_validations.dataValidation],
            }
        w("xlsx.cells.json", sheets)


if __name__ == "__main__":
    print("导出 golden fixtures →", OUT)
    export_ids()
    export_events()
    export_budget()
    export_shape()
    export_pipeline()
    export_xlsx_cells()
    print("完成。")
