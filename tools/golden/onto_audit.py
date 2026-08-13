"""导出 onto/audit.py 的 golden 向量（供 ts/test/onto.audit.test.ts 断言）。

两条链分别钉：

1. **读回传件**（`read_returned` / `_locate_anchor` / `_looks_like_data`）——
   用 openpyxl 真造几份 xlsx（正常、顶部插行、左侧插列、锚点被删、重复行、
   说明页），把每份的 `(rid, 字段) → 值` 和损伤文案原样导出。xlsx 本身也落进
   golden/，TS 侧读同一份文件，这样才真正验到"两边读同一张表读出同样的东西"。

2. **审核规则**（`ReturnAuditor` / `merge_into_oir`）—— 用真材料跑出来的
   `golden/pipeline.template.json` 当模板，配若干份人造回传内容，每条规则都
   带正反两组。审核是交付前最后一道闸，漏判等于把问题交给客户。

字节确定要冻三样东西（重跑两次 shasum 一致）：

* xlsx 里 `docProps/core.xml` 的 created/modified（openpyxl 写 `datetime.now()`）；
* 每个 zip 成员的 mtime（`zipfile` 盖当前时间）—— 与 tools/export_golden.py
  踩过的是同一个坑；两样都在 `_normalize_xlsx` 里**存完再改**；
* `PYTHONHASHSEED` —— `conflict.axis_diff` 返回的是 `pa.keys() & pb.keys()`，
  set 的迭代序随 str 哈希种子变，于是多轴口径分歧的 summary **同一份回传件
  审两次文案就可能不一样**。这是 Python 侧真实存在的不确定性，冻的是导出，
  不是掩盖：TS 侧因为 conflict.ts 固定了轴序而是确定的，两边只比内容。
"""

from __future__ import annotations

import json
import os
import sys
import zipfile
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))

import openpyxl  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from ontocopilot.onto.audit import (  # noqa: E402
    ReturnAuditor,
    merge_into_oir,
    read_returned,
)
from ontocopilot.onto.oir import (  # noqa: E402
    OIR,
    ActionType,
    BusinessRule,
    LinkType,
    ObjectType,
    OpenQuestion,
    PropertyType,
    Provenance,
    extracted,
    inferred,
    make_rid,
)
from ontocopilot.onto.template import ANCHOR_HASH, ANCHOR_RID, TemplateSpec  # noqa: E402

FROZEN = (2026, 8, 13, 20, 27, 35)
OUT_DIR = ROOT / "golden"


# ══════════════════════════════════════════════════════════════════
#  落盘后归一（主 agent 在 tools/export_golden.py 踩过的同一个坑）
# ══════════════════════════════════════════════════════════════════
def _normalize_xlsx(path: Path) -> None:
    """把 openpyxl 存出来的 xlsx 变成字节确定的。

    两处时间戳：`docProps/core.xml` 里的 created/modified（openpyxl 用
    `datetime.now()` 写），以及**每个 zip 成员的 mtime**（zipfile 盖当前时间）。
    不冻这两样，golden 每跑一次 shasum 就变一次，diff 全是噪声。

    这里是**存完再改**，不是猴补 openpyxl —— 猴补要跟着库的内部结构走，
    换个版本就悄悄失效，而失效的表现恰恰是"golden 又不稳定了"。
    """
    import re

    src = zipfile.ZipFile(path)
    members = [(i.filename, src.read(i.filename)) for i in src.infolist()]
    src.close()
    stamp = "{:04d}-{:02d}-{:02d}T{:02d}:{:02d}:{:02d}Z".format(*FROZEN)
    out = []
    for name, data in members:
        if name == "docProps/core.xml":
            text = data.decode("utf-8")
            text = re.sub(r"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</)",
                          rf"\g<1>{stamp}\g<2>", text)
            data = text.encode("utf-8")
        out.append((name, data))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in out:
            info = zipfile.ZipInfo(name, date_time=FROZEN)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            z.writestr(info, data)


# ══════════════════════════════════════════════════════════════════
#  造回传件
# ══════════════════════════════════════════════════════════════════
COLUMNS = ["definition", "owner", "confirmed"]
ROWS = [
    ("pt_amount_budget", "含税，年度累计", "王明", "已确认"),
    ("pt_amount_contract", "不含税，单次", "李强", "已确认"),
    ("pt_plan_id", "计划编号", "王明", "待确认"),
]


def _write_sheet(ws: Any, *, top_pad: int = 0, left_pad: int = 0,
                 anchor: bool = True, rows: list[tuple[str, ...]] | None = None,
                 title_row: str = "") -> None:
    """按模板的落盘形态写一张表：第 1 行填写说明、第 2 行表头、之后是数据。"""
    rows = ROWS if rows is None else rows
    r0 = 1 + top_pad
    c0 = 1 + left_pad
    if title_row:
        ws.cell(row=1, column=1, value=title_row)
    ws.cell(row=r0, column=c0, value="本表填写说明：请逐行核对")
    header = [ANCHOR_RID if anchor else "序号", ANCHOR_HASH, *COLUMNS]
    for i, name in enumerate(header):
        ws.cell(row=r0 + 1, column=c0 + i, value=name)
    for j, row in enumerate(rows):
        ws.cell(row=r0 + 2 + j, column=c0, value=row[0])
        ws.cell(row=r0 + 2 + j, column=c0 + 1, value="hash")
        for i, v in enumerate(row[1:]):
            ws.cell(row=r0 + 2 + j, column=c0 + 2 + i, value=v)


def _book(name: str, build: Any) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    build(wb)
    path = OUT_DIR / f"onto.audit.{name}.xlsx"
    wb.save(path)
    _normalize_xlsx(path)
    return path


def _make_workbooks() -> list[tuple[str, Path]]:
    out: list[tuple[str, Path]] = []

    def plain(wb: Workbook) -> None:
        _write_sheet(wb.create_sheet("01_对象清单"))

    def shifted(wb: Workbook) -> None:
        # 业务方在顶上插了一行标题、在最左边插了一列做批注
        _write_sheet(wb.create_sheet("01_对象清单"), top_pad=1, left_pad=1,
                     title_row="给张总看的汇总")

    def no_anchor(wb: Workbook) -> None:
        # 隐藏列被删掉 —— 整张表读不回来
        _write_sheet(wb.create_sheet("01_对象清单"), anchor=False)

    def duplicated(wb: Workbook) -> None:
        rows = [*ROWS, ("pt_amount_budget", "复制粘贴出来的", "赵六", "已确认")]
        _write_sheet(wb.create_sheet("01_对象清单"), rows=rows)

    def with_note_sheet(wb: Workbook) -> None:
        _write_sheet(wb.create_sheet("01_对象清单"))
        note = wb.create_sheet("我的笔记")
        note.cell(row=1, column=1, value="记一下：这批口径要问财务")
        note.cell(row=2, column=1, value="下周一之前给答复")

    def far_anchor(wb: Workbook) -> None:
        # 锚点被推到扫描窗口之外（第 9 行）—— 扫不到就是扫不到，必须报损伤
        _write_sheet(wb.create_sheet("01_对象清单"), top_pad=7)

    def multi_sheet(wb: Workbook) -> None:
        _write_sheet(wb.create_sheet("01_对象清单"))
        _write_sheet(wb.create_sheet("03_关系清单"),
                     rows=[("lt_a_b", "一对多", "王明", "已确认")])

    def blank_col(wb: Workbook) -> None:
        ws = wb.create_sheet("01_对象清单")
        _write_sheet(ws)
        # 业务方在中间插了一整列空白（列名为空）—— 不该被当成数据列
        ws.insert_cols(4)

    for name, fn in [("plain", plain), ("shifted", shifted), ("no_anchor", no_anchor),
                     ("duplicated", duplicated), ("note_sheet", with_note_sheet),
                     ("far_anchor", far_anchor), ("multi_sheet", multi_sheet),
                     ("blank_col", blank_col)]:
        out.append((name, _book(name, fn)))
    return out


def _grid(path: Path) -> list[dict[str, Any]]:
    """把 xlsx 读成 TS 侧 ReturnedSheet 的形态（openpyxl 的 max_row/max_column）。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    out: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        grid = [["" if ws.cell(row=r, column=c).value is None
                 else str(ws.cell(row=r, column=c).value)
                 for c in range(1, ws.max_column + 1)]
                for r in range(1, ws.max_row + 1)]
        out.append({"title": ws.title, "grid": grid,
                    "maxRow": ws.max_row, "maxColumn": ws.max_column})
    return out


# ══════════════════════════════════════════════════════════════════
#  审核规则的向量
# ══════════════════════════════════════════════════════════════════
def _oir() -> OIR:
    oir = OIR()
    p = Provenance("f1", "采购计划.xlsx", {"kind": "cell"}, snippet="采购包",
                   extractor="docling")
    ot = ObjectType(rid="ot_采购包", api_name=extracted("采购包", p),
                    display_name=extracted("采购包", p))
    oir.add_object(ot)
    for rid, name in [("pt_amount_budget", "计划金额"), ("pt_amount_contract", "合同金额")]:
        oir.add_property(PropertyType(rid=rid, parent=ot.rid,
                                      api_name=extracted(name, p),
                                      display_name=extracted(name, p),
                                      base_type=extracted("DECIMAL", p)))
    oir.add_link(LinkType(rid="lt_a_b", api_name=extracted("包_供", p),
                          source=ot.rid, target=ot.rid,
                          cardinality=inferred("ONE_TO_MANY")))
    oir.add_action(ActionType(rid="at_publish", api_name=extracted("发布", p)))
    oir.add_rule(BusinessRule(rid="br_1", statement=extracted("采购包必须审批", p)))
    oir.add_question(OpenQuestion(rid="oq_1", text=extracted("上限是多少？", p)))
    return oir


def _spec_from_pipeline() -> TemplateSpec:
    return TemplateSpec.from_dict(
        json.loads((OUT_DIR / "pipeline.template.json").read_text("utf-8")))


def _tiny_spec() -> TemplateSpec:
    """一份手搭的小模板，覆盖每条规则的正反两面（真模板里凑不齐全部形态）。"""
    from ontocopilot.onto.template import Cell, Role, Sheet

    sh = Sheet(name="01_对象清单", guide="", columns=[])
    rows: list[dict[str, Cell]] = []

    def cell(rid: str, field: str, value: str, role: Role, **kw: Any) -> Cell:
        return Cell(rid=rid, sheet=sh.name, field=field, value=value, role=role, **kw)

    rows.append({
        "definition": cell("pt_amount_budget", "definition", "含税，年度累计",
                           Role.REQUIRED, owner="王明", expects_prose=True),
        "confirmed": cell("pt_amount_budget", "confirmed", "", Role.REQUIRED,
                          owner="王明", options=["已确认", "待确认"]),
        "apiName": cell("pt_amount_budget", "apiName", "planAmount", Role.PREFILLED,
                        owner="王明"),
    })
    rows.append({
        "definition": cell("pt_amount_contract", "definition", "", Role.REQUIRED,
                           owner="李强", expects_prose=True),
        "confirmed": cell("pt_amount_contract", "confirmed", "", Role.REQUIRED,
                          owner="李强", options=["已确认", "待确认"]),
        "apiName": cell("pt_amount_contract", "apiName", "contractAmount",
                        Role.PREFILLED, owner="李强"),
    })
    rows.append({
        "note": cell("ot_采购包", "note", "只读，不参与完成度", Role.LOCKED),
        "displayName": cell("ot_采购包", "displayName", "采购包", Role.REQUIRED),
    })
    sh.rows = rows
    return TemplateSpec(sheets=[sh])


def _returned(pairs: dict[tuple[str, str], str]) -> dict[tuple[str, str], str]:
    return dict(pairs)


AUDIT_CASES: list[tuple[str, dict[tuple[str, str], str], bool]] = [
    # (label, 回传内容, 是否带 oir)
    ("empty", {}, False),
    ("all_missing", {("pt_amount_budget", "definition"): "",
                     ("pt_amount_budget", "confirmed"): "",
                     ("pt_amount_contract", "definition"): "",
                     ("pt_amount_contract", "confirmed"): "",
                     ("ot_采购包", "displayName"): ""}, False),
    ("all_filled", {("pt_amount_budget", "definition"): "不含税、单次结算，由财务共享中心维护",
                    ("pt_amount_budget", "confirmed"): "已确认",
                    ("pt_amount_contract", "definition"): "含税、年度累计，口径由预算科维护",
                    ("pt_amount_contract", "confirmed"): "待确认",
                    ("ot_采购包", "displayName"): "采购包（新）"}, False),
    ("enum_violation", {("pt_amount_budget", "confirmed"): "差不多确认了",
                        ("pt_amount_contract", "confirmed"): "已确认"}, False),
    ("naming_violation", {("pt_amount_budget", "apiName"): "计划金额",
                          ("pt_amount_contract", "apiName"): "contract_amount"}, False),
    ("naming_violation_with_oir", {("pt_amount_budget", "apiName"): "plan_amount"}, True),
    ("naming_ok", {("pt_amount_budget", "apiName"): "planAmountNew"}, False),
    ("perfunctory_placeholder", {("pt_amount_budget", "definition"): "无",
                                 ("pt_amount_contract", "definition"): "待定"}, False),
    ("perfunctory_too_short", {("pt_amount_budget", "definition"): "税",
                               ("pt_amount_contract", "definition"): "含税口径见附表说明"},
     False),
    ("perfunctory_unchanged_prefill",
     {("pt_amount_budget", "definition"): "含税，年度累计",
      ("pt_amount_contract", "definition"): "不含税、单次结算，由财务共享中心维护"}, False),
    ("perfunctory_copied_header", {("pt_amount_budget", "definition"): "definition"}, False),
    ("perfunctory_bulk_filled",
     {("pt_amount_budget", "definition"): "以实际发生额为准，具体口径参见财务制度",
      ("pt_amount_contract", "definition"): "以实际发生额为准，具体口径参见财务制度"}, False),
    ("enum_column_not_bulk_filled",
     {("pt_amount_budget", "confirmed"): "已确认",
      ("pt_amount_contract", "confirmed"): "已确认"}, False),
    ("divergence_tax", {("pt_amount_budget", "definition"): "含税口径，按发票金额统计",
                        ("pt_amount_contract", "definition"): "不含税口径，按净额统计"},
     False),
    ("divergence_none", {("pt_amount_budget", "definition"): "含税口径，按发票金额统计",
                         ("pt_amount_contract", "definition"): "含税口径，按发票统计"}, False),
    ("new_rows", {("pt_不认识的", "definition"): "业务方自己加的一行",
                  ("pt_amount_budget", "definition"): "含税，按发票"}, False),
    ("damage", {("__damage__", "__damage__"): "「01_对象清单」找不到锚点列 _oir_rid\n第二条损伤"},
     False),
]


def main() -> None:
    # str 的哈希随机化会让 `axis_diff` 的 `pa.keys() & pb.keys()` 每次给出不同的
    # 轴序（见文件头说明）—— 那是 Python 侧真实存在的不确定性，不冻住的话
    # golden 每跑一次就变。**冻的是导出，不是掩盖**：这条分叉照实上报。
    if os.environ.get("PYTHONHASHSEED") != "0":
        os.execve(sys.executable, [sys.executable, *sys.argv],
                  {**os.environ, "PYTHONHASHSEED": "0"})
    out: dict[str, Any] = {}

    # ── 1. 读回传件 ────────────────────────────────────────────
    books = []
    for name, path in _make_workbooks():
        returned = read_returned(path)
        books.append({
            "name": name,
            "file": path.name,
            "sheets": _grid(path),
            "returned": [{"rid": k[0], "field": k[1], "value": v}
                         for k, v in returned.items()],
        })
    out["read_returned"] = books

    # ── 2. 审核规则 ────────────────────────────────────────────
    spec = _tiny_spec()
    out["tiny_spec"] = spec.to_dict()
    cases = []
    for label, pairs, with_oir in AUDIT_CASES:
        oir = _oir() if with_oir else None
        result = ReturnAuditor().audit(spec, _returned(pairs), oir=oir)
        rec: dict[str, Any] = {
            "label": label,
            "returned": [{"rid": k[0], "field": k[1], "value": v} for k, v in pairs.items()],
            "with_oir": with_oir,
            "summary": result.summary(),
            "completeness": result.completeness,
            "findings": [c.to_dict() for c in result.findings],
            "auto_repaired": result.auto_repaired,
            "diffs": [{"rid": d.rid, "sheet": d.sheet, "field": d.field, "before": d.before,
                       "after": d.after, "role": str(d.role), "owner": d.owner,
                       "changed": d.changed, "filled": d.filled,
                       "untouched_prefill": d.untouched_prefill}
                      for d in result.diffs],
            "slips": [s.to_dict() for s in result.slips],
        }
        if with_oir:
            rec["oir_after"] = oir.to_dict() if oir else None
        cases.append(rec)
    out["audit"] = cases

    # ── 3. 真材料模板 + 空回传（完成度的下界）───────────────────
    real = _spec_from_pipeline()
    real_result = ReturnAuditor().audit(real, {}, oir=None)
    out["pipeline_template_empty"] = {
        "summary": real_result.summary(),
        "findings_count": len(real_result.findings),
        "first_findings": [c.to_dict() for c in real_result.findings[:8]],
    }

    # ── 4. merge_into_oir ──────────────────────────────────────
    merge_cases = []
    for label, pairs in [
        ("writes_back", {("pt_amount_budget", "definition"): "不含税、单次",
                         ("ot_采购包", "displayName"): "采购包（改名）"}),
        ("untouched_prefill_not_written",
         {("pt_amount_budget", "definition"): "含税，年度累计"}),
        ("dropped_no_path", {("pt_amount_budget", "confirmed"): "已确认"}),
        ("dropped_unknown_rid", {("pt_不认识的", "definition"): "x"}),
    ]:
        oir = _oir()
        result = ReturnAuditor().audit(spec, _returned(pairs), oir=None)
        changed, dropped = merge_into_oir(oir, result.diffs)
        merge_cases.append({
            "label": label,
            "returned": [{"rid": k[0], "field": k[1], "value": v} for k, v in pairs.items()],
            "changed": changed, "dropped": dropped, "oir_after": oir.to_dict(),
        })
    out["merge"] = merge_cases

    # ── 5. round(x, nd)：完成度用的是 round(x, 4)，half-even 不能猜 ───
    out["py_round"] = [
        {"x": x, "nd": nd, "out": round(x, nd)}
        for x in [0.5, 1.5, 2.5, 0.0625, 0.125, 0.12345, 0.12355, 0.1235, 0.1245,
                  1 / 3, 2 / 3, 0.9995, 0.99995, -0.0001, 0.0, 123.4565, 5 / 7]
        for nd in (0, 2, 3, 4)
    ]

    dst = OUT_DIR / "onto.audit.json"
    dst.write_text(json.dumps(out, ensure_ascii=False, indent=1), "utf-8")
    print(f"wrote {dst} ({dst.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
