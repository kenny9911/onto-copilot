"""onto/parse 的 base + text + api + sql 四个文件的 golden 导出。

TS 侧这四个模块**一个期望值都不许手写** —— 手写的是对 Python 行为的猜测，
这里导的是事实。

导什么、为什么导它：

* ``base``   —— ``content_file_id`` 的内容寻址、``make_chunk`` 的字段布局、
  ``ParsedDoc.stats()`` 里那句 ``len(v) if isinstance(v,(list,dict)) else v``
  （``title`` 是 str 所以**不**取长度，这条最容易在 TS 侧写错），
  以及注册表按扩展名派发 + 没有兜底时的 ValueError 消息（消息里有 ``!r``）。
* ``text``   —— ``_split_sections`` 的软上限与句号回切、``_heading_of``、
  ``_is_heading`` / ``_heading_level``、``_RULE_HINTS`` 的命中与**反例**
  （"金额含税" 不该命中：它是某个业务域的词，不是建模线索）。
  另外故意放了 CRLF、全角数字、代理对（emoji）三条 —— 它们在 Python 与 JS 之间
  分别踩 ``splitlines`` / ``\\d`` / 切片单位三个坑，TS 侧必须自己补齐。
* ``docx``   —— python-docx 抽出来的**原料**（段落 + 样式名、表格单元格、核心属性）
  与最终 ParsedDoc 成对导出。TS 侧的 DocxParser 吃原料、产 ParsedDoc，
  两头都被钉住，中间那段切段逻辑才是真的被测到了。
* ``api``    —— OpenAPI、普通 JSON、顶层不是对象、坏 JSON 四种入口。
* ``sql``    —— ``wire`` 是 sidecar ``/sql/parse`` 原样返回的 dict（app.py:136 那段），
  ``doc`` 是 TS 侧 DdlParser 把它接回来之后应有的样子。两者的差别（file_id /
  file_name / chunk_id 前缀）就是 TS 那层薄封装唯一该做的事。

跑法::

    .venv/bin/python tools/golden/onto_parse.py

输出 ``golden/onto.parse.json``（新文件，我独占）。重跑两次 shasum 必须一致。
"""

from __future__ import annotations

import base64
import json
import sys
import tempfile
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(ROOT / "src"))

from ontocopilot.onto.parse.api import OpenApiParser  # noqa: E402
from ontocopilot.onto.parse.base import (  # noqa: E402
    Finding,
    ParsedDoc,
    ParserRegistry,
    content_file_id,
    make_chunk,
)
from ontocopilot.onto.parse.sql import DdlParser  # noqa: E402
from ontocopilot.onto.parse.text import (  # noqa: E402
    _RULE_HINTS,
    _heading_level,
    _heading_of,
    _is_heading,
    _split_sections,
    DocxParser,
    TextParser,
)

FIXED_ID = "f_0123456789ab"  # 解析出的 file_id 全部显式给定，免得 golden 依赖临时文件字节


# ══════════════════════════════════════════════════════════════════
#  序列化：字段顺序照 dataclass 声明（TS 侧按 toEqual 比，顺序不敏感，
#  但键名一个都不能差 —— sidecar 的线上形态用的就是这套 snake_case 键）
# ══════════════════════════════════════════════════════════════════
def chunk_d(c: Any) -> dict[str, Any]:
    return {"chunk_id": c.chunk_id, "file_id": c.file_id, "file_name": c.file_name,
            "locator": c.locator, "render": c.render, "raw": c.raw, "order": c.order,
            "tags": c.tags, "context": c.context}


def finding_d(f: Finding) -> dict[str, Any]:
    return {"kind": f.kind, "message": f.message, "locator": f.locator,
            "severity": f.severity}


def doc_d(d: ParsedDoc) -> dict[str, Any]:
    return {"file_id": d.file_id, "file_name": d.file_name, "kind": d.kind,
            "chunks": [chunk_d(c) for c in d.chunks],
            "structured": d.structured,
            "findings": [finding_d(f) for f in d.findings],
            "meta": d.meta, "stats": d.stats()}


# ══════════════════════════════════════════════════════════════════
#  base
# ══════════════════════════════════════════════════════════════════
def export_base(tmp: Path) -> dict[str, Any]:
    ids = []
    for name, blob in [("a.txt", b"hello"), ("b.txt", "含税\n".encode()),
                       ("c.bin", bytes(range(256))), ("d.txt", b"")]:
        p = tmp / name
        p.write_bytes(blob)
        ids.append({"name": name, "bytes_b64": base64.b64encode(blob).decode(),
                    "file_id": content_file_id(p)})
    # 读不到内容时按文件名兜底（权限/竞态）—— 目录当路径读会抛 IsADirectoryError(OSError)
    d = tmp / "adir"
    d.mkdir()
    ids.append({"name": "adir", "bytes_b64": None, "file_id": content_file_id(d)})

    chunks = [
        {"in": {"doc_id": "s0", "file_id": "f_x", "file_name": "a.md",
                "locator": {"kind": "page", "page": 1}, "render": "正文"},
         "out": chunk_d(make_chunk(doc_id="s0", file_id="f_x", file_name="a.md",
                                   locator={"kind": "page", "page": 1}, render="正文"))},
        {"in": {"doc_id": "t0r1", "file_id": "f_y", "file_name": "b.docx",
                "locator": {"kind": "ddl", "object": "t"}, "render": "r",
                "raw": {"a": 1}, "order": 7, "tags": ["table", "rule"]},
         "out": chunk_d(make_chunk(doc_id="t0r1", file_id="f_y", file_name="b.docx",
                                   locator={"kind": "ddl", "object": "t"}, render="r",
                                   raw={"a": 1}, order=7, tags=["table", "rule"]))},
    ]

    # stats：structured 里 list/dict 取长度，str/int 原样 —— "title" 那条是判据
    stats = []
    for structured in [
        {},
        {"sections": 3, "chars": 120},
        {"endpoints": [1, 2, 3], "schemas": {"A": {}, "B": {}}, "title": "采购中台 API"},
        {"tables": [], "raw": [1, 2, 3]},          # raw 键被排除
        {"keys": ["a", "b"], "note": None, "ok": True},
    ]:
        doc = ParsedDoc(file_id="f_x", file_name="x.json", kind="openapi")
        doc.chunks.append(make_chunk(doc_id="s0", file_id="f_x", file_name="x.json",
                                     locator={}, render="a"))
        doc.findings.append(Finding("k", "m", {}))
        doc.structured = structured
        stats.append({"structured": structured, "out": doc.stats()})

    reg = (ParserRegistry().register(DdlParser()).register(OpenApiParser())
           .register(DocxParser()).register(TextParser(), fallback=True))
    dispatch = [{"name": n, "kind": reg.for_path(Path(n)).kind} for n in
                ["a.ddl", "a.sql", "a.SQL", "a.json", "a.yaml", "a.yml", "a.docx",
                 "a.DOCX", "a.md", "a.txt", "a.markdown", "a.rst", "a.xlsx", "a",
                 ".bashrc", "a.tar.gz", "a."]]

    bare = ParserRegistry().register(DdlParser())
    try:
        bare.for_path(Path("a.xyz"))
        msg_ascii = ""
    except ValueError as exc:
        msg_ascii = str(exc)
    try:
        bare.for_path(Path("a.中文"))
        msg_cjk = ""
    except ValueError as exc:
        msg_cjk = str(exc)

    return {"content_file_id": ids, "make_chunk": chunks, "stats": stats,
            "dispatch": dispatch,
            "no_parser_message": {"ascii": msg_ascii, "cjk": msg_cjk}}


# ══════════════════════════════════════════════════════════════════
#  text
# ══════════════════════════════════════════════════════════════════
RULE_CASES = [
    "一个执行计划可拆入多个采购包",
    "每个订单必须对应一个客户",
    "references the parent table",
    "Each order must reference exactly one customer",
    "an order can contain multiple line items",
    "金额含税",                      # 反例：业务域词不是建模线索
    "按不含税记录",                  # 反例
    "今天天气不错",                  # 反例
    "PRIMARY KEY 约束",
    "每个ORDER必须",                # 中文紧邻英文：Python 的 \b 认 CJK 为词字符
    "多个each多个",                  # 同上，纯英文线索被中文包住
    # 下面两条是**隔离** \b 差异用的：句子里没有任何中文线索词，命中与否全看
    # \beach\b / \bmany\b 认不认「汉字是词字符」。Python 认（→ 不命中），
    # JS 的 \w 只认 ASCII（→ 会命中）。TS 侧必须自己把 \b 换成 Unicode 版本。
    "订单each订单",
    "很多many很多",
    "订单 each 订单",
    "由采购员生成",
    "由采购员在系统里生成",          # 由.{0,8}生成 的边界（8 个字以内）
    "由采购员在系统的界面里生成",    # 超过 8 个字，不该命中
]

HEADING_CASES = [
    (None, "3.2 采购包组建"),
    ("Heading 1", "采购业务流程说明"),
    ("Heading 2", "3.2 采购包组建"),
    ("标题 3", "金额口径"),
    ("heading 9", "x"),
    ("Normal", "普通段落"),
    ("Normal", "第三章 采购"),
    ("Normal", "第二节 组包"),
    ("Normal", "一、总则"),
    ("Normal", "## 二级标题"),
    ("Normal", "#没有空格"),
    ("Normal", "1.2.3 三级编号"),
    ("Normal", "１.２ 全角数字"),     # Python \d 认全角，JS 不认
    ("标题２", "全角样式号"),          # int("２") 在 Python 是 2，JS 的 Number 给 NaN
    (None, ""),
    ("Normal", "  4. 带前导空格"),
    # BOM 在 Python 的 \s 里**不是**空白（isspace() 为假），JS 的 \s 里是 ——
    # 用 JS 原生 \s 写 _HEADING 的话这一条会凭空变成标题。
    ("Normal", "﻿# 标题"),
    ("Normal", "　一、全角空格开头"),   # 表意空格两边都算空白
]

SPLIT_CASES = [
    "# 标题\n正文一\n正文二\n# 标题二\n正文三",
    "一、总则\n第一条 甲方应当……\n二、附则\n第二条 本办法自发布之日起施行。",
    "没有任何标题的一段话。第二句话。",
    "",
    "   \n\n  ",
    # 软上限（1200）+ 句号回切：前半段长于一半才在句号处切
    ("正" * 700 + "。" + "文" * 700 + "。"),
    # 句号出现在前半段（不到一半），走硬切分支
    ("句。" + "长" * 1400),
    # 真正走到"句号回切"分支的：切点落在串中间，后半截要留给下一块
    "\n".join(["正" * 399 + "。"] * 3 + ["文" * 500]),
    # 切点是 "\n\n" 而不是句号（rfind 取两者的较大者）
    "甲" * 700 + "\n\n" + "乙" * 600,
    # 代理对：cp 计数下 cut(600) 不过半(601) → 走硬切；若按 UTF-16 下标算
    # 则 cut(700) 过半(651) → 走回切，两条分支的结果完全不同。
    "🙂" * 100 + "正" * 500 + "。" + "文" * 601,
    # _heading_of 的 first[:60] 也是按 code point 切的
    "# " + "🙂" * 70 + "\n正文",
    "第一段\r\n第二段\r\n# 标\r\n题下正文",     # CRLF：Python splitlines 吃掉 \r
    "行一\x0b行二\x0c行三 行四",           # Python splitlines 的冷门分隔符
    "带🙂emoji的一行\n" + "尾" * 3,
]


#: `_read_text` 的编码猜测。Node 的 TextDecoder 走 WHATWG/ICU，与 Python 的 codec
#: 在**拒绝哪些字节**上不一样（而且 ICU 有些非法字节是**静默丢掉**而不是报错），
#: 所以这一组要逐条钉：选中的编码不同 = 后面所有切片的文本都不同。
READ_TEXT_CASES: list[bytes] = [
    b"",
    b"hello\n",
    "含税·年度累计".encode(),
    "﻿# 标题".encode("utf-8-sig"),
    "一、总则\n每个订单必须对应一个客户。".encode("gb18030"),
    "採購合約".encode("big5"),
    bytes([0x80, 0x41, 0x0A, 0xFF]),    # 两个中文编码都拒绝 → latin-1
    bytes([0xA1, 0x40]),                # 结构合法：gb18030 给 PUA，big5 给全角空格
    bytes([0x81, 0x30, 0x81, 0x30]),    # gb18030 的四字节形式
    "café".encode("latin-1"),
    bytes([0xE9]),                       # 半个 GBK 双字节 → 三个编码都不成
    bytes([0xC8, 0x40]),                 # big5 的 0xC8 段无映射
]


def export_text(tmp: Path) -> dict[str, Any]:
    from ontocopilot.onto.parse.tabular import _read_text

    read_text = []
    for i, blob in enumerate(READ_TEXT_CASES):
        p = tmp / f"enc{i}.bin"
        p.write_bytes(blob)
        text, enc = _read_text(p)
        read_text.append({"bytes_b64": base64.b64encode(blob).decode(),
                          "text": text, "encoding": enc})

    docs = []
    for name, blob, enc in [
        ("说明.md", "# 采购说明\n一个执行计划可拆入多个采购包。\n\n## 口径\n"
                    "预算控制以年度累计含税金额为准。".encode(), "utf-8"),
        ("gbk.txt", "一、总则\n每个订单必须对应一个客户。".encode("gb18030"), "gb18030"),
        ("bom.txt", "﻿# 标题\n正文".encode("utf-8-sig"), "utf-8-sig"),
        ("latin.txt", bytes([0x80, 0x41, 0x0a, 0xff]), "latin-1"),
        ("emoji.md", "# 🙂 标题\n每个🙂必须有一个🙃。".encode(), "utf-8"),
    ]:
        p = tmp / name
        p.write_bytes(blob)
        docs.append({"name": name, "bytes_b64": base64.b64encode(blob).decode(),
                     "encoding_hint": enc, "file_id": FIXED_ID,
                     "doc": doc_d(TextParser().parse(p, file_id=FIXED_ID))})

    return {
        "rule_hints": [{"s": s, "hit": bool(_RULE_HINTS.search(s))} for s in RULE_CASES],
        "is_heading": [{"style": st, "text": t, "out": _is_heading(st, t)}
                       for st, t in HEADING_CASES],
        "heading_level": [{"style": st, "text": t, "out": _heading_level(st, t)}
                          for st, t in HEADING_CASES],
        "heading_of": [{"block": b, "out": _heading_of(b)} for b in SPLIT_CASES],
        "split_sections": [{"text": t, "out": _split_sections(t)} for t in SPLIT_CASES],
        "read_text": read_text,
        "docs": docs,
    }


# ══════════════════════════════════════════════════════════════════
#  docx —— 原料与成品成对导出
# ══════════════════════════════════════════════════════════════════
def docx_content(path: Path) -> dict[str, Any]:
    """TS 侧 DocxExtractor 应该交出来的东西 —— python-docx 能给、JS 给不了的部分。

    ``modified`` 在这里就 ``str()`` 掉：``str(datetime)`` 的格式属于 Python，
    留在 Python 侧算，TS 只做"非空就报泄漏"的判断。
    """
    import docx

    d = docx.Document(str(path))
    core = d.core_properties
    return {
        "paragraphs": [{"text": p.text, "style": p.style.name if p.style else None}
                       for p in d.paragraphs],
        "tables": [[[c.text for c in r.cells] for r in t.rows] for t in d.tables],
        "core": {"creator": core.author,
                 "last_modified_by": core.last_modified_by,
                 "modified": str(core.modified) if core.modified else None,
                 "title": core.title},
    }


def build_flow_docx(path: Path) -> None:
    """conftest.py 的 ``flow_docx``，一字不改地搬过来（那份是测试的权威输入）。"""
    import docx

    d = docx.Document()
    d.core_properties.author = "wubin"
    d.add_heading("采购业务流程说明", level=1)
    d.add_heading("3.2 采购包组建", level=2)
    d.add_paragraph(
        "执行计划下达后，采购员按物料类别、供应商能力、交付窗口三个因子组包。"
        "一个执行计划可拆入多个采购包，一个采购包也可包含多个执行计划的行项。"
        "组包完成后进入询比价环节。")
    d.add_heading("4.1 金额口径", level=2)
    d.add_paragraph(
        "预算控制以年度累计含税金额为准；合同签订时按不含税单次金额记录，"
        "两者之间的换算由财务共享中心维护税率表。")
    t = d.add_table(rows=3, cols=3)
    for c, h in enumerate(["字段", "所属对象", "口径"]):
        t.cell(0, c).text = h
    t.cell(1, 0).text = "plan_amount"
    t.cell(1, 1).text = "pbpHeader"
    t.cell(1, 2).text = "含税年度累计"
    t.cell(2, 0).text = "plan_amount"
    t.cell(2, 1).text = "clmContract"
    t.cell(2, 2).text = "不含税单次"
    d.save(path)


def build_rules_docx(path: Path) -> None:
    """tests/test_parse.py 的英文规则句用例。"""
    import docx

    d = docx.Document()
    d.add_heading("Rules", level=1)
    d.add_paragraph("Each order must reference exactly one customer; "
                    "an order can contain multiple line items.")
    d.save(path)


def build_edge_docx(path: Path) -> None:
    """三条 TS 侧独有风险：空表（整表被丢）、列多于表头、超软上限的连续段落。"""
    import docx

    d = docx.Document()
    d.core_properties.author = ""              # 空作者 → meta 里不该出现
    d.core_properties.last_modified_by = "李强"  # 只有这一条泄漏
    d.add_heading("标题", level=1)
    d.add_paragraph("短段。")
    d.add_paragraph("正" * 800)
    d.add_paragraph("文" * 800)                 # 累计过 1200 → 前面的先落一块
    d.add_paragraph("")                          # 空段落被跳过
    t0 = d.add_table(rows=2, cols=2)             # 全空表 → 整张丢掉
    t0.cell(0, 0).text = ""
    t1 = d.add_table(rows=3, cols=3)
    t1.cell(0, 0).text = "字段"                  # 表头只有一列有值
    t1.cell(1, 0).text = "a"
    t1.cell(1, 1).text = "b"                     # 越过表头长度的值被丢
    t1.cell(2, 2).text = "c"
    d.save(path)


def export_docx(tmp: Path) -> list[dict[str, Any]]:
    out = []
    for name, build in [("流程说明.docx", build_flow_docx), ("rules.docx", build_rules_docx),
                        ("edge.docx", build_edge_docx)]:
        p = tmp / name
        build(p)
        out.append({"name": name, "file_id": FIXED_ID, "content": docx_content(p),
                    "doc": doc_d(DocxParser().parse(p, file_id=FIXED_ID))})
    # 仓库里那份真材料 —— 构造出来的 docx 覆盖不到真实 Word 的样式名与合并单元格
    real = ROOT / "materials" / "流程说明.docx"
    if real.exists():
        out.append({"name": "materials/流程说明.docx", "file_id": FIXED_ID,
                    "content": docx_content(real),
                    "doc": doc_d(DocxParser().parse(real, file_id=FIXED_ID))})
    return out


# ══════════════════════════════════════════════════════════════════
#  api
# ══════════════════════════════════════════════════════════════════
CONFTEST_OPENAPI = {
    "openapi": "3.0.0",
    "info": {"title": "采购中台 API"},
    "paths": {
        "/purchase-plans/{id}/submit": {
            "post": {"operationId": "submitPurchasePlan", "summary": "提交采购计划",
                     "requestBody": {"content": {"application/json": {
                         "schema": {"$ref": "#/components/schemas/SubmitReq"}}}}}},
        "/purchase-packages": {
            "post": {"operationId": "createPurchasePackage", "summary": "创建采购包"}},
        "/suppliers": {"get": {"operationId": "listSuppliers"}},
    },
    "components": {"schemas": {
        "Plan": {"type": "object", "required": ["planId"], "properties": {
            "planId": {"type": "string"},
            "planAmount": {"type": "number", "format": "double",
                           "description": "计划金额（含税，年度累计）"},
            "effectiveDate": {"type": "string", "format": "date"},
            "status": {"type": "string", "enum": ["DRAFT", "SUBMITTED"]},
        }},
        "SubmitReq": {"type": "object", "properties": {"remark": {"type": "string"}}},
    }},
}

#: 边界：没有 operationId（合成 id）、DELETE、路径参数、HEAD（不收）、
#: ops 不是 dict、op 不是 dict、summary 超 120、definitions 版 schema、
#: enum 里混布尔与数字（str(x) 给 "True"）、无写操作端点
EDGE_OPENAPI = {
    "info": {},
    "paths": {
        "/a/{id}/b": {"delete": {"summary": "删", "tags": ["x"]},
                      "head": {"operationId": "ignored"},
                      "patch": {"description": "只有 description " + "长" * 130}},
        "/": {"put": {}},
        "//weird//": {"post": {}},
        "/notadict": "oops",
        "/opnotadict": {"post": 3},
    },
    "definitions": {
        "Legacy": {"required": ["a"], "properties": {
            "a": {"type": "integer"},
            "b": {"type": "string", "format": "date-time", "description": "时间"},
            "c": {"type": "boolean", "enum": [True, False, 1, 2.5, None]},
            "d": {"type": "string", "format": "uuid"},
            "e": {"type": "array", "items": {"type": "string"}},
            "f": "notadict",
        }},
        "NotObject": {"type": "string"},
        "Arr": ["x"],
    },
}

READONLY_OPENAPI = {"info": {"title": "只读"}, "paths": {"/x": {"get": {}}},
                    "components": {"schemas": {}}}

PLAIN_JSON = {"名单": ["甲", "乙"], "数量": 3, "嵌套": {"a": 1.0, "b": None, "c": True},
              "长文本": "很长" * 400}


def export_api(tmp: Path) -> list[dict[str, Any]]:
    cases: list[tuple[str, str]] = [
        ("openapi.json", json.dumps(CONFTEST_OPENAPI, ensure_ascii=False)),
        ("edge.json", json.dumps(EDGE_OPENAPI, ensure_ascii=False)),
        ("readonly.json", json.dumps(READONLY_OPENAPI, ensure_ascii=False)),
        ("plain.json", json.dumps(PLAIN_JSON, ensure_ascii=False)),
        ("toplevel_list.json", json.dumps([1, 2], ensure_ascii=False)),
        ("broken.json", "{不是 JSON"),
    ]
    real = ROOT / "materials" / "openapi.json"
    if real.exists():
        cases.append(("materials/openapi.json", real.read_text(encoding="utf-8")))

    out = []
    for name, text in cases:
        p = tmp / Path(name).name
        p.write_text(text, encoding="utf-8")
        out.append({"name": name, "text": text, "file_id": FIXED_ID,
                    "doc": doc_d(OpenApiParser().parse(p, file_id=FIXED_ID))})
    return out


# ══════════════════════════════════════════════════════════════════
#  sql —— wire（sidecar 原样返回）+ doc（TS 侧接回来之后应有的样子）
# ══════════════════════════════════════════════════════════════════
CONFTEST_DDL = """-- 采购中台物理模型
CREATE TABLE pbp_header (
  plan_id      VARCHAR(32) NOT NULL PRIMARY KEY,
  plan_name    VARCHAR(200),
  plan_amount  DECIMAL(18,2),   -- 含税·年度累计·CNY
  created_at   TIMESTAMP
);

CREATE TABLE clm_contract (
  contract_id  VARCHAR(32) NOT NULL PRIMARY KEY,
  plan_id      VARCHAR(32) NOT NULL,
  plan_amount  DECIMAL(18,2),   -- 不含税·单次·CNY
  CONSTRAINT fk_plan FOREIGN KEY (plan_id) REFERENCES pbp_header(plan_id)
);
"""

COMPOSITE_DDL = """CREATE TABLE line (
  a VARCHAR(8) COMMENT '行内注释口径',
  b VARCHAR(8),
  PRIMARY KEY (a, b),
  CONSTRAINT fk_ab FOREIGN KEY (a, b) REFERENCES head(x, y)
);
"""

NO_TABLE_DDL = "SELECT 1;\n"
BROKEN_DDL = "CREATE TABLE ( ( ( ;\n"


def wire_of(doc: ParsedDoc) -> dict[str, Any]:
    """sidecar/app.py:136-143 那段的原样复制 —— 线上形态没有 file_id/file_name/context。"""
    return {
        "file_id": doc.file_id, "file_name": doc.file_name, "kind": doc.kind,
        "structured": doc.structured,
        "chunks": [{"chunk_id": c.chunk_id, "locator": c.locator, "render": c.render,
                    "raw": c.raw, "order": c.order, "tags": c.tags} for c in doc.chunks],
        "findings": [{"kind": f.kind, "message": f.message, "locator": f.locator,
                      "severity": f.severity} for f in doc.findings],
    }


def export_sql(tmp: Path) -> list[dict[str, Any]]:
    cases = [("schema.ddl", CONFTEST_DDL, None), ("composite.ddl", COMPOSITE_DDL, None),
             ("empty.ddl", NO_TABLE_DDL, None), ("broken.ddl", BROKEN_DDL, None),
             ("mysql.ddl", CONFTEST_DDL, "mysql")]
    real = ROOT / "materials" / "schema.ddl"
    if real.exists():
        cases.append(("materials/schema.ddl", real.read_text(encoding="utf-8"), None))

    out = []
    for name, sql, dialect in cases:
        base = Path(name).name
        # wire：sidecar 那边 file_id 恒为 "sidecar"，file_name 是请求里给的名字
        wp = tmp / base
        wp.write_text(sql, encoding="utf-8")
        wire = wire_of(DdlParser(dialect=dialect).parse(wp, file_id="sidecar"))
        # doc：同一份 DDL 用真实 file_id 解析 —— TS 侧那层薄封装要把 wire 映成这个
        doc = doc_d(DdlParser(dialect=dialect).parse(wp, file_id=FIXED_ID))
        out.append({"name": name, "sql": sql, "dialect": dialect,
                    "file_id": FIXED_ID, "file_name": base, "wire": wire, "doc": doc})
    return out


def main() -> None:
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        payload = {
            "base": export_base(tmp),
            "text": export_text(tmp),
            "docx": export_docx(tmp),
            "api": export_api(tmp),
            "sql": export_sql(tmp),
        }
    out = ROOT / "golden" / "onto.parse.json"
    out.write_text(json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=1) + "\n",
                   encoding="utf-8")
    print(f"wrote {out} ({out.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
