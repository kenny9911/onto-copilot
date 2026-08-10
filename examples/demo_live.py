"""真实 LLM 端到端演示 —— 接自定义聚合网关，跑完整往返闭环。

    原始材料 → LLM 抽取(真实调用) → 跨厂商 Critic 评审 → 冲突检测
    → 澄清排序 → 编译 xlsx 模板 → 模拟业务方填写 → 回传审核 → 打回单

与 ``demo_pipeline.py`` 的区别：那个用 ScriptedBackend 验编排骨架，这个用真实
模型验**抽取质量、跨厂商评委、结构化输出、真实成本**。

需要 ``.env``（见 ``.env.example``）::

    python examples/demo_live.py
"""

from __future__ import annotations

import asyncio
import json
from pathlib import Path

from ontocopilot.kernel.backends import OpenAICompatBackend
from ontocopilot.kernel.budget import Budget
from ontocopilot.kernel.config import llm_config
from ontocopilot.kernel.critic import CriticContext, CriticPanel, LLMCritic, RuleCritic
from ontocopilot.kernel.dag import Difficulty
from ontocopilot.kernel.journal import InMemoryBlobStore, InMemoryJournal
from ontocopilot.kernel.llm import GATEWAY_MODELS, ModelGateway, gateway_routing
from ontocopilot.kernel.recorder import Recorder
from ontocopilot.onto.audit import ReturnAuditor, merge_into_oir, read_returned
from ontocopilot.onto.clarify import ClarificationEngine, apply_decision
from ontocopilot.onto.conflict import auto_repair, detect_all
from ontocopilot.onto.oir import (
    OIR,
    BaseType,
    Cardinality,
    LinkType,
    ObjectType,
    PropertyType,
    Provenance,
    Status,
    extracted,
    inferred,
    make_rid,
)
from ontocopilot.onto.template import Role, compile_template, write_xlsx

RUN_ID = "run_live_ont112"
OUT = Path("artifacts")

# ══════════════════════════════════════════════════════════════════
#  原始材料 —— 故意保留真实梳理表里的杂乱（中英混排、口径藏在括号里）
# ══════════════════════════════════════════════════════════════════
MATERIALS = {
    "实体梳理.xlsx": """业务对象实体梳理（第 2/9/44/81 行）
行2 | 采购需求计划 | 采购业务计划头 | pbpHeader | 年度滚动，按物料大类拆分
行9 | 采购包    | 采购包头     | purchasePackage | 由多个执行计划行合并而成
行44| 采购合同   | clmContract  | 计划金额（含税，年度累计，CNY）
行81| 采购订单头  | poHeader     | 由采购包下达""",
    "schema.ddl": """CREATE TABLE pbp_header (
  plan_id      VARCHAR(32) PRIMARY KEY,
  plan_name    VARCHAR(200),
  plan_amount  DECIMAL(18,2)   -- 含税·年度累计
);
CREATE TABLE clm_contract (
  contract_id  VARCHAR(32) PRIMARY KEY,
  plan_id      VARCHAR(32),
  plan_amount  DECIMAL(18,2),  -- 不含税·单次
  CONSTRAINT fk_plan FOREIGN KEY (plan_id) REFERENCES pbp_header(plan_id)
);""",
    "openapi.json": """POST /purchase-plans/{id}/submit  operationId=submitPurchasePlan
POST /purchase-packages           operationId=createPurchasePackage
GET  /suppliers                   operationId=listSuppliers""",
}

ENDPOINTS = [
    {"operationId": "submitPurchasePlan", "method": "post",
     "path": "/purchase-plans/{id}/submit", "file_id": "f1", "file_name": "openapi.json",
     "pointer": "$.paths./purchase-plans/{id}/submit"},
    {"operationId": "createPurchasePackage", "method": "post", "path": "/purchase-packages",
     "file_id": "f1", "file_name": "openapi.json", "pointer": "$.paths./purchase-packages"},
]

EXTRACT_SCHEMA = {
    "type": "object",
    "required": ["objects", "properties", "links"],
    "properties": {
        "objects": {
            "type": "array",
            "items": {
                "type": "object",
                "required": ["api_name", "display_name", "source_file", "source_locator"],
                "properties": {
                    "api_name": {"type": "string"},
                    "display_name": {"type": "string"},
                    "source_file": {"type": "string"},
                    "source_locator": {"type": "string"},
                },
            },
        },
        "properties": {
            "type": "array",
            "items": {
                "type": "object",
                "required": ["parent_api_name", "api_name", "display_name", "base_type",
                             "definition", "source_file", "source_locator"],
                "properties": {
                    "parent_api_name": {"type": "string"},
                    "api_name": {"type": "string"},
                    "display_name": {"type": "string"},
                    "base_type": {"type": "string",
                                  "enum": [t.value for t in BaseType]},
                    "definition": {"type": "string"},
                    "source_file": {"type": "string"},
                    "source_locator": {"type": "string"},
                },
            },
        },
        "links": {
            "type": "array",
            "items": {
                "type": "object",
                "required": ["api_name", "from_api_name", "to_api_name", "cardinality",
                             "source_file", "source_locator"],
                "properties": {
                    "api_name": {"type": "string"},
                    "from_api_name": {"type": "string"},
                    "to_api_name": {"type": "string"},
                    "cardinality": {"type": "string",
                                    "enum": ["ONE_TO_ONE", "ONE_TO_MANY", "MANY_TO_MANY"]},
                    "source_file": {"type": "string"},
                    "source_locator": {"type": "string"},
                },
            },
        },
    },
}

SYSTEM = """你是 OntoCopilot 的抽取器，为 FDE 工程师从业务材料中抽取本体候选。

规范：
- apiName 用 lowerCamelCase，物理名与业务名冲突时以物理名为准，业务名进 displayName
- definition 写这个字段到底统计什么。材料里若标注了含税/不含税、时间粒度，必须原样保留
- source_file / source_locator 必须指向你实际读到该信息的位置，不许编造
- 同一个概念在不同材料里口径不同时，**分别抽出、各自保留原口径**，不要擅自统一"""


def _prov(file_name: str, locator: str, snippet: str) -> Provenance:
    kind = ("ddl" if file_name.endswith(".ddl")
            else "json" if file_name.endswith(".json") else "cell")
    loc = ({"kind": "ddl", "object": locator} if kind == "ddl"
           else {"kind": "json", "pointer": locator} if kind == "json"
           else {"kind": "cell", "sheet": "业务对象实体梳理", "row": locator, "col": "F"})
    return Provenance("f_" + file_name[:3], file_name, loc, snippet=snippet,
                      extractor="llm", confidence=0.8)


#: 责任人按业务域分派。真实系统里来自项目组织架构表，这里用关键词近似。
DOMAIN_OWNERS = (("plan", "王明"), ("package", "王明"), ("contract", "李强"),
                 ("po", "李强"), ("supplier", "张莉"))


def _owner_of(api_name: str) -> str:
    low = api_name.lower()
    return next((o for k, o in DOMAIN_OWNERS if k in low), "王明")


def build_oir(data: dict) -> OIR:
    """把 LLM 抽取结果落成 OIR。结构组装是确定性的，不交给模型。"""
    oir = OIR()
    by_api: dict[str, str] = {}
    for o in data.get("objects", []):
        ev = _prov(o["source_file"], o["source_locator"], o["display_name"])
        rid = make_rid("ot", o["api_name"])
        oir.add_object(ObjectType(
            rid=rid, api_name=extracted(o["api_name"], ev),
            display_name=extracted(o["display_name"], ev),
            primary_key=inferred([]), owner=_owner_of(o["api_name"])))
        by_api[o["api_name"].lower()] = rid

    for i, p in enumerate(data.get("properties", [])):
        parent = by_api.get(p["parent_api_name"].lower())
        if parent is None:
            continue
        ev = _prov(p["source_file"], p["source_locator"], p["definition"] or p["display_name"])
        oir.add_property(PropertyType(
            rid=make_rid("pt", f"{p['parent_api_name']}_{p['api_name']}_{i}"),
            parent=parent,
            api_name=extracted(p["api_name"], ev),
            display_name=extracted(p["display_name"], ev),
            base_type=extracted(BaseType(p["base_type"]), ev),
            definition=extracted(p["definition"], ev),
            owner=_owner_of(p["parent_api_name"])))

    for l in data.get("links", []):
        src, tgt = by_api.get(l["from_api_name"].lower()), by_api.get(l["to_api_name"].lower())
        if not src or not tgt:
            continue
        ev = _prov(l["source_file"], l["source_locator"], l["api_name"])
        oir.add_link(LinkType(
            rid=make_rid("lt", l["api_name"]), api_name=extracted(l["api_name"], ev),
            source=src, target=tgt,
            cardinality=extracted(Cardinality(l["cardinality"]), ev),
            join_key=inferred(None)))

    # 主键：名字里带 id 且是第一个的，作为候选主键（确定性启发式）
    for o in oir.objects.values():
        pk = [r for r in o.properties
              if oir.properties[r].api_name.value.lower().endswith("id")]
        if pk:
            o.primary_key = inferred(pk[:1])
    return oir


async def main() -> None:
    cfg = llm_config()
    journal, blobs = InMemoryJournal(), InMemoryBlobStore()
    rec = Recorder(RUN_ID, journal, blobs)
    budget = Budget(tokens=2_000_000, usd=5.0)
    backend = OpenAICompatBackend(cfg.base_url, cfg.api_key)
    gw = ModelGateway(backend, rec, routing=gateway_routing(), budget=budget)

    bar = "═" * 76
    print(bar)
    print("  OntoCopilot 真实 LLM 端到端 ·", cfg.base_url)
    print(f"  凭证 {cfg.redacted_key}" + ("  ⚠ 明文 HTTP" if cfg.insecure_transport else ""))
    print(bar)

    # ── 1. 真实抽取 ────────────────────────────────────────────
    corpus = "\n\n".join(f"⟦{n}⟧\n{t}" for n, t in MATERIALS.items())
    print("\n▸ 1/6 抽取 · 3 份异构材料 → OIR")
    comp = await gw.call(
        "EXTRACT",
        f"从下面的材料中抽取业务对象与属性。\n\n{corpus}",
        system=SYSTEM, difficulty=Difficulty.HIGH, schema=EXTRACT_SCHEMA, max_tokens=6000,
    )
    oir = build_oir(comp.data)
    print(f"  模型 {comp.model} · {comp.usage.tok_in}→{comp.usage.tok_out} tok · ${comp.usd:.4f}")
    print(f"  产出 {oir.stats()}")
    for p in list(oir.properties.values())[:6]:
        print(f"    · {p.api_name.value:16} {p.definition.value[:34]:36} "
              f"← {p.definition.evidence[0].cite() if p.definition.evidence else '无出处'}")

    # ── 2. 跨厂商 Critic ───────────────────────────────────────
    print("\n▸ 2/6 评审 · 规则视角 + 跨厂商语义评委")
    panel = CriticPanel({
        "schema": RuleCritic("schema", lambda d: []),
        "provenance": LLMCritic("provenance", [
            "每个 definition 都指向了材料里真实存在的位置",
            "含税/不含税、时间粒度等口径限定词被原样保留，没有被擅自统一",
            "apiName 全部符合 lowerCamelCase",
        ], instruction="出处格式形如 文件名!位置 或 文件名#对象名。核对不到就判不通过。"),
    }, rec)
    generator = GATEWAY_MODELS["frontier"]
    verdicts = await panel.judge(
        comp.data, ["schema", "provenance"],
        CriticContext(node_id="CRITIC", gateway=gw, generator=generator,
                      evidence_render=corpus, samples=1),
    )
    for v in verdicts:
        judge = "规则" if v.lens == "schema" else gw.routing.judge_for(generator).name
        print(f"  [{v.lens}] {'通过' if v.passed else '未通过'} · 评委 {judge} · {v.note}")
        for f in v.findings[:3]:
            print(f"      ! {f.code} {f.claim[:70]}")

    # ── 3. 冲突 + 澄清 ─────────────────────────────────────────
    print("\n▸ 3/6 冲突检测与澄清排序")
    conflicts = detect_all(oir, endpoints=ENDPOINTS)
    repaired = auto_repair(oir, conflicts)
    cs = ClarificationEngine(max_questions=3).rank(conflicts, oir)
    kinds: dict[str, int] = {}
    for c in conflicts:
        kinds[str(c.kind)] = kinds.get(str(c.kind), 0) + 1
    print(f"  冲突 {len(conflicts)} 条 {kinds} · 自动修 {len(repaired)}")
    print(f"  路由 {cs.summary()}")
    for q in cs.questions:
        print(f"    ── {q.title[:72]}")
        print(f"       影响 {q.impact} · 得分 {q.score:.3f} · "
              f"{'不可逆' if not q.reversible else '可逆'}")

    div = next((c for c in conflicts if str(c.kind) == "semantic_divergence"), None)
    if div:
        applied = apply_decision(oir, div, "split_two_properties", note="与合同部当面对齐")
        print(f"  FDE 拍板 → {applied['label'][:64]}")

    # ── 4. 编译模板 ────────────────────────────────────────────
    print("\n▸ 4/6 编译模板")
    spec = compile_template(oir, conflicts)
    path = write_xlsx(spec, OUT / "采购中台实体模板_v1.xlsx", project="ONT-112")
    s = spec.stats()
    print(f"  {path} · {s['total_cells']} 格 · 预填 {s['prefill_rate']:.0%} · "
          f"业务必填 {s['business_required']} · 责任人 {s['owners']}")

    # ── 5. 模拟业务方填写（含真实世界的糟糕填法）────────────────
    print("\n▸ 5/6 模拟业务方回传")
    from openpyxl import load_workbook

    ret_path = OUT / "采购中台实体模板_回传_v1.xlsx"
    wb = load_workbook(path)
    filled = {"good": 0, "perfunctory": 0, "blank": 0}
    for ws in wb.worksheets:
        if ws.title.startswith("00_") or ws.cell(row=2, column=1).value != "_oir_rid":
            continue
        headers = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
        for r in range(3, ws.max_row + 1):
            rid = ws.cell(row=r, column=1).value
            if not rid:
                continue
            for c, name in enumerate(headers, start=1):
                cell = spec.by_rid().get((str(rid), str(name)))
                if cell is None or cell.role is not Role.REQUIRED:
                    continue
                i = (r + c) % 5
                if i == 0:  # 抄列名
                    ws.cell(row=r, column=c, value=str(name)); filled["perfunctory"] += 1
                elif i == 1:  # 留空
                    filled["blank"] += 1
                elif i == 2 and cell.value:  # 原样交回预填
                    ws.cell(row=r, column=c, value=cell.value); filled["perfunctory"] += 1
                else:
                    v = ({"definition": f"不含税、单次结算、CNY（{rid} 由财务共享中心维护税率表）",
                          "description": f"{rid} 在采购流程中的业务实体",
                          "owner": "王明" if r % 2 else "李强",
                          "cardinality": "MANY_TO_MANY", "confirmed": "已确认",
                          "correct": "正确", "effects": "CREATE"}
                         .get(str(name), "已确认"))
                    ws.cell(row=r, column=c, value=v); filled["good"] += 1
    wb.save(ret_path)
    print(f"  {ret_path} · 认真填 {filled['good']} · 敷衍 {filled['perfunctory']} · "
          f"留空 {filled['blank']}")

    # ── 6. 回传审核 ────────────────────────────────────────────
    print("\n▸ 6/6 回传审核")
    result = ReturnAuditor().audit(spec, read_returned(ret_path), oir=oir)
    merged = merge_into_oir(oir, result.diffs)
    print(f"  完成度 {result.completeness:.0%}（达标线 95%）")
    print(f"  单元格 比对 {len(result.diffs)} · 改动 "
          f"{sum(1 for d in result.diffs if d.changed)}")
    print(f"  问题 {result.counts} · 自动修 {len(result.auto_repaired)}")
    print(f"  写回 OIR {len(merged)} 处（原样交回的预填不算填写）")
    for slip in result.slips:
        print(f"    打回 · {slip.owner} · {len(slip.items)} 项")
        for kind, items in slip.to_dict()["by_kind"].items():
            print(f"        {kind}: {len(items)} 条，例：{items[0][:56]}")

    (OUT / "audit.json").write_text(
        json.dumps(result.summary(), ensure_ascii=False, indent=1), encoding="utf-8")

    # ── 账单 ───────────────────────────────────────────────────
    snap = budget.snapshot()
    print(f"\n▸ 账单 · {snap['spent']['tokens']:.0f} tokens · "
          f"${snap['spent']['usd']:.4f}（网关回报的真实成本）· {snap['level_label']}")
    print(f"▸ 事件日志 {len(list(journal.read(RUN_ID)))} 条 · 产物目录 {OUT}/")
    print(bar)
    await backend.aclose()


if __name__ == "__main__":
    asyncio.run(main())
